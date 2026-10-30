"""Exercise the actual exe with synthetic data and no Python/conda on PATH."""

from __future__ import annotations

import argparse
import io
import json
import os
import socket
import subprocess
import time
from pathlib import Path
from uuid import uuid4

import httpx
import yaml
from docx import Document
from generate_doc2md_samples import _write_native_pdf
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--exe", type=Path, default=ROOT / "dist/DocQA/DocQA.exe")
    args = parser.parse_args()
    exe = args.exe.resolve()
    home = ROOT / "runtime" / ("exe-smoke-便携测试 " + uuid4().hex[:8])
    home.mkdir(parents=True)
    environment = {
        key: value for key, value in os.environ.items()
        if key.upper() in {"SYSTEMROOT", "WINDIR", "SYSTEMDRIVE", "TEMP", "TMP", "COMSPEC"}
    }
    environment["PATH"] = str(Path(os.environ["SYSTEMROOT"]) / "System32")
    checked = subprocess.run(
        [str(exe), "--check"], cwd=home, env=environment, capture_output=True, timeout=60,
    )
    assert checked.returncode == 0, checked.stderr.decode("utf-8", errors="replace")
    assert b'"frozen": true' in checked.stdout
    configs = home / "configs"
    configs.mkdir()
    for name in ("web.demo.example.yaml", "service.demo.example.yaml"):
        value = yaml.safe_load((ROOT / "configs" / name).read_text(encoding="utf-8"))
        if name.startswith("web"):
            value.update(reuse_workers=True, query_workers=1)
        (configs / name).write_text(yaml.safe_dump(value), encoding="utf-8")
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    url = f"http://127.0.0.1:{port}"
    log = home / "launcher.log"
    checks = ["native dependencies without Python/conda PATH"]
    inputs = home / "synthetic-inputs"
    inputs.mkdir()
    _write_native_pdf(inputs / "native.pdf")
    Image.new("RGB", (400, 300), "white").save(inputs / "scan.pdf", "PDF")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "PPTX TEST 30s"
    presentation.save(inputs / "sample.pptx")
    workbook = Workbook()
    workbook.active.append(["Model", "Interval"])
    workbook.active.append(["SD5993", 30])
    workbook.save(inputs / "sample.xlsx")
    (inputs / "sample.csv").write_text("Model,Interval\nSD5993,30\n", encoding="utf-8")
    parse_config = home / "parse-check.yaml"
    parse_config.write_text("office_render: never\nocr_mode: disabled\n", encoding="utf-8")
    for source in sorted(inputs.iterdir()):
        parsed = subprocess.run([
            str(exe), "--home", str(home), "--cli", "doc2md", "parse", str(source),
            "--output", str(home / "parsed-check"), "--allow-root", str(inputs),
            "--config", str(parse_config),
        ], cwd=home, env=environment, capture_output=True, timeout=60)
        assert parsed.returncode == 0, (source.name, parsed.stdout, parsed.stderr)
    checks.append("frozen CLI: native PDF, scanned PDF/PDFium, PPTX, XLSX and CSV parsing")
    for module, error in (("docqa.web.worker", b"INVALID_REQUEST"),
                          ("docqa.web.visual", b"VISION_CHECK_FAILED")):
        rejected = subprocess.run([str(exe), "--internal-worker", module], input=b"{}",
            cwd=home, env=environment, capture_output=True, timeout=30)
        assert rejected.returncode == 4 and error in rejected.stdout, rejected.stdout
    checks.append("single-request Agent and visual-worker protocol routing")
    with log.open("wb") as output:
        process = subprocess.Popen([
            str(exe), "--demo", "--home", str(home), "--no-browser", "--port", str(port),
        ], cwd=home, env=environment, stdout=output, stderr=output, stdin=subprocess.DEVNULL,
            creationflags=subprocess.CREATE_NO_WINDOW)
        try:
            with httpx.Client(base_url=url, timeout=30, trust_env=False) as client:
                deadline = time.monotonic() + 90
                while True:
                    assert process.poll() is None, log.read_text(encoding="utf-8", errors="replace")
                    try:
                        response = client.get("/api/v1/health/live")
                        if response.status_code == 200:
                            break
                    except httpx.TransportError:
                        pass
                    assert time.monotonic() < deadline, "Frozen server startup timed out"
                    time.sleep(0.25)
                bootstrap = client.get("/api/v1/bootstrap").json()
                assert bootstrap["settings"]["synthetic"]
                client.headers["X-DocQA-Token"] = bootstrap["csrf_token"]
                assert client.get("/").status_code == 200
                for path in ("/static/app.js", "/static/js/main.js", "/static/styles.css"):
                    assert client.get(path).status_code == 200, path
                checks.append("frozen demo parsing, SQLite-vec build, Web and static assets")

                def wait(path, states):
                    limit = time.monotonic() + 90
                    while time.monotonic() < limit:
                        response = client.get(path)
                        assert response.status_code == 200, response.text
                        value = response.json()
                        if value["state"] in states:
                            return value
                        assert value["state"] not in {"failed", "cancelled", "upload_failed"}, value
                        time.sleep(0.3)
                    raise AssertionError("Frozen task timed out: " + path)

                for round_number in range(2):
                    cid = client.post("/api/v1/conversations", json={}).json()["id"]
                    response = client.post(f"/api/v1/conversations/{cid}/messages", json={
                        "question": "SD5993 心跳如何启动？", "client_request_id": uuid4().hex,
                    })
                    assert response.status_code == 202, response.text
                    job = wait("/api/v1/jobs/" + response.json()["job_id"], {"succeeded"})
                    assert job["answer"]["citations"], job
                    assert job["answer"]["synthetic"]
                    checks.append(f"persistent frozen Agent request {round_number + 1}, with citations")
                document = Document()
                document.add_paragraph("FROZEN_UPLOAD SD5993 upload ingestion test")
                data = io.BytesIO()
                document.save(data)
                raw = data.getvalue()
                response = client.post("/api/v1/uploads", json={
                    "name": "打包上传验证.docx", "size": len(raw),
                })
                assert response.status_code == 201, response.text
                path = "/api/v1/uploads/" + response.json()["id"]
                assert client.put(path + "/content", content=raw,
                    headers={"Content-Type": "application/octet-stream"}).status_code == 200
                response = client.post(path + "/commit", json={})
                assert response.status_code == 202, response.text
                result = wait(path, {"ready", "ready_with_gaps"})
                assert result["doc_id"]
                checks.append("frozen ingestion child: upload -> parse -> index -> ready")
                documents = client.get("/api/v1/documents").json()["documents"]
                assert any(d["doc_id"] == result["doc_id"] for d in documents)
                checks.append("library refresh after publication")
                chrome = Path("C:/Program Files/Google/Chrome/Application/chrome.exe")
                if chrome.is_file():
                    from playwright.sync_api import expect, sync_playwright

                    with sync_playwright() as playwright:
                        browser = playwright.chromium.launch(executable_path=str(chrome), headless=True)
                        page = browser.new_page(viewport={"width": 1440, "height": 1000})
                        errors = []
                        page.on("pageerror", lambda error: errors.append(str(error)))
                        page.route("**/*", lambda route: route.continue_()
                            if route.request.url.startswith(url + "/") else route.abort())
                        page.goto(url)
                        expect(page.locator("html")).to_have_attribute("data-app-state", "bound")
                        expect(page.locator("#document-count")).not_to_have_text("0")
                        page.locator("#theme-button").click()
                        expect(page.locator("html")).to_have_attribute("data-theme", "dark")
                        page.locator("#question-input").fill("SD5993 心跳如何启动？")
                        page.locator("#question-input").press("Enter")
                        expect(page.locator(".message-actions")).to_have_count(1, timeout=30000)
                        assert not errors, errors
                        page.screenshot(path=str(home / "frozen-ui.png"))
                        browser.close()
                    checks.append("frozen browser UI: JavaScript, theme, Enter submission and answer")
        finally:
            # This PID belongs to this test, never an existing user server.
            killed = subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False,
                creationflags=subprocess.CREATE_NO_WINDOW)
            if killed.returncode and process.poll() is None:
                process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)
    (home / "result.json").write_text(json.dumps({"passed": checks}, indent=2), encoding="utf-8")
    print(json.dumps({"passed": checks, "output": str(home)}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
