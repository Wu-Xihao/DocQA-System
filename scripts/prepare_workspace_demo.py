"""Create three synthetic DOCX sources and build an isolated, model-free workspace."""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402

from docqa.service.config import load_options  # noqa: E402
from docqa.service.pipeline import Service  # noqa: E402


def main():
    options = load_options(ROOT / "configs/service.demo.example.yaml")
    if options.retrieval.mode != "demo":
        raise RuntimeError("This sample launcher requires demo models")
    folder = ROOT / "runtime/workspace-demo-sources"
    folder.mkdir(parents=True, exist_ok=True)
    (ROOT / "runtime/workspace-demo-uploads").mkdir(parents=True, exist_ok=True)
    samples = [
        (
            "心跳协议",
            "SD5993 心跳启动流程：调用 FE_HEART_Start 后发送心跳报文，等待对端确认。心跳周期为30秒。",
        ),
        ("故障恢复", "心跳超时后执行连接恢复：先断开旧会话，再重新建立连接，最后重发注册请求。"),
        ("图文流程", "DOCQA_IMAGE 文档处理流程包括解析、检索、问答三个阶段。图示仅用于资料核对。"),
    ]
    for name, content in samples:
        path = folder / (name + ".docx")
        if path.exists():
            continue  # Never overwrite a sample the user has edited.
        document = Document()
        document.add_heading(name, 1)
        document.add_paragraph(content)
        if name == "图文流程":
            image = Image.new("RGB", (900, 260), "#f7f5ff")
            draw = ImageDraw.Draw(image)
            for i, label in enumerate(["1. Parse", "2. Retrieve", "3. Answer"]):
                x = 40 + i * 290
                draw.rounded_rectangle(
                    (x, 65, x + 230, 190), radius=20, fill="white", outline="#a090dc", width=3
                )
                draw.text((x + 60, 115), label, fill="#493d71", font_size=22)
                if i < 2:
                    draw.line((x + 232, 125, x + 280, 125), fill="#6657d9", width=4)
                    draw.polygon([(x + 280, 125), (x + 268, 119), (x + 268, 131)], fill="#6657d9")
            picture = ROOT / "runtime/workspace-demo-flow.png"
            image.save(picture)
            document.add_picture(str(picture))
        document.save(path)
    service = Service(ROOT / "runtime/workspace-demo-service", options=options)
    service.add_source(folder)
    service.run_once()
    status = service.index_status()
    print(
        {key: status[key] for key in ("documents", "chunks", "ingestion_status") if key in status}
    )


if __name__ == "__main__":
    main()
