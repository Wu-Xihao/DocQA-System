"""Generate synthetic, redistributable doc2md acceptance inputs.

Run from the project root:
    python scripts/generate_doc2md_samples.py --output runtime/manual-input

No ReportLab, model service, Tesseract or LibreOffice is needed to generate these
inputs. Existing named files are never overwritten without --overwrite.
"""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.shared import Inches as DocxInches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

SAMPLE_FILES = {
    "native_pdf": "text.pdf",
    "scanned_pdf": "scan.pdf",
    "docx": "sample.docx",
    "pptx": "sample.pptx",
    "vector_pptx": "vector.pptx",
    "broken_pdf": "broken.pdf",
    "unsupported": "unsupported.txt",
    "diagram": "diagram.png",
    "readme": "README_样例说明.md",
}


def _write_native_pdf(path: Path) -> None:
    """Write a real two-page PDF with only native Helvetica text."""
    streams = [
        b"BT /F1 22 Tf 45 740 Td (DocQA Native PDF Test) Tj "
        b"/F1 14 Tf 0 -45 Td (NATIVE-TEXT-001) Tj "
        b"0 -30 Td (Physical page 1: native text must remain searchable.) Tj ET",
        b"BT /F1 22 Tf 45 740 Td (Physical page 2) Tj "
        b"/F1 14 Tf 0 -45 Td (NATIVE-TEXT-002) Tj "
        b"0 -30 Td (No OCR or vision service is needed for this text.) Tj ET",
    ]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 2 /Kids [4 0 R 6 0 R] >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, stream in enumerate(streams):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + index * 2} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    data = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_offset = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    path.write_bytes(data)


def _font(size: int):
    for name in ("arial.ttf", "DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"):
        try:
            return ImageFont.truetype(name, size=size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def _draw_samples(stage: Path) -> None:
    with Image.new("RGB", (1200, 800), "white") as scan:
        draw = ImageDraw.Draw(scan)
        draw.text((70, 70), "SCANNED DOCUMENT", font=_font(70), fill="black")
        draw.text((70, 210), "OCR TEST", font=_font(64), fill="black")
        draw.text((70, 340), "SCAN-ID-001", font=_font(48), fill="black")
        draw.text(
            (70, 465), "Raster text only. No native PDF text layer.", font=_font(37), fill="black"
        )
        draw.text(
            (70, 580), "Expected: OCR text and page 1 source locator.", font=_font(35), fill="black"
        )
        scan.save(stage / "scan.pdf", "PDF", resolution=144)

    with Image.new("RGB", (1100, 440), "white") as diagram:
        draw = ImageDraw.Draw(diagram)
        draw.text((45, 30), "SYNTHETIC WORKFLOW - TEST IMAGE", font=_font(36), fill="#15385B")
        positions = [(50, 150, 320, 300), (420, 150, 690, 300), (790, 150, 1060, 300)]
        for bounds, label in zip(positions, ["START", "CHECK", "DONE"], strict=True):
            draw.rounded_rectangle(bounds, radius=14, outline="#1D658A", width=5, fill="#E8F3F8")
            x, y, right, bottom = bounds
            draw.text(
                ((x + right) / 2, (y + bottom) / 2),
                label,
                anchor="mm",
                font=_font(38),
                fill="black",
            )
        for start, end in ((320, 420), (690, 790)):
            draw.line((start, 225, end - 8, 225), fill="#1D658A", width=5)
            draw.polygon([(end - 20, 212), (end, 225), (end - 20, 238)], fill="#1D658A")
        draw.text(
            (45, 365),
            "Generated fixture; not evidence of model accuracy.",
            font=_font(28),
            fill="#555555",
        )
        diagram.save(stage / "diagram.png")


def _write_docx(stage: Path) -> None:
    document = Document()
    document.add_heading("doc2md 功能测试文档", level=1)
    document.add_paragraph("顺序标记：DOCX-BEFORE-TABLE。此段应位于表格之前。")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    values = [("模块", "职责"), ("doc2md", "多模态文档解析"), ("DocBase", "知识检索")]
    for row, values_row in zip(table.rows, values, strict=True):
        for cell, value in zip(row.cells, values_row, strict=True):
            cell.text = value
    document.add_paragraph("顺序标记：DOCX-BEFORE-IMAGE。此段应位于图片之前。")
    document.add_picture(str(stage / "diagram.png"), width=DocxInches(5.7))
    document.add_paragraph("图 1 人工生成的测试流程图")
    document.add_paragraph("顺序标记：DOCX-AFTER-IMAGE。此段应位于图片之后。")
    document.add_paragraph("此文件原生对象定位不承诺 Word 分页；物理页码须由实际转换渲染产生。")
    document.save(stage / "sample.docx")


def _presentation():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    return presentation


def _text(slide, text, x, y, width, height, *, size=22):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.text_frame.word_wrap = True
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string("15385B")
    return shape


def _write_pptx(stage: Path) -> None:
    presentation = _presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _text(slide, "doc2md 幻灯片测试", 0.5, 0.3, 12.3, 0.8, size=32)
    _text(
        slide,
        "原生文字标记：PPTX-TEXT-001；图像、表格应保留幻灯片位置。",
        0.55,
        1.25,
        12.1,
        0.65,
        size=22,
    )
    table = slide.shapes.add_table(3, 2, Inches(0.6), Inches(2.2), Inches(4.2), Inches(1.8)).table
    for row_index, row in enumerate([("对象", "期望"), ("文字", "原生提取"), ("图片", "可追溯")]):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(20)
    slide.shapes.add_picture(
        str(stage / "diagram.png"), Inches(5.5), Inches(2.15), width=Inches(7.2)
    )
    _text(slide, "此图仅为人工生成的样例，不代表 VL 识别效果已经验收。", 0.6, 6.3, 12, 0.7, size=20)
    presentation.save(stage / "sample.pptx")

    vector = _presentation()
    slide = vector.slides.add_slide(vector.slide_layouts[6])
    _text(slide, "原生矢量流程图测试（不含嵌入图片）", 0.5, 0.35, 12.3, 0.8, size=30)
    boxes = []
    for x, label in zip(
        (0.8, 4.8, 8.8),
        ("开始\nVECTOR-START", "校验\nVECTOR-CHECK", "结束\nVECTOR-END"),
        strict=True,
    ):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.6), Inches(3.2), Inches(1.6)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("E8F3F8")
        box.line.color.rgb = RGBColor.from_string("1D658A")
        box.text = label
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(24)
            paragraph.font.color.rgb = RGBColor.from_string("15385B")
        boxes.append(box)
    for left, right in zip(boxes[:-1], boxes[1:], strict=True):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            left.left + left.width,
            left.top + left.height // 2,
            right.left,
            right.top + right.height // 2,
        )
        connector.line.width = Pt(3)
        connector.line.color.rgb = RGBColor.from_string("1D658A")
        arrow = OxmlElement("a:tailEnd")
        arrow.set("type", "triangle")
        connector._element.spPr.get_or_add_ln().append(arrow)
    _text(
        slide,
        "预期：记录原生形状/连接线；完整视觉关系须经过幻灯片渲染及真实 VL 验证。",
        0.6,
        5.8,
        12.1,
        1.0,
        size=21,
    )
    vector.save(stage / "vector.pptx")


README = """# doc2md 人工生成测试样例

本目录由 `scripts/generate_doc2md_samples.py` 生成。所有内容均为合成测试数据，未使用用户照片或业务资料。

| 文件 | 样例内容 | 人工核对要点 |
| --- | --- | --- |
| `text.pdf` | 两页纯原生英文文本 | 搜索 `NATIVE-TEXT-001` / `NATIVE-TEXT-002`，物理页码分别为 1、2；无需 OCR 或 VL；不应生成无必要的整页视觉任务 |
| `scan.pdf` | 一页仅含位图文字的扫描型 PDF | 原生无文字层；应保留第 1 页图片；真实 OCR 后包含 `SCANNED DOCUMENT`、`OCR TEST`、`SCAN-ID-001` |
| `sample.docx` | 中文标题、段落、3×2 表格和人工流程图 | 保持 BEFORE-TABLE → 表格 → BEFORE-IMAGE → 图片 → AFTER-IMAGE 的前后顺序；原生定位不得伪造 Word 页码 |
| `sample.pptx` | 一页原生中文标题/文字、3×2 表格和嵌入图片 | `PPTX-TEXT-001`、表格与图均应有第 1 张幻灯片位置；图片应可回溯 |
| `vector.pptx` | 一页原生框、文字及两条带箭头连接线，无嵌入图片 | 原生文本标记 `VECTOR-START`、`VECTOR-CHECK`、`VECTOR-END`；检查矢量未被静默遗漏。未启用渲染时应明确提示视觉语义未完成 |
| `broken.pdf` | 刻意损坏的 PDF | 应报告 `PARSE_CORRUPT`，不得发布为成功结果；批处理仍应继续处理其他文件 |
| `unsupported.txt` | 非本阶段支持的输入格式 | 应报告 `UNSUPPORTED_FORMAT`，不得当 PDF 或 Office 处理 |
| `diagram.png` | START → CHECK → DONE 人工绘制位图 | 是 Office 文件中的图片来源；独立 PNG 是否接收以当前支持矩阵为准 |

## 样例的边界

- 这些文件用于可复现的功能验收，不证明真实 OCR 或 VL 模型的准确率。
- `scan.pdf` 使用英文大字，便于先用 `eng` OCR 语言包进行最小核验；中文扫描识别需要另备中文样例及 `chi_sim` 语言包。
- 生成过程不调用模型、不做 OCR、不使用 LibreOffice，也不读取原业务文件。
- `vector.pptx` 的完整图形关系与原文视觉校验，需要实际 Office 渲染工具及真实 VL 服务；只有文本成功不能等同于图形语义完成。
- 如需重复生成，应选择新目录，或明确传入 `--overwrite`。默认遇到同名文件会退出且保持原文件。

详细操作与通过条件见项目 `Document` 下的 doc2md 实现及手动测试文档。
"""


def generate_samples(output: Path, *, overwrite: bool = False) -> dict[str, Path]:
    """Generate all inputs, rejecting existing target names by default."""
    output = Path(output).resolve()
    paths = {key: output / name for key, name in SAMPLE_FILES.items()}
    if output.exists() and not output.is_dir():
        raise NotADirectoryError(f"样例输出路径不是目录：{output}")
    conflicts = [path.name for path in paths.values() if path.exists()]
    if conflicts and not overwrite:
        raise FileExistsError(
            "已存在同名样例；请选择新目录或显式使用 --overwrite：" + ", ".join(conflicts)
        )
    output.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".doc2md-samples-", dir=output) as temporary:
        stage = Path(temporary)
        _write_native_pdf(stage / "text.pdf")
        _draw_samples(stage)
        _write_docx(stage)
        _write_pptx(stage)
        (stage / "broken.pdf").write_bytes(
            b"%PDF-1.4\nINTENTIONALLY BROKEN DOCQA TEST PDF\n%%EOF\n"
        )
        (stage / "unsupported.txt").write_text(
            "This file intentionally tests unsupported input handling.\n", encoding="utf-8"
        )
        (stage / "README_样例说明.md").write_text(README, encoding="utf-8")
        for key, destination in paths.items():
            # Exclusive mode also rejects a conflicting file created after the
            # initial preflight check. Overwrite requires explicit caller intent.
            with (stage / SAMPLE_FILES[key]).open("rb") as source:
                with destination.open("wb" if overwrite else "xb") as target:
                    shutil.copyfileobj(source, target)
    return paths


def main() -> None:
    parser = argparse.ArgumentParser(description="生成不含业务资料的 doc2md 人工测试样例")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("runtime/manual-input"),
        help="输出目录（默认 runtime/manual-input）",
    )
    parser.add_argument("--overwrite", action="store_true", help="显式允许覆盖已存在的同名样例")
    args = parser.parse_args()
    try:
        samples = generate_samples(args.output, overwrite=args.overwrite)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    print(
        json.dumps({key: str(path) for key, path in samples.items()}, ensure_ascii=False, indent=2)
    )


if __name__ == "__main__":
    main()
