"""Native XLSX extraction: worksheets become sheet-range-located table blocks.

Cell values are the cached results Excel saved (data_only); formulas never run
here. Merged ranges repeat their anchor value so every row reads on its own.
Charts and pictures are recorded as coverage gaps, not guessed from cells.
"""

from __future__ import annotations

from zipfile import BadZipFile

from docqa.doc2md.common import ParseContext
from docqa.doc2md.parsers.sheets import SheetBuilder, cell_text
from docqa.domain.documents import ParsedContent
from docqa.domain.errors import DocQAError


def _rows(sheet) -> list[list[str]]:
    rows = [[cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    for merged in sheet.merged_cells.ranges:
        anchor = (
            rows[merged.min_row - 1][merged.min_col - 1]
            if merged.min_row <= len(rows) and merged.min_col <= len(rows[merged.min_row - 1])
            else ""
        )
        for r in range(merged.min_row, min(merged.max_row, len(rows)) + 1):
            row = rows[r - 1]
            for c in range(merged.min_col, min(merged.max_col, len(row)) + 1):
                row[c - 1] = anchor
    return rows


def parse_xlsx(ctx: ParseContext) -> ParsedContent:
    """Parse a validated XLSX snapshot; sheet order follows the workbook."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    result = ParsedContent(
        metadata={
            "format": "xlsx",
            "native_coverage": "worksheet cell values (cached formula results), merged ranges",
            "needs_office_render": False,
            "physical_page_mapping": "unavailable",
        }
    )
    builder = SheetBuilder(ctx, "xlsx", result)
    try:
        workbook = load_workbook(str(ctx.source), data_only=True, read_only=False)
    except (InvalidFileException, BadZipFile, KeyError, ValueError, TypeError, OSError) as exc:
        raise DocQAError(
            "DOCUMENT_PARSE_FAILED", "XLSX workbook is damaged, encrypted, or unreadable"
        ) from exc
    try:
        sheets = workbook.worksheets
        if len(sheets) > ctx.max_pages:
            builder.warn("SHEET_LIMIT", f"Only the first {ctx.max_pages} worksheets are parsed")
            sheets = sheets[: ctx.max_pages]
        for index, sheet in enumerate(sheets, start=1):
            hidden = sheet.sheet_state != "visible"
            if hidden:
                builder.warn("XLSX_HIDDEN_SHEET", "Hidden worksheets are parsed and marked")
            if getattr(sheet, "_images", None) or getattr(sheet, "_charts", None):
                builder.warn(
                    "XLSX_VISUAL_NOT_PARSED", "Worksheet charts/pictures are not parsed as evidence"
                )
            builder.add_sheet(
                sheet.title, _rows(sheet), sheet_index=index, metadata={"hidden": hidden}
            )
        if workbook.chartsheets:
            builder.warn(
                "XLSX_VISUAL_NOT_PARSED", "Worksheet charts/pictures are not parsed as evidence"
            )
    finally:
        workbook.close()
    return result
