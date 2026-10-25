"""Native PPTX shape-tree extraction with slide/object anchors and coverage warnings."""

from __future__ import annotations

from hashlib import sha256
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from docqa.doc2md.common import ParseContext
from docqa.domain.documents import (
    Block,
    FigureOccurrence,
    ParsedContent,
    ParseWarning,
    SourceLocator,
)
from docqa.domain.errors import DocQAError

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


class _PptxParser:
    def __init__(self, ctx: ParseContext, presentation):
        self.ctx = ctx
        self.presentation = presentation
        self.result = ParsedContent(
            page_count=len(presentation.slides),
            metadata={
                "format": "pptx",
                "needs_office_render": False,
                "reading_order": "native shape-tree (z-order), not inferred visual reading order",
                "slide_size_emu": [int(presentation.slide_width), int(presentation.slide_height)],
            },
        )
        self._assets = {}
        self._warnings: set[tuple[str, str]] = set()

    def warn(self, code: str, message: str, *, render: bool = False):
        key = code, message
        if key not in self._warnings:
            self.result.warnings.append(ParseWarning(code=code, message=message))
            self._warnings.add(key)
        if render:
            self.result.metadata["needs_office_render"] = True
            key = (
                "VISUAL_RENDER_REQUIRED",
                "PPTX native shapes require an Office rendering baseline",
            )
            if key not in self._warnings:
                self.result.warnings.append(ParseWarning(code=key[0], message=key[1]))
                self._warnings.add(key)

    def locator(self, slide, slide_number, shape, path, transform):
        sx, sy, tx, ty, exact = transform
        raw = [
            int(shape.left or 0),
            int(shape.top or 0),
            int(shape.width or 0),
            int(shape.height or 0),
        ]
        bbox = None
        clipped = False
        if exact and not shape.rotation:
            x0, y0 = raw[0] * sx + tx, raw[1] * sy + ty
            x1, y1 = (raw[0] + raw[2]) * sx + tx, (raw[1] + raw[3]) * sy + ty
            width = int(self.presentation.slide_width)
            height = int(self.presentation.slide_height)
            if width > 0 and height > 0:
                clipped = x0 < 0 or y0 < 0 or x1 > width or y1 > height
                bbox = tuple(
                    max(0.0, min(1.0, value))
                    for value in (x0 / width, y0 / height, x1 / width, y1 / height)
                )
        return SourceLocator(
            locator_type="pptx_shape",
            slide=slide_number,
            bbox=bbox,
            precision="exact_object",
            native_ref={
                "part": str(slide.part.partname),
                "slide_id": int(slide.slide_id),
                "shape_path": path,
                "shape_path_index_base": 1,
                "shape_id": int(shape.shape_id),
                "shape_name": shape.name,
                "raw_rect_emu": raw,
                "rotation_degrees": float(shape.rotation),
                "parent_to_slide_transform": [sx, sy, tx, ty],
                "bbox_transform_exact": exact and not bool(shape.rotation),
                "clipped_by_slide": clipped,
            },
            display_label=f"Slide {slide_number}, shape {'/'.join(map(str, path))}: {shape.name}",
            doc_id=self.ctx.doc_id,
            revision_id=self.ctx.revision_id,
        )

    def append_block(self, kind, locator, text="", **kwargs):
        self.result.blocks.append(
            Block(
                block_id=self.ctx.identity(
                    "block", "pptx", locator.native_ref, len(self.result.blocks)
                ),
                type=kind,
                text=text,
                order=len(self.result.blocks),
                locators=[locator],
                **kwargs,
            )
        )

    def picture(self, shape, locator):
        try:
            data = shape.image.blob
            asset = self.ctx.save_image(data, f"p{locator.slide:04d}_figure")
        except DocQAError as exc:
            if exc.code == "RESOURCE_LIMIT_EXCEEDED":
                raise
            self.warn(
                "EMBEDDED_IMAGE_UNSUPPORTED",
                "A PPTX image cannot be decoded as raster data",
                render=True,
            )
            return
        except (ValueError, KeyError, AttributeError):
            self.warn(
                "IMAGE_REFERENCE_BROKEN", "PPTX image relationship cannot be read", render=True
            )
            return
        self._assets.setdefault(asset.asset_id, asset)
        crop = {
            name: float(getattr(shape, f"crop_{name}", 0))
            for name in ("left", "top", "right", "bottom")
        }
        transform = next(shape.element.iter(_A + "xfrm"), None)
        flips = {
            name: transform.get(name, "0") if transform is not None else "0"
            for name in ("flipH", "flipV")
        }
        transformed = (
            bool(shape.rotation)
            or any(value != 0 for value in crop.values())
            or any(value not in {"0", "false"} for value in flips.values())
            or not locator.native_ref["bbox_transform_exact"]
            or locator.native_ref["clipped_by_slide"]
        )
        if transformed:
            self.warn(
                "IMAGE_DISPLAY_TRANSFORM",
                "PPTX image crop/rotation/group transform requires rendering",
                render=True,
            )
        figure_id = self.ctx.identity("figure", "pptx", locator.native_ref)
        doc_pr = next(shape.element.iter(_P + "cNvPr"), None)
        caption = (
            (doc_pr.get("descr") or doc_pr.get("title") or shape.name)
            if doc_pr is not None
            else shape.name
        )
        figure = FigureOccurrence(
            figure_id=figure_id,
            asset_id=asset.asset_id,
            original_asset_id=asset.derived_from,
            locator=locator,
            caption=caption,
            analysis_required=not transformed,
            semantic_state="disabled" if transformed else "pending",
            metadata={
                "raw_sha256": sha256(data).hexdigest(),
                "crop": crop,
                "flips": flips,
                "rotation_degrees": float(shape.rotation),
                "visible_region_policy": "office_render" if transformed else "native_image",
            },
        )
        self.result.figures.append(figure)
        self.append_block("figure", locator, figure_ids=[figure_id])

    def group_transform(self, group, parent):
        sx, sy, tx, ty, exact = parent
        transform = next(group.element.iter(_A + "xfrm"), None)
        if transform is None:
            return sx, sy, tx, ty, False
        off = transform.find(_A + "off")
        ext = transform.find(_A + "ext")
        child_off = transform.find(_A + "chOff")
        child_ext = transform.find(_A + "chExt")
        if any(node is None for node in (off, ext, child_off, child_ext)):
            return sx, sy, tx, ty, False
        cx, cy = int(child_ext.get("cx", "0")), int(child_ext.get("cy", "0"))
        if not cx or not cy:
            return sx, sy, tx, ty, False
        gsx, gsy = int(ext.get("cx", "0")) / cx, int(ext.get("cy", "0")) / cy
        gtx = int(off.get("x", "0")) - int(child_off.get("x", "0")) * gsx
        gty = int(off.get("y", "0")) - int(child_off.get("y", "0")) * gsy
        exact = (
            exact
            and not bool(group.rotation)
            and all(transform.get(name, "0") in {"0", "false"} for name in ("flipH", "flipV"))
        )
        return sx * gsx, sy * gsy, sx * gtx + tx, sy * gty + ty, exact

    def walk_shapes(self, slide, number, shapes, path=None, transform=(1.0, 1.0, 0.0, 0.0, True)):
        for index, shape in enumerate(shapes):
            shape_path = [*(path or []), index + 1]
            locator = self.locator(slide, number, shape, shape_path, transform)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                self.warn(
                    "GROUP_VISUAL_LAYOUT",
                    "Grouped PPTX content needs visual relationship validation",
                    render=True,
                )
                self.walk_shapes(
                    slide, number, shape.shapes, shape_path, self.group_transform(shape, transform)
                )
                continue
            if shape.has_text_frame:
                text = shape.text_frame.text.replace("\v", "\n").strip()
                if text:
                    title = slide.shapes.title
                    is_title = title is not None and shape.element is title.element
                    paragraphs = [
                        {"text": paragraph.text, "level": int(paragraph.level)}
                        for paragraph in shape.text_frame.paragraphs
                    ]
                    self.append_block(
                        "heading" if is_title else "paragraph",
                        locator,
                        text,
                        heading_level=1 if is_title else None,
                        metadata={"paragraphs": paragraphs, "shape_type": int(shape.shape_type)},
                    )
            if shape.has_table:
                table = shape.table
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                layout = []
                for row_index, row in enumerate(table.rows):
                    for col_index, cell in enumerate(row.cells):
                        layout.append(
                            {
                                "row": row_index,
                                "column": col_index,
                                "is_merge_origin": cell.is_merge_origin,
                                "is_spanned": cell.is_spanned,
                                "span_width": cell.span_width,
                                "span_height": cell.span_height,
                            }
                        )
                self.append_block(
                    "table",
                    locator,
                    "\n".join("\t".join(row) for row in rows),
                    table_rows=rows,
                    metadata={"cell_layout": layout},
                )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                self.picture(shape, locator)
            elif shape.shape_type not in {
                MSO_SHAPE_TYPE.TEXT_BOX,
                MSO_SHAPE_TYPE.PLACEHOLDER,
                MSO_SHAPE_TYPE.TABLE,
            }:
                self.warn(
                    "NATIVE_VISUAL_NOT_EXTRACTED",
                    "PPTX shape/connector/chart/SmartArt/media requires visual rendering",
                    render=True,
                )
            # Placeholder pictures and picture fills are not exposed as normal
            # Picture objects.  Do not silently treat them as extracted pixels.
            elif any(True for _ in shape.element.iter(_A + "blip")):
                self.warn(
                    "PICTURE_FILL_NOT_EXTRACTED",
                    "PPTX picture fill requires visual rendering",
                    render=True,
                )

    def run(self):
        if len(self.presentation.slides) > self.ctx.max_pages:
            raise DocQAError("RESOURCE_LIMIT_EXCEEDED", "PPTX exceeds configured slide limit")
        for number, slide in enumerate(self.presentation.slides, start=1):
            self.walk_shapes(slide, number, slide.shapes)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    self.warn(
                        "SPEAKER_NOTES_NOT_PARSED",
                        "PPTX speaker notes are outside native slide extraction",
                    )
            # Master/layout/background graphics can be visible without being in
            # slide.shapes.  Record this coverage gap for the rendering pass.
            own_background = slide.element.find(_P + "cSld/" + _P + "bg")
            inherited = list(slide.slide_layout.shapes) + list(
                slide.slide_layout.slide_master.shapes
            )
            if own_background is not None or any(
                shape.shape_type not in {MSO_SHAPE_TYPE.PLACEHOLDER} for shape in inherited
            ):
                self.warn(
                    "INHERITED_VISUAL_CONTENT",
                    "PPTX background/layout/master visuals require rendering",
                    render=True,
                )
        self.result.assets = list(self._assets.values())
        return self.result


def parse_pptx(ctx: ParseContext) -> ParsedContent:
    """Parse a validated PPTX snapshot; slide numbers are physical, 1-based."""
    try:
        return _PptxParser(ctx, Presentation(str(ctx.source))).run()
    except DocQAError:
        raise
    except (
        PackageNotFoundError,
        BadZipFile,
        XMLSyntaxError,
        KeyError,
        ValueError,
        TypeError,
        AttributeError,
        OSError,
    ) as exc:
        raise DocQAError(
            "DOCUMENT_PARSE_FAILED", "PPTX package is damaged, encrypted, or unreadable"
        ) from exc
