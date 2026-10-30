"""Portable Windows entry point; application data stays outside the bundle."""

from __future__ import annotations

import argparse
import json
import os
import runpy
import shutil
import socket
import sys
import threading
import webbrowser
from pathlib import Path
from urllib.request import urlopen

from docqa.processes import WORKER_MODULES

TEMPLATES = {
    "doc2md.example.yaml": "local.doc2md.yaml",
    "docbase.example.yaml": "local.database.yaml",
    "agent.example.yaml": "local.agent.yaml",
    "service.example.yaml": "local.service.yaml",
    "web.example.yaml": "local.web.yaml",
    ".env.example": ".env",
}


def bundle_root() -> Path:
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[2]))


def default_home() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parents[2]


def initialize(home: Path) -> None:
    """Copy only public templates; existing user configuration is never replaced."""
    configs = home / "configs"
    configs.mkdir(parents=True, exist_ok=True)
    templates = bundle_root() / "configs"
    for source, target in TEMPLATES.items():
        destination = configs / target
        if not destination.exists():
            # Exclusive creation also protects parallel first launches.
            try:
                with destination.open("xb") as output:
                    output.write((templates / source).read_bytes())
            except FileExistsError:
                pass
    for name in ("data/private/documents", "data/private/web-uploads", "runtime"):
        (home / name).mkdir(parents=True, exist_ok=True)


def prepare_demo(home: Path) -> Path:
    """Build a small synthetic workspace without network or private input."""
    from docx import Document

    from docqa.service.config import load_options
    from docqa.service.pipeline import Service

    for name in ("service.demo.example.yaml", "web.demo.example.yaml"):
        destination = home / "configs" / name
        if not destination.exists():
            shutil.copyfile(bundle_root() / "configs" / name, destination)
    sources = home / "runtime/workspace-demo-sources"
    sources.mkdir(parents=True, exist_ok=True)
    (home / "runtime/workspace-demo-uploads").mkdir(parents=True, exist_ok=True)
    path = sources / "DocQA-demo.docx"
    if not path.exists():
        document = Document()
        document.add_heading("DocQA 演示资料", 1)
        document.add_paragraph(
            "SD5993 心跳启动流程：调用 FE_HEART_Start 后发送心跳报文，等待对端确认。"
            "心跳周期为30秒。超时后先断开旧会话，再重新连接。"
        )
        document.save(path)
    service = Service(
        home / "runtime/workspace-demo-service",
        options=load_options(home / "configs/service.demo.example.yaml"),
    )
    service.add_source(sources)
    report = service.run_once()
    if report.get("failed"):
        raise RuntimeError("DEMO_PREPARATION_FAILED")
    return home / "configs/web.demo.example.yaml"


def open_when_ready(url: str, stopped: threading.Event) -> None:
    for _ in range(120):
        if stopped.is_set():
            return
        try:
            with urlopen(url + "/api/v1/health/live", timeout=0.5) as response:
                if json.load(response).get("status") == "live":
                    webbrowser.open(url)
                    return
        except (OSError, ValueError):
            pass
        stopped.wait(0.25)


def self_check() -> dict:
    """Check packaged assets, native libraries and profile sources, without APIs."""
    import tempfile
    from importlib.metadata import version

    import numpy as np
    import pypdfium2 as pdfium
    import sqlite_vec
    from PIL import Image

    from docqa.docbase.analyzer import Analyzer
    from docqa.docbase.pipeline import implementation_profile
    from docqa.docbase.vector_db import create
    from docqa.web.app import STATIC

    # These checks intentionally create only temporary, synthetic data.
    with tempfile.TemporaryDirectory(prefix="docqa-check-") as directory:
        sample = Path(directory) / "vectors.sqlite"
        create(sample, np.array([[1.0, 0.0]], dtype=np.float32))
        Image.new("RGB", (8, 8)).save(Path(directory) / "sample.png")
    assert (STATIC / "index.html").is_file()
    assert (STATIC / "js/main.js").is_file()
    assert (bundle_root() / "configs/web.example.yaml").is_file()
    assert pdfium.PdfDocument and sqlite_vec.load
    assert Analyzer().profile()["code_sha256"] and implementation_profile()
    for package in ("pydantic", "pdfplumber", "pypdfium2", "python-docx", "python-pptx", "Pillow"):
        assert version(package)
    return {"status": "ok", "frozen": bool(getattr(sys, "frozen", False)), "network": False}


def main(arguments: list[str] | None = None) -> int:
    arguments = list(sys.argv[1:] if arguments is None else arguments)
    # Workers must bypass dotenv, initialization, argument parsing and browser startup.
    if arguments[:1] == ["--internal-worker"]:
        if len(arguments) < 2 or arguments[1] not in WORKER_MODULES:
            return 2
        sys.argv = [arguments[1], *arguments[2:]]
        runpy.run_module(arguments[1], run_name="__main__", alter_sys=True)
        return 0
    parser = argparse.ArgumentParser(description="DocQA Windows launcher / DocQA 工作台")
    parser.add_argument("--demo", action="store_true", help="离线合成演示，不调用模型")
    parser.add_argument("--no-browser", action="store_true", help="不自动打开浏览器")
    parser.add_argument("--port", type=int, help="Web 端口，默认使用配置值")
    parser.add_argument("--home", type=Path, default=default_home(), help="配置与数据目录")
    parser.add_argument("--config", type=Path, help="指定 Web YAML")
    parser.add_argument("--check", action="store_true", help="检查打包依赖，不调用模型")
    parser.add_argument("--cli", nargs=argparse.REMAINDER, help="执行原有 docqa CLI 命令")
    args = parser.parse_args(arguments)
    if args.check:
        print(json.dumps(self_check(), ensure_ascii=True))
        return 0
    home = args.home.resolve()
    initialize(home)
    os.chdir(home)
    from dotenv import load_dotenv

    load_dotenv(home / "configs/.env", override=False)
    if args.cli is not None:
        from docqa.cli.main import app

        app(args=args.cli, prog_name="DocQA.exe --cli")
        return 0
    if args.demo and args.config:
        parser.error("--demo and --config cannot be combined")
    config = args.config or home / "configs/local.web.yaml"
    from docqa.web.config import load_options

    # Choose demo port separately, keeping real deployments on their configured port.
    if args.demo:
        config = prepare_demo(home)
    options = load_options(config)
    if args.port is not None:
        options.port = args.port
        options = type(options).model_validate(options.model_dump())
    url = f"http://127.0.0.1:{options.port}"
    import portalocker

    lock = portalocker.Lock(home / f"runtime/desktop-{options.port}.lock", timeout=0)
    try:
        lock.acquire()
    except portalocker.exceptions.LockException:
        print(f"DocQA 已在运行：{url}")
        if not args.no_browser:
            webbrowser.open(url)
        return 0
    stopped = threading.Event()
    try:
        # Reserve the listening socket before starting a browser readiness watcher.
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
            if os.name == "nt":
                listener.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
            try:
                listener.bind(("127.0.0.1", options.port))
            except OSError:
                print(f"端口 {options.port} 已占用，请关闭旧服务或使用 --port 8766。")
                return 2
            listener.listen(128)
            import uvicorn

            from docqa.web.app import create_app

            application = create_app(options)
            print(f"DocQA：{url}\n数据目录：{home}\n按 Ctrl+C 停止。", flush=True)
            if not args.demo:
                print("首次运行请填写 configs/.env 的 API 配置。", flush=True)
            if not args.no_browser:
                threading.Thread(target=open_when_ready, args=(url, stopped), daemon=True).start()
            server = uvicorn.Server(uvicorn.Config(
                application, host="127.0.0.1", port=options.port, workers=1,
                proxy_headers=False, server_header=False, log_level="warning",
            ))
            server.run(sockets=[listener])
    finally:
        stopped.set()
        lock.release()
    return 0


def entry() -> None:
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        raise SystemExit(130) from None
    except Exception as error:
        # Avoid dumping credentials or model payloads into a public console.
        print(f"DocQA 启动失败（{getattr(error, 'code', type(error).__name__)}）。")
        print("请检查配置、目录权限和端口。可从终端运行 DocQA.exe --check 检查依赖。")
        if getattr(sys, "frozen", False) and len(sys.argv) == 1 and sys.stdin.isatty():
            input("按 Enter 退出…")
        raise SystemExit(4) from None
