"""Offline scale and robustness stress test on a generated mixed-format corpus.

Measures parse throughput, index build time and retrieval latency (sequential,
concurrent and with filter conditions), then feeds malformed inputs to check
each is rejected with an explicit error instead of crashing. Demo providers only:
the numbers describe local processing, never real model latency or quality.
"""

from __future__ import annotations

import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook

from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.docbase.retrieval import query
from docqa.domain.errors import DocQAError
from docqa.quality import percentile, write_report

TOPICS = ["心跳", "证书", "超时", "重连", "加密", "日志", "电源", "温度"]


def _make_sources(folder: Path, documents: int, rows: int) -> list[Path]:
    paths = []
    for number in range(documents):
        topic = TOPICS[number % len(TOPICS)]
        kind = ("docx", "xlsx", "csv")[number % 3]
        path = folder / f"stress-{number:05d}.{kind}"
        if kind == "docx":
            document = Document()
            document.add_heading(f"{topic}规范 {number}", 1)
            for paragraph in range(5):
                document.add_paragraph(
                    f"STRESS_{number}_{paragraph}：{topic}处理步骤 {paragraph}，"
                    f"阈值 {number % 97} 毫秒，失败时记录 E{number:05d}。"
                )
            document.save(path)
        elif kind == "xlsx":
            book = Workbook()
            sheet = book.active
            sheet.append(["编号", "主题", "数值", "说明"])
            for row in range(rows):
                sheet.append([f"S{number}-{row}", topic, row * 1.5, f"{topic}记录 {row}"])
            book.save(path)
        else:
            lines = ["编号,主题,数值"] + [f"C{number}-{r},{topic},{r}" for r in range(rows)]
            path.write_text("\n".join(lines), encoding="utf-8")
        paths.append(path)
    return paths


def _malformed(folder: Path) -> dict[str, Path]:
    cases = {}

    def write(name, data: bytes):
        path = folder / name
        path.write_bytes(data)
        cases[name] = path

    write("empty.pdf", b"")
    write("not-really.pdf", b"plain text pretending to be a PDF")
    write("truncated.docx", b"PK\x03\x04truncated")
    write("binary.csv", bytes(range(256)) * 4)
    write("legacy-ole.xlsx", bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 600)
    bomb = folder / "bomb.xlsx"
    with ZipFile(bomb, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", "A" * 30_000_000)
    cases["bomb.xlsx"] = bomb
    traversal = folder / "traversal.docx"
    with ZipFile(traversal, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../evil.txt", "x")
    cases["traversal.docx"] = traversal
    return cases


def _latencies(function, inputs, workers):
    def timed(item):
        started = time.perf_counter()
        result = function(item)
        return (time.perf_counter() - started) * 1000, result

    started = time.perf_counter()
    if workers == 1:
        values = [timed(item) for item in inputs]
    else:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            values = list(executor.map(timed, inputs))
    wall = time.perf_counter() - started
    times = [ms for ms, _ in values]
    return {
        "requests": len(times),
        "workers": workers,
        "p50_ms": percentile(times, 0.5),
        "p95_ms": percentile(times, 0.95),
        "max_ms": round(max(times), 3),
        "throughput_qps": round(len(times) / wall, 2) if wall else None,
        "empty_results": sum(not result.hits for _, result in values),
    }


def stress(output, *, documents=60, rows=200, queries=40, workers=8):
    if not 3 <= documents <= 3000 or not 1 <= rows <= 20000:
        raise DocQAError("INVALID_REQUEST", "Use 3..3000 documents and 1..20000 rows")
    if not 1 <= queries <= 1000 or not 1 <= workers <= 32:
        raise DocQAError("INVALID_REQUEST", "Use 1..1000 queries and 1..32 workers")
    root = Path(output).resolve() / ("stress-" + uuid4().hex[:12])
    sources = root / "sources"
    sources.mkdir(parents=True)
    options = ParseOptions(allowed_roots=[sources], office_render="never")
    paths = _make_sources(sources, documents, rows)

    started = time.perf_counter()
    artifacts, parse_times = [], []
    for path in paths:
        begin = time.perf_counter()
        artifacts.append(Path(parse_document(path, root / "parsed", options=options).artifact_dir))
        parse_times.append(time.perf_counter() - begin)
    parse_seconds = time.perf_counter() - started

    retrieval = DocBaseOptions(mode="demo")
    started = time.perf_counter()
    built = build_index(artifacts, root / "index", options=retrieval)
    build_seconds = time.perf_counter() - started

    questions = [
        f"{TOPICS[i % len(TOPICS)]} 处理步骤 STRESS_{i % documents}_{i % 5}" for i in range(queries)
    ]

    def search(question):
        return query(root / "index", question, options=retrieval)

    def filtered(question):
        return query(
            root / "index",
            question,
            options=retrieval,
            filters={"doc_types": ["xlsx", "csv"], "content_types": ["table"]},
        )

    search(questions[0])  # Warm the verified-generation cache, as a running server does.
    latency = {
        "sequential": _latencies(search, questions, 1),
        "concurrent": _latencies(search, questions, workers),
        "filtered": _latencies(filtered, questions, 1),
    }

    robustness = []
    bad = root / "malformed"
    bad.mkdir()
    for name, path in _malformed(bad).items():
        try:
            parse_document(path, root / "parsed-bad", options=ParseOptions(allowed_roots=[bad]))
            robustness.append({"input": name, "rejected": False, "error_code": None})
        except DocQAError as error:
            robustness.append({"input": name, "rejected": True, "error_code": error.code})
        except Exception as error:  # Anything else is a crash: report it, do not hide it.
            robustness.append(
                {"input": name, "rejected": False, "error_code": "CRASH:" + type(error).__name__}
            )
    edge_queries = {
        "max_utf8_question": "检" * 5333,
        "special_characters": "\"'<script>%_*?[]{}\\ OR 1=1 --",
        "emoji_and_mixed": "心跳 heartbeat 🚀 ＳＤ５９９３",
    }
    for name, question in edge_queries.items():
        try:
            result = search(question)
            robustness.append({"input": name, "rejected": False, "status": result.status})
        except DocQAError as error:
            robustness.append({"input": name, "rejected": True, "error_code": error.code})
    for name, question in {"empty_question": "  ", "oversize_question": "检" * 6000}.items():
        try:
            search(question)
            robustness.append({"input": name, "rejected": False, "error_code": None})
        except DocQAError as error:
            robustness.append({"input": name, "rejected": True, "error_code": error.code})

    must_reject = {
        "empty.pdf",
        "not-really.pdf",
        "truncated.docx",
        "binary.csv",
        "legacy-ole.xlsx",
        "bomb.xlsx",
        "traversal.docx",
        "empty_question",
        "oversize_question",
    }
    crashes = [r for r in robustness if str(r.get("error_code") or "").startswith("CRASH")]
    unexpected = [
        r["input"] for r in robustness if (r["input"] in must_reject) != bool(r["rejected"])
    ]
    report = {
        "schema_version": "docqa.stress.v1",
        "synthetic": True,
        "external_requests": 0,
        "documents": documents,
        "rows_per_sheet": rows,
        "chunks": built.chunks,
        "parse": {
            "seconds": round(parse_seconds, 3),
            "docs_per_second": round(documents / parse_seconds, 2),
            "p95_seconds": percentile(parse_times, 0.95),
        },
        "build_seconds": round(build_seconds, 3),
        "retrieval": latency,
        "robustness": robustness,
        "passed": not crashes and not unexpected,
        "unexpected": unexpected,
        "caveat": "Demo providers: excludes real embedding/reranker network latency.",
    }
    return write_report(root / "stress.json", report) | {"report": str(root / "stress.json")}
