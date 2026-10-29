"""Source-bound study artifacts and fixed-layout exports from checked answers."""

from __future__ import annotations

import csv
import io
import json
import math
import textwrap
from pathlib import Path
from xml.etree import ElementTree as ET

from docqa.domain.errors import DocQAError

LABELS = {
    "table": "要点表格",
    "relation": "关系图",
    "mindmap": "思维导图",
    "report": "研究文档",
    "slides": "演示文稿",
}


def build_content(answer, kind, plan=None):
    if answer.status not in {"answered", "partial", "conflicting_evidence"} or not answer.claims:
        raise DocQAError("ARTIFACT_EVIDENCE_MISSING", "This answer has no supported claims")
    if not answer.synthetic and not answer.verified:
        raise DocQAError("ARTIFACT_EVIDENCE_MISSING", "Answer has not passed support checks")
    topics = {
        t["topic_id"]: t["text"]
        for event in answer.trace
        if event.get("stage") == "plan"
        for t in event.get("topics", [])
    }
    facts = [
        {
            "id": claim.claim_id,
            "text": claim.text,
            "kind": claim.kind,
            "topics": claim.topic_ids,
            "sources": [c.citation_id for c in answer.citations if claim.claim_id in c.claim_ids],
        }
        for claim in answer.claims
    ]
    citations = [
        {
            "id": c.citation_id,
            "name": c.source_name,
            "quote": c.quote,
            "location": "；".join(loc.display_label for loc in c.locators),
        }
        for c in answer.citations
    ]
    content = {
        "facts": facts,
        "citations": citations,
        "warnings": answer.warnings,
        "unanswered_topics": answer.unanswered_topics,
        "conflicts": [c.model_dump(mode="json") for c in answer.conflicts],
        "quality": answer.status,
        "synthetic": answer.synthetic,
        "topic_labels": topics,
    }
    if kind == "table":
        content["columns"] = ["已核验内容", "类型", "引用来源"]
        content["rows"] = [
            [
                {"value": f["text"], "claim_ids": [f["id"]], "sources": f["sources"]},
                {"value": f["kind"], "sources": f["sources"]},
                {"value": "、".join(f["sources"]), "sources": f["sources"]},
            ]
            for f in facts
        ]
    elif kind == "relation":
        if not answer.diagram_edges:
            raise DocQAError(
                "ARTIFACT_RELATIONS_MISSING", "Generate and verify explicit relationships first"
            )
        labels = list(
            dict.fromkeys(
                label for e in answer.diagram_edges for label in (e.source_label, e.target_label)
            )
        )
        content["nodes"] = [{"id": str(i), "label": label} for i, label in enumerate(labels)]
        content["edges"] = [
            {
                "source": str(labels.index(e.source_label)),
                "target": str(labels.index(e.target_label)),
                "direction": e.direction,
                "claim_ids": e.claim_ids,
                "sources": list(
                    dict.fromkeys(
                        c.citation_id
                        for c in answer.citations
                        if set(c.claim_ids) & set(e.claim_ids)
                    )
                ),
            }
            for e in answer.diagram_edges
        ]
    elif kind in {"mindmap", "slides"}:
        from docqa.domain.study import StudyPlan, StudySection

        if plan is None:
            groups = {}
            for fact in facts:
                label = topics.get(fact["topics"][0], "研究要点") if fact["topics"] else "研究要点"
                groups.setdefault(label, []).append(fact["id"])
            plan = StudyPlan(
                sections=[
                    StudySection(title=label[:80], claim_ids=ids) for label, ids in groups.items()
                ]
            )
        plan.check_claims(f["id"] for f in facts)
        by_id = {
            fact["id"]: dict(
                fact,
                label=plan.claim_labels.get(
                    fact["id"], fact["text"][:44] + ("…" if len(fact["text"]) > 44 else "")
                ),
            )
            for fact in facts
        }
        sections = [
            {
                "label": section.title,
                "kind": "editorial_group",
                "children": [by_id[i] for i in section.claim_ids],
            }
            for section in plan.sections
        ]
        content["outline"] = [section.title for section in plan.sections]
        content["tree"] = {"label": answer.question, "kind": "editorial_root", "children": sections}
        slides = []
        for section in sections:
            cards = []
            for fact in section["children"]:
                for i, part in enumerate(
                    textwrap.wrap(
                        fact["text"], width=300, replace_whitespace=False, drop_whitespace=False
                    )
                    or [fact["text"]]
                ):
                    cards.append(
                        {
                            "body": part,
                            "sources": fact["sources"],
                            "claim_ids": [fact["id"]],
                            "label": "推断"
                            if fact["kind"] == "inference"
                            else "来源要点" + (f" · 续{i + 1}" if i else ""),
                        }
                    )
            for i in range(0, len(cards), 2):
                selected = cards[i : i + 2]
                slides.append(
                    {
                        "title": section["label"] + (f" · {i // 2 + 1}" if len(cards) > 2 else ""),
                        "body": "\n\n".join(card["body"] for card in selected),
                        "cards": selected,
                        "layout": "columns" if len(selected) > 1 else "focus",
                        "sources": list(
                            dict.fromkeys(s for card in selected for s in card["sources"])
                        ),
                        "claim_ids": list(
                            dict.fromkeys(c for card in selected for c in card["claim_ids"])
                        ),
                    }
                )
        if len(slides) > 80:
            raise DocQAError("ARTIFACT_TOO_LARGE", "Slide deck exceeds the export page budget")
        content["slides"] = slides
    if kind == "mindmap":
        nodes, edges, positions, lines, heights, height = graph_layout(
            {"kind": kind, "content": content}
        )
        content["map_layout"] = {
            "font_size": 24,
            "line_height": 28,
            "nodes": nodes,
            "edges": edges,
            "positions": positions,
            "lines": lines,
            "heights": heights,
            "width": 2100,
            "height": height,
        }
    return content


def markdown(artifact):
    value = artifact["content"]
    lines = ["# " + artifact["title"], "", f"资料版本：{artifact['generation_id']}", ""]
    if value["synthetic"]:
        lines += ["> 演示成果：未经过真实模型语义验收。", ""]
    if value["quality"] != "answered":
        lines += ["> 本成果有来源覆盖缺口；仅整理已支持内容，未解析部分不能据此推断。", ""]
    groups = (
        value.get("tree", {}).get("children") if artifact["kind"] in {"mindmap", "slides"} else None
    )
    if groups:
        for group in groups:
            lines += ["## " + group["label"], ""]
            for fact in group["children"]:
                lines += [fact["text"] + " " + " ".join(f"[{s}]" for s in fact["sources"]), ""]
    else:
        for fact in value["facts"]:
            lines += [fact["text"] + " " + " ".join(f"[{s}]" for s in fact["sources"]), ""]
    if value["conflicts"]:
        lines += ["## 来源冲突", ""] + ["- " + c["description"] for c in value["conflicts"]] + [""]
    if value["unanswered_topics"]:
        lines += (
            ["## 待解决要点", ""]
            + ["- " + value["topic_labels"].get(t, t) for t in value["unanswered_topics"]]
            + [""]
        )
    lines += ["## 来源", ""]
    for citation in value["citations"]:
        lines += [
            f"[{citation['id']}] {citation['name']} · {citation['location']}",
            "",
            "> " + citation["quote"].replace("\n", "\n> "),
            "",
        ]
    return "\n".join(lines)


def graph_layout(artifact):
    value = artifact["content"]
    if artifact["kind"] == "relation":
        nodes = value["nodes"]
        edges = value["edges"]
    elif artifact["kind"] == "mindmap":
        colors = ["#7c5ce0", "#168e88", "#d78033", "#3e7cc4", "#b75a92", "#63973c"]
        nodes = [
            {
                "id": "root",
                "label": value["tree"]["label"],
                "level": 0,
                "color": "#6852cf",
                "sources": [],
            }
        ]
        edges, positions, lines, heights = [], {}, {}, {}
        totals = [100, 100]
        for i, group in enumerate(value["tree"]["children"]):
            side = 0 if totals[0] <= totals[1] else 1
            color, gid = colors[i % len(colors)], "g" + str(i)
            children = group["children"]
            ids = list(dict.fromkeys(s for fact in children for s in fact["sources"]))
            node = {"id": gid, "label": group["label"], "level": 1, "color": color, "sources": ids}
            nodes.append(node)
            edges.append(
                {
                    "source": "root",
                    "target": gid,
                    "direction": "undirected",
                    "color": color,
                    "sources": ids,
                }
            )
            top = totals[side]
            for j, fact in enumerate(children):
                fid = f"{gid}f{j}"
                label = fact.get("label", fact["text"][:44]) + " " + " ".join(fact["sources"])
                lines[fid] = textwrap.wrap(label, 12) or [""]
                heights[fid] = max(96, len(lines[fid]) * 28 + 36)
                nodes.append(
                    {
                        "id": fid,
                        "label": label,
                        "level": 2,
                        "color": color,
                        "sources": fact["sources"],
                        "claim_ids": [fact["id"]],
                        "detail": fact["text"],
                    }
                )
                positions[fid] = (230 if side == 0 else 1870, totals[side] + heights[fid] / 2)
                totals[side] += heights[fid] + 36
                edges.append(
                    {
                        "source": gid,
                        "target": fid,
                        "direction": "undirected",
                        "color": color,
                        "sources": fact["sources"],
                    }
                )
            lines[gid] = textwrap.wrap(group["label"], 12) or [""]
            heights[gid] = max(96, len(lines[gid]) * 28 + 36)
            totals[side] = max(totals[side], top + heights[gid] + 36)
            positions[gid] = (650 if side == 0 else 1450, (top + totals[side] - 36) / 2)
            totals[side] += 60
        lines["root"] = textwrap.wrap(nodes[0]["label"], 12) or [""]
        heights["root"] = max(108, len(lines["root"]) * 28 + 40)
        height = math.ceil(max(480, *totals, heights["root"] + 160))
        positions["root"] = (1050, height / 2)
        if len(nodes) > 200 or height * 2100 > 20_000_000:
            raise DocQAError("ARTIFACT_TOO_LARGE", "Mindmap exceeds safe canvas budget")
        return nodes, edges, positions, lines, heights, height
    else:
        raise DocQAError("EXPORT_UNSUPPORTED", "This artifact is not a graph")
    if len(nodes) > 200:
        raise DocQAError("ARTIFACT_TOO_LARGE", "Graph exceeds node budget")
    positions, lines, heights = {}, {}, {}
    top = 100
    for start in range(0, len(nodes), 3):
        row = nodes[start : start + 3]
        for node in row:
            lines[node["id"]] = textwrap.wrap(node["label"], 20) or [""]
            heights[node["id"]] = max(96, len(lines[node["id"]]) * 22 + 32)
        row_height = max(heights[n["id"]] for n in row)
        for i, node in enumerate(row):
            positions[node["id"]] = (190 + i * 360, top + row_height / 2)
        top += row_height + 55
    height = max(480, math.ceil(top))
    if height * 1100 > 20_000_000:
        raise DocQAError("ARTIFACT_TOO_LARGE", "Graph exceeds export pixel budget")
    return nodes, edges, positions, lines, heights, height


def graph_svg(artifact):
    nodes, edges, positions, lines, heights, height = graph_layout(artifact)
    width = 2100 if artifact["kind"] == "mindmap" else 1100
    svg = ET.Element(
        "svg",
        {
            "xmlns": "http://www.w3.org/2000/svg",
            "viewBox": f"0 0 {width} {height}",
            "width": str(width),
            "height": str(height),
        },
    )
    ET.SubElement(svg, "rect", {"width": "100%", "height": "100%", "fill": "#f8f7fe"})
    ET.SubElement(
        svg, "text", {"x": "30", "y": "32", "font-size": "16", "fill": "#302751"}
    ).text = (
        "DocQA · "
        + LABELS[artifact["kind"]]
        + " · "
        + ("演示成果" if artifact["content"]["synthetic"] else artifact["content"]["quality"])
        + " · "
        + artifact["generation_id"][:12]
    )
    for edge in edges:
        a, b = positions[edge["source"]], positions[edge["target"]]
        ET.SubElement(
            svg,
            "path",
            {
                "d": (
                    f"M{a[0]},{a[1]} C{(a[0] + b[0]) / 2},{a[1]} {(a[0] + b[0]) / 2},{b[1]} {b[0]},{b[1]}"
                    if artifact["kind"] == "mindmap"
                    else f"M{a[0]},{a[1]} L{b[0]},{b[1]}"
                ),
                "stroke": edge.get("color", "#a69ad8"),
                "stroke-width": "2",
                "fill": "none",
                "stroke-dasharray": "6 5" if edge.get("direction") == "unknown" else "none",
            },
        )
        if edge.get("direction") == "directed":
            angle = math.atan2(b[1] - a[1], b[0] - a[0])
            x, y = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
            points = f"{x},{y} {x - 12 * math.cos(angle - 0.4)},{y - 12 * math.sin(angle - 0.4)} {x - 12 * math.cos(angle + 0.4)},{y - 12 * math.sin(angle + 0.4)}"
            ET.SubElement(svg, "polygon", {"points": points, "fill": "#6657d9"})
    for node in nodes:
        x, y = positions[node["id"]]
        group = ET.SubElement(svg, "g")
        ET.SubElement(group, "title").text = node.get("detail", node["label"])
        ET.SubElement(
            group,
            "rect",
            {
                "x": str(x - 156),
                "y": str(y - heights[node["id"]] / 2),
                "width": "312",
                "height": str(heights[node["id"]]),
                "rx": "18",
                "fill": "#6852cf" if node.get("level") == 0 else "white",
                "stroke": node.get("color", "#d7d1f1"),
            },
        )
        text = ET.SubElement(
            group,
            "text",
            {
                "x": str(x),
                "y": str(y - (len(lines[node["id"]]) - 1) * 11),
                "text-anchor": "middle",
                "fill": "white" if node.get("level") == 0 else "#302751",
                "font-family": "Microsoft YaHei, sans-serif",
                "font-size": "14",
            },
        )
        for i, line in enumerate(lines[node["id"]]):
            ET.SubElement(text, "tspan", {"x": str(x), "dy": "0" if i == 0 else "22"}).text = line
    return ET.tostring(svg, encoding="utf-8", xml_declaration=True)


def export_bytes(artifact, format_name):
    if format_name == "json":
        return json.dumps(artifact, ensure_ascii=False, indent=2).encode(), "application/json"
    if format_name in {"md", "markdown"}:
        return markdown(artifact).encode(), "text/markdown; charset=utf-8"
    if format_name == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(["来源支持的内容", "类型", "引用来源", "质量状态", "模式"])
        for fact in artifact["content"]["facts"]:
            text = fact["text"]
            if text.lstrip().startswith(("=", "+", "-", "@")):
                text = "'" + text
            writer.writerow(
                [
                    text,
                    fact["kind"],
                    " ".join(fact["sources"]),
                    artifact["content"]["quality"],
                    "演示，未做真实语义校验" if artifact["content"]["synthetic"] else "真实模型",
                ]
            )
        return ("\ufeff" + output.getvalue()).encode(), "text/csv; charset=utf-8"
    if format_name == "svg":
        return graph_svg(artifact), "image/svg+xml"
    if format_name == "docx":
        from docx import Document
        from docx.shared import Pt

        document = Document()
        document.styles["Normal"].font.name = "Microsoft YaHei"
        document.styles["Normal"].font.size = Pt(11)
        for line in markdown(artifact).splitlines():
            if line.startswith("# "):
                document.add_heading(line[2:], 0)
            elif line.startswith("## "):
                document.add_heading(line[3:], 1)
            elif line:
                document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    if format_name == "pptx":
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt

        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        deck = artifact["content"].get("slides")
        if deck is None:
            raise DocQAError("EXPORT_UNSUPPORTED", "PPTX export requires a slide artifact")
        pages = [
            {
                "title": artifact["title"],
                "layout": "cover",
                "body": "依据已校验原文整理\n"
                + (
                    "来源覆盖有缺口"
                    if artifact["content"]["quality"] != "answered"
                    else "来源可追溯"
                ),
                "sources": [],
            },
            *(
                [
                    {
                        "title": "内容导航",
                        "body": "\n".join(
                            f"{i + 1:02d}   {label}"
                            for i, label in enumerate(artifact["content"].get("outline", []))
                        ),
                        "sources": [],
                        "layout": "agenda",
                    }
                ]
                if artifact["content"].get("outline")
                else []
            ),
            *deck,
        ]
        sources = artifact["content"]["citations"]
        for i in range(0, len(sources), 4):
            pages.append(
                {
                    "title": "引用来源",
                    "body": "\n\n".join(
                        f"[{c['id']}] {c['name']}\n{c['location']}" for c in sources[i : i + 4]
                    ),
                    "sources": [],
                }
            )
        if artifact["content"]["conflicts"]:
            for conflict in artifact["content"]["conflicts"]:
                for part in textwrap.wrap(conflict["description"], 420):
                    pages.append({"title": "来源冲突（须保留）", "body": part, "sources": []})
        if artifact["content"]["unanswered_topics"]:
            missing = "；".join(
                artifact["content"]["topic_labels"].get(t, t)
                for t in artifact["content"]["unanswered_topics"]
            )
            for part in textwrap.wrap(missing, 420):
                pages.append({"title": "尚待解决的要点", "body": part, "sources": []})
        mode_label = (
            "演示成果，未做真实语义校验"
            if artifact["content"]["synthetic"]
            else artifact["content"]["quality"]
        )
        for number, page in enumerate(pages, 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            cover = page.get("layout") == "cover"
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = (
                RGBColor(30, 35, 68) if cover else RGBColor(247, 248, 253)
            )
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = RGBColor(124, 92, 224)
            accent.line.fill.background()

            def textbox(left, top, width, height, text, size, *, pale=False, bold=False):
                frame = slide.shapes.add_textbox(
                    Inches(left), Inches(top), Inches(width), Inches(height)
                ).text_frame
                frame.word_wrap = True
                frame.text = text
                for paragraph in frame.paragraphs:
                    paragraph.font.name, paragraph.font.size = "Microsoft YaHei", Pt(size)
                    paragraph.font.bold = bold
                    paragraph.font.color.rgb = (
                        RGBColor(225, 221, 245)
                        if cover
                        else (RGBColor(110, 104, 142) if pale else RGBColor(44, 45, 70))
                    )
                return frame

            textbox(0.65, 0.4, 12, 0.35, "DOCQA RESEARCH / " + str(number).zfill(2), 11, pale=True)
            textbox(
                0.65, 1.3 if cover else 0.95, 12, 1.4, page["title"], 34 if cover else 28, bold=True
            )
            cards = page.get("cards", [])
            if cards:
                for i, card in enumerate(cards):
                    left, width = (0.65 + i * 6.15, 5.85) if len(cards) > 1 else (0.65, 12)
                    shape = slide.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE,
                        Inches(left),
                        Inches(2.3),
                        Inches(width),
                        Inches(4),
                    )
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    shape.line.color.rgb = RGBColor(219, 216, 239)
                    textbox(left + 0.2, 2.45, width - 0.4, 0.3, card["label"], 11, pale=True)
                    textbox(left + 0.2, 2.9, width - 0.4, 2.9, card["body"], 17)
                    textbox(
                        left + 0.2,
                        5.9,
                        width - 0.4,
                        0.35,
                        " ".join("[" + ref + "]" for ref in card["sources"]),
                        10,
                        pale=True,
                    )
            else:
                textbox(0.65, 3.1 if cover else 2.4, 12, 3.6, page["body"], 21 if cover else 18)
            textbox(
                0.65,
                6.8,
                12,
                0.4,
                " ".join(page["sources"]) + f"   ·   DocQA {number} · {mode_label}",
                10,
                pale=True,
            )
        output = io.BytesIO()
        presentation.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    if format_name == "png":
        from PIL import Image, ImageDraw, ImageFont

        nodes, edges, positions, lines, heights, height = graph_layout(artifact)
        width = 2100 if artifact["kind"] == "mindmap" else 1100
        image = Image.new("RGB", (width, height), "#f8f7fe")
        draw = ImageDraw.Draw(image)
        font_path = next(
            (
                p
                for p in [
                    Path("C:/Windows/Fonts/msyh.ttc"),
                    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                ]
                if p.is_file()
            ),
            None,
        )
        font = ImageFont.truetype(str(font_path), 14) if font_path else ImageFont.load_default()
        draw.text(
            (30, 20),
            "DocQA · "
            + LABELS[artifact["kind"]]
            + " · "
            + (
                "演示成果，未做真实语义校验"
                if artifact["content"]["synthetic"]
                else artifact["content"]["quality"]
            ),
            fill="#302751",
            font=font,
        )
        for edge in edges:
            a, b = positions[edge["source"]], positions[edge["target"]]
            if edge.get("direction") == "unknown":
                distance = math.dist(a, b)
                for start in range(0, math.ceil(distance), 12):
                    segment = [
                        tuple(
                            a[i] + (b[i] - a[i]) * min(distance, p) / max(distance, 1)
                            for i in range(2)
                        )
                        for p in (start, start + 6)
                    ]
                    draw.line(segment, fill="#a69ad8", width=2)
            elif artifact["kind"] == "mindmap":
                middle = (a[0] + b[0]) / 2
                points = []
                for step in range(41):
                    t, u = step / 40, 1 - step / 40
                    points.append(
                        (
                            u**3 * a[0]
                            + 3 * u * u * t * middle
                            + 3 * u * t * t * middle
                            + t**3 * b[0],
                            u**3 * a[1] + 3 * u * u * t * a[1] + 3 * u * t * t * b[1] + t**3 * b[1],
                        )
                    )
                draw.line(points, fill=edge.get("color", "#a69ad8"), width=3)
            else:
                draw.line([a, b], fill="#a69ad8", width=2)
            if edge.get("direction") == "directed":
                x, y = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
                angle = math.atan2(b[1] - a[1], b[0] - a[0])
                draw.polygon(
                    [
                        (x, y),
                        (x - 12 * math.cos(angle - 0.4), y - 12 * math.sin(angle - 0.4)),
                        (x - 12 * math.cos(angle + 0.4), y - 12 * math.sin(angle + 0.4)),
                    ],
                    fill="#6657d9",
                )
        for node in nodes:
            x, y = positions[node["id"]]
            node_height = heights[node["id"]]
            draw.rounded_rectangle(
                (x - 156, y - node_height / 2, x + 156, y + node_height / 2),
                radius=18,
                fill="#6852cf" if node.get("level") == 0 else "white",
                outline=node.get("color", "#d7d1f1"),
            )
            for i, line in enumerate(lines[node["id"]]):
                draw.text(
                    (x, y - (len(lines[node["id"]]) - 1) * 11 + i * 22),
                    line,
                    anchor="mm",
                    fill="white" if node.get("level") == 0 else "#302751",
                    font=font,
                )
        output = io.BytesIO()
        image.save(output, "PNG")
        return output.getvalue(), "image/png"
    raise DocQAError("EXPORT_UNSUPPORTED", "Unsupported export format")
