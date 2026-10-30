"""Offline end-to-end synthetic corpus baseline; never opens user configuration."""

import json
import time
from pathlib import Path
from uuid import uuid4

from docx import Document

from docqa.agent.config import AgentOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.domain.errors import DocQAError
from docqa.quality import evaluate, percentile, write_report
from docqa.web.runner import HeadlessRunner
from docqa.web.warm_runner import WarmRunner


def benchmark(output, *, documents=4, repeat=3):
    if not 2 <= documents <= 100 or not 2 <= repeat <= 10:
        raise DocQAError("INVALID_REQUEST", "Use 2..100 documents and 2..10 repetitions")
    root = Path(output).resolve() / uuid4().hex
    root.mkdir(parents=True)
    artifacts = []
    for number in range(documents):
        source = root / f"synthetic-{number}.docx"
        document = Document()
        document.add_heading(f"Synthetic protocol {number}", 1)
        document.add_paragraph(
            f"SYNTHETIC_{number} startup procedure: send HELLO, wait for ACK, then mark the connection ready."
        )
        document.save(source)
        artifacts.append(Path(parse_document(source, root / "parsed").artifact_dir))
    options = AgentOptions(mode="demo", retrieval=DocBaseOptions(mode="demo"))
    built = build_index(artifacts, root / "index", options=options.retrieval)
    cases = root / "cases.json"
    cases.write_text(
        json.dumps(
            [
                {
                    "id": "startup",
                    "question": "SYNTHETIC_0 startup HELLO ACK",
                    "expected_status": ["answered"],
                    "must_include": ["HELLO", "ACK"],
                }
            ]
        ),
        encoding="utf-8",
    )
    gate = evaluate(root / "index", cases, root / "quality", options)
    request = {"question": "SYNTHETIC_0 startup HELLO ACK", "generation_id": built.generation_id}
    rows = []
    for kind, cls in [("cold", HeadlessRunner), ("reused", WarmRunner)]:
        runner = cls(root / "index", root / "unused", options, root / "runs", timeout_seconds=45)
        try:
            for iteration in range(repeat):
                identity = uuid4().hex
                runner.agent_state = root / "agent" / identity
                started = time.perf_counter()
                result = runner.run(
                    identity, request, cancelled=lambda: False, progress=lambda _: None
                )
                timing = json.loads((root / "runs" / identity / "timing.json").read_text("utf-8"))
                rows.append(
                    {
                        "runner": kind,
                        "iteration": iteration + 1,
                        "seconds": round(time.perf_counter() - started, 6),
                        "status": result["status"],
                        "reused": timing.get("reused", False),
                        "process_id": timing.get("process_id"),
                    }
                )
        finally:
            if isinstance(runner, WarmRunner):
                runner.close()
    summary = {}
    for kind in ("cold", "reused"):
        values = [r["seconds"] for r in rows if r["runner"] == kind]
        summary[kind] = {
            "p50": percentile(values, 0.5),
            "p95": percentile(values, 0.95),
            "first_seconds": values[0],
        }
    return write_report(
        root / "baseline.json",
        {
            "schema_version": "docqa.benchmark.v1",
            "synthetic": True,
            "external_requests": 0,
            "semantic_quality_accepted": False,
            "documents": documents,
            "chunks": built.chunks,
            "quality_gate_passed": gate["passed"],
            "samples_per_runner": repeat,
            "latency": summary,
            "samples": rows,
            "caveat": "Local process/corpus baseline only; excludes real provider latency and large-corpus acceptance",
        },
    ) | {"report": str(root / "baseline.json"), "cases": str(cases), "index": str(root / "index")}
