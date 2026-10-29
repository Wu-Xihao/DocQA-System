"""Incomplete sources retain honest quality while admitting only native evidence."""

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from docqa.agent.config import AgentOptions
from docqa.agent.pipeline import answer_question
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document, verify_artifact
from docqa.docbase.admission import COVERAGE_WARNING
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.docbase.retrieval import query
from docqa.docbase.storage import load_generation
from docqa.domain.documents import VisionDescription
from docqa.domain.errors import DocQAError
from docqa.service.config import ServiceOptions
from docqa.service.pipeline import Service


class UncertainVision:
    def fingerprint(self):
        return {"adapter": "uncertain-test-double", "synthetic": True}

    def describe(self, *args, **kwargs):
        return VisionDescription(
            status="partial",
            observations=["UNSAFE_VL_SENTINEL"],
            uncertainties=["Unreadable label"],
        )


class SampleOCR:
    def fingerprint(self):
        return {"adapter": "ocr-test-double"}

    def recognize(self, *args, **kwargs):
        return "UNSAFE_OCR_SENTINEL"


def write_source(path, *, text=True):
    document = Document()
    if text:
        document.add_paragraph("NATIVE_ONLY_HEARTBEAT 每隔30秒发送一次心跳。")
    image = BytesIO()
    Image.new("RGB", (40, 30), "red").save(image, format="PNG")
    image.seek(0)
    document.add_picture(image)
    document.save(path)


@pytest.fixture
def parsed(tmp_path):
    path = tmp_path / "source.docx"
    write_source(path)
    return parse_document(
        path,
        tmp_path / "parsed",
        options=ParseOptions(ocr_mode="auto", office_render="never"),
        ocr_client=SampleOCR(),
        vision_client=UncertainVision(),
    )


def retrieval(**changes):
    return DocBaseOptions(
        mode="demo", recall_routes=["dense", "bm25"], incomplete_policy="native_only", **changes
    )


def test_default_policy_still_rejects_incomplete_sources(parsed, tmp_path):
    with pytest.raises(DocQAError) as error:
        build_index(
            [Path(parsed.artifact_dir)], tmp_path / "index", options=DocBaseOptions(mode="demo")
        )
    assert error.value.code == "SOURCE_NOT_READY"


def test_native_projection_preserves_source_integrity_and_excludes_derived_content(
    parsed, tmp_path
):
    artifact = Path(parsed.artifact_dir)
    manifest = (artifact / "manifest.json").read_bytes()
    built = build_index([artifact], tmp_path / "index", options=retrieval())
    assert built.documents == 1
    assert COVERAGE_WARNING in built.warnings
    with load_generation(Path(built.generation_dir)) as generation:
        assert generation.members[0]["source_quality"]["index_mode"] == "native_only"
        assert generation.manifest["partial_source_count"] == 1
        assert not generation.manifest["production_ready"]
        assert all(
            span.origin == "native" and span.figure_id is None
            for chunk in generation.chunks
            for span in chunk.source_spans
        )
        assert all(
            not chunk.figure_ids
            and "UNSAFE_VL_SENTINEL" not in chunk.index_text
            and "UNSAFE_OCR_SENTINEL" not in chunk.index_text
            for chunk in generation.chunks
        )
    assert verify_artifact(artifact)["status"] != "complete"
    assert (artifact / "manifest.json").read_bytes() == manifest
    found = query(tmp_path / "index", "NATIVE_ONLY_HEARTBEAT", options=retrieval())
    assert found.hits and COVERAGE_WARNING in found.warnings
    assert found.status == "ok"  # Source coverage differs from model-route degradation.


def test_partial_answer_has_native_citation_and_visible_coverage_notice(parsed, tmp_path):
    build_index([Path(parsed.artifact_dir)], tmp_path / "index", options=retrieval())
    answer = answer_question(
        tmp_path / "index",
        "NATIVE_ONLY_HEARTBEAT",
        output=tmp_path / "agent",
        options=AgentOptions(mode="demo", retrieval=retrieval()),
    )
    assert answer.status == "partial", answer.model_dump()
    assert answer.citations and all(c.origin == "native" for c in answer.citations)
    assert COVERAGE_WARNING in answer.warnings
    assert "来源仍有解析缺口" in answer.answer_markdown
    assert not answer.production_ready


def test_image_only_artifact_cannot_be_force_admitted(tmp_path):
    path = tmp_path / "image-only.docx"
    write_source(path, text=False)
    result = parse_document(
        path,
        tmp_path / "parsed",
        options=ParseOptions(ocr_mode="disabled", office_render="never"),
        vision_client=UncertainVision(),
    )
    with pytest.raises(DocQAError) as error:
        build_index([Path(result.artifact_dir)], tmp_path / "index", options=retrieval())
    assert error.value.code == "NO_INDEXABLE_CONTENT"


def test_tampered_artifact_still_fails_before_admission(parsed, tmp_path):
    path = Path(parsed.artifact_dir) / "document.md"
    path.write_bytes(path.read_bytes() + b"tampered")
    with pytest.raises(DocQAError) as error:
        build_index([Path(parsed.artifact_dir)], tmp_path / "index", options=retrieval())
    assert error.value.code == "ARTIFACT_CORRUPT"


def test_service_switches_policy_using_existing_artifact_without_reparse(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    write_source(source / "source.docx")
    service = Service(
        tmp_path / "state",
        options=ServiceOptions(
            allowed_source_roots=[source],
            settle_seconds=0,
            parse=ParseOptions(ocr_mode="disabled", office_render="never"),
            retrieval=DocBaseOptions(mode="demo", recall_routes=["dense", "bm25"]),
        ),
    )
    service.add_source(source)
    service.run_once()
    assert service.index_status()["documents"] == 0
    original = service.list_documents()[0]["latest_attempt_artifact"]
    service.options.retrieval.incomplete_policy = "native_only"
    service.parse_function = lambda *args, **kwargs: pytest.fail("must reuse saved artifacts")
    service.run_once()
    status = service.index_status()
    assert status["documents"] == status["partial_documents"] == 1
    assert status["blocked_documents"] == 0
    assert status["ingestion_status"] == "ready_with_gaps"
    assert service.list_documents()[0]["state"] == "review_required"
    assert service.list_documents()[0]["latest_attempt_artifact"] == original
    service.options.retrieval.incomplete_policy = "reject"
    service.run_once()
    assert service.index_status()["documents"] == 0


def test_query_cannot_silently_use_different_admission_policy(parsed, tmp_path):
    build_index([Path(parsed.artifact_dir)], tmp_path / "index", options=retrieval())
    strict = retrieval().model_copy(update={"incomplete_policy": "reject"})
    with pytest.raises(DocQAError) as error:
        query(tmp_path / "index", "NATIVE_ONLY_HEARTBEAT", options=strict)
    assert error.value.code == "PROFILE_MISMATCH"
