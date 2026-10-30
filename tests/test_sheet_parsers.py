"""Generated XLSX/CSV inputs verify sheet-range locations, merged cells and limits."""

from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from docqa.doc2md.common import ParseContext
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.parsers import PARSERS
from docqa.doc2md.parsers.delimited import parse_csv
from docqa.doc2md.parsers.sheets import SEGMENT_ROWS, column_letter, trim
from docqa.doc2md.parsers.xlsx import parse_xlsx
from docqa.doc2md.pipeline import parse_document
from docqa.doc2md.security import inspect_snapshot, validate_source
from docqa.domain.documents import SourceLocator
from docqa.domain.errors import DocQAError
from docqa.domain.formats import ALL_FORMATS, NATIVE_FORMATS, normalize_types


def context(source: Path) -> ParseContext:
    return ParseContext(
        source=source,
        work_dir=source.parent / "out",
        doc_id="test-document",
        revision_id="test-revision",
        source_name=source.name,
    )


def workbook(path: Path, rows=3) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "参数"
    sheet.append(["型号", "电压", "日期"])
    for index in range(rows):
        sheet.append([f"SD{5990 + index}", 3.3 + index, datetime(2026, 1, 1) + timedelta(days=index)])
    sheet.merge_cells("A10:B10")
    sheet["A10"] = "合并说明"
    hidden = book.create_sheet("隐藏")
    hidden.append(["内部", 1])
    hidden.sheet_state = "hidden"
    book.create_sheet("空表")
    book.save(path)
    return path


def tables(result):
    return [block for block in result.blocks if block.type == "table"]


def test_column_letters_and_trim():
    assert [column_letter(i) for i in (1, 26, 27, 52, 703)] == ["A", "Z", "AA", "AZ", "AAA"]
    rows, row0, column0 = trim([["", "", ""], ["", "x", ""], ["", "", "y"], ["", "", ""]])
    assert (rows, row0, column0) == ([["x", ""], ["", "y"]], 2, 2)
    assert trim([["", ""]]) == ([], 1, 1)


def test_xlsx_sheets_become_located_tables_with_merged_values(tmp_path):
    result = parse_xlsx(context(workbook(tmp_path / "参数.xlsx")))
    headings = [block.text for block in result.blocks if block.type == "heading"]
    assert headings == ["参数", "隐藏", "空表"]
    first = tables(result)[0]
    assert first.table_rows[0] == ["型号", "电压", "日期"]
    assert first.table_rows[1] == ["SD5990", "3.3", "2026-01-01"]
    assert ["合并说明", "合并说明", ""] in first.table_rows
    locator = first.locators[0]
    assert locator.locator_type == "sheet_range" and locator.page is None
    assert locator.native_ref["sheet"] == "参数"
    assert locator.native_ref["range"] == "A2:C10"
    assert locator.native_ref["header_range"] == "A1:C1"
    assert "参数 A2:C10" in locator.display_label
    assert {w.code for w in result.warnings} == {"XLSX_HIDDEN_SHEET"}
    hidden = [b for b in result.blocks if b.type == "heading" and b.text == "隐藏"][0]
    assert hidden.metadata["hidden"] is True


def test_long_sheets_split_with_repeated_header(tmp_path):
    result = parse_xlsx(context(workbook(tmp_path / "long.xlsx", rows=SEGMENT_ROWS * 2 + 5)))
    segments = [b for b in tables(result) if b.metadata["sheet"] == "参数"]
    assert len(segments) == 3
    assert all(segment.table_rows[0] == ["型号", "电压", "日期"] for segment in segments)
    assert [s.metadata["repeated_header"] for s in segments] == [False, True, True]
    assert segments[1].locators[0].native_ref["range"] == f"A{SEGMENT_ROWS + 2}:C{2 * SEGMENT_ROWS + 1}"
    data_rows = sum(len(segment.table_rows) - 1 for segment in segments)
    assert data_rows == SEGMENT_ROWS * 2 + 5  # The merged note overwrites row 10.


def test_csv_detects_encoding_and_delimiter(tmp_path):
    source = tmp_path / "清单.csv"
    source.write_bytes("名称;数量\n螺栓;12\n垫片;7\n".encode("gb18030"))
    result = parse_csv(context(source))
    assert result.metadata["encoding"] == "gb18030" and result.metadata["delimiter"] == ";"
    assert tables(result)[0].table_rows == [["名称", "数量"], ["螺栓", "12"], ["垫片", "7"]]
    assert tables(result)[0].locators[0].native_ref == {
        "sheet": "清单",
        "range": "A2:B3",
        "header_range": "A1:B1",
    }


def test_csv_rejects_binary_and_unknown_encoding(tmp_path):
    source = tmp_path / "bad.csv"
    source.write_bytes(b"a,b\n\x00\x01")
    with pytest.raises(DocQAError) as error:
        parse_csv(context(source))
    assert error.value.code == "FORMAT_MISMATCH"
    source.write_bytes(b"\xff\xfe\xfa")
    with pytest.raises(DocQAError):
        parse_csv(context(source))


def test_preflight_and_validation_accept_new_formats(tmp_path):
    options = ParseOptions(allowed_roots=[tmp_path])
    book = workbook(tmp_path / "a.xlsx")
    assert validate_source(book, options)[1:] == ("xlsx", False)
    assert inspect_snapshot(book, "xlsx", options) == {"external_resources": False}
    fake = tmp_path / "fake.csv"
    fake.write_bytes(book.read_bytes())
    with pytest.raises(DocQAError) as error:
        inspect_snapshot(fake, "csv", options)
    assert error.value.code == "FORMAT_MISMATCH"
    renamed = tmp_path / "doc.xlsx"
    from docx import Document

    Document().save(renamed)
    with pytest.raises(DocQAError) as error:
        inspect_snapshot(renamed, "xlsx", options)
    assert error.value.code == "FORMAT_MISMATCH"


def test_sheet_locator_requires_sheet():
    with pytest.raises(ValueError):
        SourceLocator(
            locator_type="sheet_range", precision="section_only", display_label="x", native_ref={}
        )


def test_registry_and_format_lists_stay_consistent():
    assert set(PARSERS) == set(NATIVE_FORMATS)
    assert set(NATIVE_FORMATS) < set(ALL_FORMATS)
    assert normalize_types([".XLSX", "csv", "csv"]) == ["xlsx", "csv"]
    with pytest.raises(ValueError):
        normalize_types(["exe"])


def test_xlsx_end_to_end_markdown_and_source_map(tmp_path):
    (tmp_path / "src").mkdir()
    source = workbook(tmp_path / "src" / "参数.xlsx")
    result = parse_document(
        source, tmp_path / "parsed", options=ParseOptions(allowed_roots=[tmp_path / "src"])
    )
    assert result.status in {"complete", "partial", "review_required"}
    markdown = Path(result.markdown_path).read_text(encoding="utf-8")
    assert "| 型号 | 电压 | 日期 |" in markdown
    assert "SD5990" in markdown and "工作表 参数 A2:C10" in markdown
