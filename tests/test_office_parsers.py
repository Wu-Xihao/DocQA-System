"""Small generated Office packages verify native extraction and honest coverage limits."""

from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as WordInches
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from docqa.doc2md.common import ParseContext
from docqa.doc2md.parsers.docx import parse_docx
from docqa.doc2md.parsers.pptx import parse_pptx
from docqa.domain.errors import DocQAError


@pytest.fixture
def picture(tmp_path):
    path = tmp_path / "small.png"
    image = Image.new("RGB", (80, 40), "white")
    image.putpixel((1, 1), (20, 60, 120))
    image.save(path)
    return path


def context(source: Path, **kwargs) -> ParseContext:
    return ParseContext(
        source=source,
        work_dir=source.parent / (source.stem + "-out"),
        doc_id="test-document",
        revision_id="test-revision",
        source_name=source.name,
        **kwargs,
    )


def codes(result):
    return {warning.code for warning in result.warnings}


def assert_order_and_locations(result):
    assert [block.order for block in result.blocks] == list(range(len(result.blocks)))
    assert len({block.block_id for block in result.blocks}) == len(result.blocks)
    assert all(block.locators for block in result.blocks)
    assert all(
        locator.doc_id == "test-document" and locator.revision_id == "test-revision"
        for block in result.blocks
        for locator in block.locators
    )


def test_docx_interleaves_headings_paragraphs_tables_links_and_images(tmp_path, picture):
    source = tmp_path / "body.docx"
    document = Document()
    document.add_heading("接口设计", level=2)
    paragraph = document.add_paragraph("before ")
    link = OxmlElement("w:hyperlink")
    relationship = document.part.relate_to(
        "https://example.invalid/never-fetch", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "hyperlink text"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    paragraph.add_run().add_picture(str(picture), width=WordInches(1))
    paragraph.add_run("after")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "timeout"
    table.cell(1, 1).text = "30 s"
    document.add_paragraph("last")
    document.save(source)

    result = parse_docx(context(source))

    assert [block.type for block in result.blocks] == [
        "heading",
        "paragraph",
        "figure",
        "paragraph",
        "table",
        "paragraph",
    ]
    assert result.blocks[0].heading_level == 2
    assert result.blocks[1].text == "before hyperlink text"
    assert (
        result.blocks[1].metadata["hyperlinks"][0]["target"]
        == "https://example.invalid/never-fetch"
    )
    assert result.blocks[3].text == "after"
    assert result.blocks[4].table_rows == [["name", "value"], ["timeout", "30 s"]]
    assert result.blocks[-1].text == "last"
    assert result.page_count is None
    assert all(locator.page is None for block in result.blocks for locator in block.locators)
    assert all(
        locator.native_ref["part"] == "/word/document.xml"
        for block in result.blocks
        for locator in block.locators
    )
    assert result.figures[0].locator.native_ref["object_path"].endswith("a:blip")
    assert len(result.assets) == 1
    assert result.assets[0].width == 80
    assert not result.metadata["needs_office_render"]
    assert_order_and_locations(result)


def test_docx_duplicate_image_has_distinct_occurrences_and_one_asset(tmp_path, picture):
    source = tmp_path / "repeated.docx"
    document = Document()
    document.add_picture(str(picture), width=WordInches(1))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(picture), width=WordInches(1))
    document.save(source)

    result = parse_docx(context(source))

    assert len(result.assets) == 1
    assert len(result.figures) == 2
    assert len({figure.figure_id for figure in result.figures}) == 2
    assert result.figures[0].asset_id == result.figures[1].asset_id
    assert "/w:tbl/" in result.figures[1].locator.native_ref["object_path"]
    assert [block.type for block in result.blocks] == ["figure", "table", "figure"]
    assert_order_and_locations(result)
    # A second parse of the same immutable revision yields the same identities.
    again = parse_docx(context(source))
    assert [figure.figure_id for figure in again.figures] == [
        figure.figure_id for figure in result.figures
    ]


def test_docx_tracks_native_merged_cell_structure(tmp_path):
    source = tmp_path / "merge.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
    table.cell(1, 2).text = "ordinary"
    document.save(source)

    result = parse_docx(context(source))

    assert result.blocks[0].table_rows[0] == ["merged", ""]
    assert result.blocks[0].metadata["cell_layout"][0]["grid_span"] == 2
    assert result.blocks[0].table_rows[1][2] == "ordinary"


def test_docx_revisions_are_not_silently_blended(tmp_path):
    source = tmp_path / "revision.docx"
    document = Document()
    paragraph = document.add_paragraph("existing ")
    for tag, text_tag, value in [
        ("w:del", "w:delText", "removed text"),
        ("w:ins", "w:t", "new text"),
    ]:
        revision = OxmlElement(tag)
        run = OxmlElement("w:r")
        text = OxmlElement(text_tag)
        text.text = value
        run.append(text)
        revision.append(run)
        paragraph._p.append(revision)
    document.save(source)

    result = parse_docx(context(source))

    assert "REVISIONS_UNRESOLVED" in codes(result)
    assert "VISUAL_RENDER_REQUIRED" in codes(result)
    assert result.metadata["needs_office_render"]
    assert "removed text" not in " ".join(block.text for block in result.blocks)
    assert "new text" in " ".join(block.text for block in result.blocks)


def test_docx_external_image_is_not_fetched(tmp_path):
    source = tmp_path / "external.docx"
    document = Document()
    drawing = OxmlElement("w:drawing")
    blip = OxmlElement("a:blip")
    blip.set(qn("r:link"), "rId999")
    drawing.append(blip)
    document.add_paragraph().add_run()._r.append(drawing)
    document.save(source)

    result = parse_docx(context(source))

    assert "EXTERNAL_IMAGE_NOT_FETCHED" in codes(result)
    assert result.figures == []
    assert result.metadata["needs_office_render"]


def test_docx_cropped_image_cannot_be_interpreted_as_fully_visible(tmp_path, picture):
    source = tmp_path / "crop.docx"
    document = Document()
    document.add_picture(str(picture), width=WordInches(1))
    blip_fill = document.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill
    src_rect = OxmlElement("a:srcRect")
    src_rect.set("l", "50000")
    blip_fill.append(src_rect)
    document.save(source)

    result = parse_docx(context(source))

    assert "IMAGE_DISPLAY_TRANSFORM" in codes(result)
    assert not result.figures[0].analysis_required
    assert result.figures[0].semantic_state == "disabled"
    assert result.figures[0].metadata["visible_region_policy"] == "office_render"


def test_docx_headers_and_text_boxes_have_explicit_coverage_warning(tmp_path):
    source = tmp_path / "coverage.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "header fact"
    document.add_paragraph("body fact")
    text_box = OxmlElement("w:txbxContent")
    inner = OxmlElement("w:p")
    text_box.append(inner)
    document.paragraphs[0]._p.append(text_box)
    document.save(source)

    result = parse_docx(context(source))

    assert "NON_BODY_CONTENT_NOT_PARSED" in codes(result)
    assert "COMPLEX_OBJECT_NOT_PARSED" in codes(result)
    assert result.metadata["needs_office_render"]
    assert result.blocks[0].text == "body fact"


def test_pptx_native_text_tables_and_repeated_images(tmp_path, picture):
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "System architecture"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text = "first\nsecond"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "workers"
    table.cell(1, 1).text = "2"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(3), width=Inches(2))
    other = presentation.slides.add_slide(presentation.slide_layouts[6])
    other.shapes.add_picture(str(picture), Inches(2), Inches(2), width=Inches(2))
    presentation.save(source)

    result = parse_pptx(context(source))

    assert result.page_count == 2
    assert [block.type for block in result.blocks] == [
        "heading",
        "paragraph",
        "table",
        "figure",
        "figure",
    ]
    assert result.blocks[0].heading_level == 1
    assert result.blocks[1].text == "first\nsecond"
    assert result.blocks[2].table_rows == [["name", "value"], ["workers", "2"]]
    assert len(result.figures) == 2
    assert len(result.assets) == 1
    assert len({figure.figure_id for figure in result.figures}) == 2
    assert [figure.locator.slide for figure in result.figures] == [1, 2]
    assert result.figures[0].locator.native_ref["part"] == "/ppt/slides/slide1.xml"
    assert result.figures[0].locator.native_ref["shape_path"] == [4]
    assert result.figures[0].locator.bbox == pytest.approx((0.1, 0.4, 0.3, 0.53333333333))
    assert_order_and_locations(result)


def test_pptx_nested_groups_recurse_and_map_group_coordinates(tmp_path, picture):
    source = tmp_path / "group.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text = "inside group"
    inner.shapes.add_picture(str(picture), Inches(2), Inches(2), width=Inches(1))
    # python-pptx updates a group's extents when adding children.  Explicitly
    # rescale the outer group afterwards to exercise child-space transforms.
    outer.left = Inches(3)
    outer.top = Inches(1)
    outer.width = Inches(4)
    outer.height = Inches(3)
    presentation.save(source)

    result = parse_pptx(context(source))

    assert any(block.text == "inside group" for block in result.blocks)
    figure = result.figures[0]
    assert figure.locator.native_ref["shape_path"] == [1, 1, 2]
    assert figure.locator.bbox is not None
    assert figure.locator.bbox == pytest.approx((0.5, 0.4, 0.7, 0.533333333333))
    assert figure.locator.native_ref["bbox_transform_exact"]
    assert "GROUP_VISUAL_LAYOUT" in codes(result)
    assert "VISUAL_RENDER_REQUIRED" in codes(result)
    assert result.metadata["needs_office_render"]
    assert_order_and_locations(result)


@pytest.mark.parametrize("display_change", ["crop", "rotate", "group_rotate"])
def test_pptx_transformed_original_image_is_not_interpreted(tmp_path, picture, display_change):
    source = tmp_path / f"{display_change}.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if display_change == "group_rotate":
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(str(picture), Inches(1), Inches(1), width=Inches(1))
        group.rotation = 45
    else:
        shape = slide.shapes.add_picture(str(picture), Inches(1), Inches(1), width=Inches(1))
        if display_change == "crop":
            shape.crop_left = 0.5
        else:
            shape.rotation = 45
    presentation.save(source)

    result = parse_pptx(context(source))

    assert "IMAGE_DISPLAY_TRANSFORM" in codes(result)
    assert not result.figures[0].analysis_required
    assert result.figures[0].semantic_state == "disabled"
    if display_change != "crop":
        assert result.figures[0].locator.bbox is None


def test_pptx_pure_shape_flowchart_requires_visual_baseline(tmp_path):
    source = tmp_path / "flowchart.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    first.text = "input"
    second = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(1), Inches(2), Inches(1)
    )
    second.text = "output"
    slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(1.5), Inches(5), Inches(1.5)
    )
    presentation.save(source)

    result = parse_pptx(context(source))

    assert [block.text for block in result.blocks] == ["input", "output"]
    assert result.figures == []  # It is not an embedded image.
    assert "NATIVE_VISUAL_NOT_EXTRACTED" in codes(result)
    assert "VISUAL_RENDER_REQUIRED" in codes(result)
    assert result.metadata["needs_office_render"]


def test_pptx_chart_and_speaker_notes_have_explicit_coverage_warning(tmp_path):
    source = tmp_path / "chart.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("count", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data
    )
    slide.notes_slide.notes_text_frame.text = "speaker only fact"
    presentation.save(source)

    result = parse_pptx(context(source))

    assert "VISUAL_RENDER_REQUIRED" in codes(result)
    assert "SPEAKER_NOTES_NOT_PARSED" in codes(result)
    assert "speaker only fact" not in " ".join(block.text for block in result.blocks)


def test_pptx_enforces_slide_limit(tmp_path):
    source = tmp_path / "limit.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)

    with pytest.raises(DocQAError, match="slide limit") as error:
        parse_pptx(context(source, max_pages=1))
    assert error.value.code == "RESOURCE_LIMIT_EXCEEDED"


@pytest.mark.parametrize(("extension", "parser"), [("docx", parse_docx), ("pptx", parse_pptx)])
def test_office_damage_is_a_typed_error(tmp_path, extension, parser):
    source = tmp_path / ("broken." + extension)
    source.write_bytes(b"not an Office package")

    with pytest.raises(DocQAError) as error:
        parser(context(source))
    assert error.value.code == "DOCUMENT_PARSE_FAILED"
