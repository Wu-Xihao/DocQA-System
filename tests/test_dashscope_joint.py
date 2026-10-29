"""Native HTTP contract + real parsed assets/SQLite-vec; all model responses are isolated fakes."""

from __future__ import annotations

import json
import math
from pathlib import Path
from types import SimpleNamespace

import httpx
import numpy as np
import pytest
from docx import Document
from pydantic import ValidationError

from docqa.agent.config import AgentOptions
from docqa.agent.retrieval import CandidateLedger, Subquery
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions, ServiceOptions
from docqa.docbase.dashscope import DashScopeEmbeddingClient, shared_encoder
from docqa.docbase.pipeline import build_index
from docqa.docbase.providers import make_providers
from docqa.docbase.retrieval import query
from docqa.docbase.storage import load_generation, read_current
from docqa.domain.errors import DocQAError
from docqa.service.config import ServiceOptions as BackgroundOptions
from docqa.web.config import WebOptions
from docqa.web.runner import HeadlessRunner
from docqa.web.workspace import Workspace


def joint_config(**changes):
    return ServiceOptions.model_validate(
        {
            "protocol": "dashscope.text-embedding",
            "endpoint": "https://models.invalid/joint",
            "model": "text-embedding-v4",
            "revision": "test-unpinned",
            "tokenizer_revision": "test-unreported",
            "dimension": 64,
            "vocab_size": 10000,
            "batch_size": 25,
            "retries": 0,
            **changes,
        }
    )


def options(backend="numpy", **changes):
    config = joint_config()
    return DocBaseOptions(
        dense=config,
        sparse=config.model_copy(),
        vector_backend=backend,
        reranker=ServiceOptions(
            protocol="gitee.rerank",
            endpoint="https://models.invalid/rerank",
            model="test-reranker",
            revision="1",
            tokenizer_revision="1",
            retries=0,
        ),
        top_k=3,
        **changes,
    )


def embedding_row(index=0, **changes):
    return {
        "text_index": index,
        "embedding": [3.0, 4.0] + [0.0] * 62,
        "sparse_embedding": [{"index": 2, "value": 1.5}, {"index": 10, "value": 0.5}],
        **changes,
    }


@pytest.mark.parametrize("index_key", ["text_index", "index", "both"])
def test_native_batch_contract_order_caps_and_credentials(index_key, monkeypatch):
    calls = []
    monkeypatch.setenv("TEST_DASHSCOPE_KEY", "isolated-key")
    monkeypatch.setenv("TEST_DASHSCOPE_ENDPOINT", "https://models.invalid/joint")

    def respond(request):
        assert request.headers["authorization"] == "Bearer isolated-key"
        payload = json.loads(request.content)
        calls.append(payload)
        assert set(payload) == {"model", "input", "parameters"}
        assert payload["parameters"] == {
            "dimension": 64,
            "output_type": "dense&sparse",
            "text_type": "query",
        }
        assert len(payload["input"]["texts"]) <= 10
        rows = []
        for index, text in enumerate(payload["input"]["texts"]):
            row = embedding_row(index, embedding=[float(text.split(":")[-1]) + 1, 4.0] + [0.0] * 62)
            if index_key in {"index", "both"}:
                row["index"] = index
            if index_key == "index":
                del row["text_index"]
            rows.append(row)
        return httpx.Response(200, json={"output": {"embeddings": list(reversed(rows))}})

    client = DashScopeEmbeddingClient(
        joint_config(
            api_key_env="TEST_DASHSCOPE_KEY",
            endpoint="",
            endpoint_env="TEST_DASHSCOPE_ENDPOINT",
            query_prefix="q:",
        ),
        transport=httpx.MockTransport(respond),
    )
    values = client.encode([str(i) for i in range(25)], "query")
    assert [row["dense"][0] for row in values] == list(range(1, 26))
    assert all(row["sparse"] == {2: 1.5, 10: 0.5} for row in values)
    assert len(calls) == 3
    assert client.encode([], "query") == [] and len(calls) == 3


BAD_ROWS = [
    {"embedding": [0.0] * 64},
    {"embedding": [1.0]},
    {"embedding": [math.nan] + [1.0] * 63},
    {"embedding": [True] + [1.0] * 63},
    {"text_index": True},
    {"text_index": -1},
    {"text_index": 1},
    {"index": 1},
    {"index": False},
    {"truncated": True},
    {"finish_reason": "length"},
    {"sparse_embedding": None},
    {"sparse_embedding": []},
    {"sparse_embedding": [{"index": 2, "value": 1}, {"index": 2, "value": 2}]},
    *[{"sparse_embedding": [{"index": token, "value": 1}]} for token in [-1, 10000, True, "2"]],
    *[{"sparse_embedding": [{"index": 2, "value": weight}]} for weight in [-1, 0, True, math.inf]],
]


@pytest.mark.parametrize("changes", BAD_ROWS)
def test_joint_response_rejects_incomplete_nonfinite_or_ambiguous_rows(changes):
    response = {"output": {"embeddings": [embedding_row(**changes)]}}
    client = DashScopeEmbeddingClient(
        joint_config(),
        transport=httpx.MockTransport(
            lambda _: httpx.Response(200, content=json.dumps(response).encode()),
        ),
    )
    with pytest.raises(DocQAError) as caught:
        client.encode(["synthetic"])
    assert caught.value.code == "MODEL_RESPONSE_INVALID"


@pytest.mark.parametrize(
    "response",
    [
        [],
        {},
        {"embeddings": [embedding_row()]},
        {"output": {"embeddings": []}},
        {"output": {"embeddings": [embedding_row(), embedding_row()]}},
        {"truncated": True, "output": {"embeddings": [embedding_row()]}},
        {"output": {"finish_reason": "length", "embeddings": [embedding_row()]}},
    ],
)
def test_native_envelope_cannot_hide_missing_or_duplicate_results(response):
    client = DashScopeEmbeddingClient(
        joint_config(),
        transport=httpx.MockTransport(
            lambda _: httpx.Response(200, json=response),
        ),
    )
    with pytest.raises(DocQAError, match="contract"):
        client.encode(["synthetic"])


@pytest.mark.parametrize(
    "change",
    [
        {"sparse": joint_config(model="other-model")},
        {"recall_routes": ["dense", "bm25"]},
        {"sparse": ServiceOptions()},
        {"dense": joint_config(dimension=None), "sparse": joint_config(dimension=None)},
    ],
)
def test_joint_configuration_cannot_mix_encoders_or_disable_sparse(change):
    base = options().model_dump()
    with pytest.raises(ValidationError):
        DocBaseOptions.model_validate({**base, **change})


def test_missing_endpoint_or_key_fails_without_network(monkeypatch):
    monkeypatch.delenv("TEST_MISSING_ENDPOINT", raising=False)
    with pytest.raises(DocQAError) as caught:
        DashScopeEmbeddingClient(joint_config(endpoint_env="TEST_MISSING_ENDPOINT"))
    assert caught.value.code == "MODEL_NOT_CONFIGURED"
    monkeypatch.delenv("TEST_MISSING_KEY", raising=False)
    client = DashScopeEmbeddingClient(
        joint_config(api_key_env="TEST_MISSING_KEY"),
        transport=httpx.MockTransport(lambda _: pytest.fail("must not send")),
    )
    with pytest.raises(DocQAError) as caught:
        client.encode(["synthetic"])
    assert caught.value.code == "MODEL_CREDENTIAL_MISSING"


@pytest.fixture
def corpus(tmp_path):
    source, parsed = tmp_path / "source", tmp_path / "parsed"
    source.mkdir()
    artifacts, ids = [], {}
    for marker, category, text in [
        ("ALPHA", "operations", "heartbeat starts after registration"),
        ("BETA", "operations", "heartbeat recovery after timeout"),
        ("GAMMA", "architecture", "layered modules implement a document system"),
    ]:
        doc = Document()
        doc.add_paragraph(f"{marker} {text}")
        path = source / f"{marker}.docx"
        doc.save(path)
        result = parse_document(
            path,
            parsed,
            options=ParseOptions(
                allowed_roots=[source],
                category=category,
                office_render="never",
                ocr_mode="disabled",
            ),
        )
        assert result.status == "complete"
        artifacts.append(Path(result.artifact_dir))
        ids[marker] = result.doc_id
    return source, parsed, artifacts, ids


class NativeServer:
    def __init__(self):
        self.calls, self.fail = [], False

    def __call__(self, request):
        payload = json.loads(request.content)
        self.calls.append((request.url.path, payload))
        if request.url.path == "/rerank":
            return httpx.Response(
                200,
                json={
                    "results": [
                        {
                            "index": i,
                            "relevance_score": 3.0
                            if "BETA" in text
                            else 2.0
                            if "ALPHA" in text
                            else 1.0,
                        }
                        for i, text in reversed(list(enumerate(payload["documents"])))
                    ]
                },
            )
        if self.fail:
            return httpx.Response(
                200, json={"output": {"embeddings": [embedding_row(sparse_embedding=[])]}}
            )
        rows = []
        for i, text in enumerate(payload["input"]["texts"]):
            a, b = (0.0, 1.0) if "GAMMA" in text else (1.0, 0.0)
            weight = 2.0 if "BETA" in text else 1.0
            rows.append(
                embedding_row(
                    i,
                    embedding=[a, b] + [0.0] * 62,
                    sparse_embedding=[
                        {"index": 3 if "GAMMA" in text else 2, "value": weight},
                        {"index": 10, "value": 0.2},
                    ],
                )
            )
        return httpx.Response(200, json={"output": {"embeddings": list(reversed(rows))}})

    def embeddings(self, text_type):
        return [
            payload
            for path, payload in self.calls
            if path == "/joint" and payload["parameters"]["text_type"] == text_type
        ]


@pytest.mark.parametrize("backend", ["numpy", "sqlite_vec"])
def test_build_three_route_query_agent_filters_and_incremental_cache(corpus, tmp_path, backend):
    source, parsed, artifacts, ids = corpus
    config, server = options(backend), NativeServer()
    providers = make_providers(config, transport=httpx.MockTransport(server))
    assert shared_encoder(*providers[:2]) is not None
    index = tmp_path / "knowledge"
    first = build_index(artifacts, index, options=config, providers=providers)
    assert first.documents == first.chunks == 3
    assert len(server.embeddings("document")) == 1
    assert len(list((index / "cache/dense_sparse").glob("*.json"))) == 3
    with load_generation(Path(first.generation_dir)) as generation:
        assert np.allclose(np.linalg.norm(generation.vectors, axis=1), 1.0)
        assert generation.sparse["postings"] and all(generation.sparse["nonzero_counts"])
        assert generation.profiles["dense"]["protocol"] == "dashscope.text-embedding"
        assert generation.profiles["sparse"]["model"] == "text-embedding-v4"
    result = query(index, "heartbeat", options=config, providers=providers)
    assert len(server.embeddings("query")) == 1  # Not one request per vector type.
    assert set(result.routes) == {"dense", "bm25", "sparse", "reranker"}
    assert result.status == "ok" and result.hits[0].source_name == "BETA.docx"
    assert all(hit.rerank_score is not None for hit in result.hits)
    for hit in result.hits:
        assert hit.rrf_score == pytest.approx(
            sum(1 / (60 + rank) for rank in hit.route_ranks.values())
        )
    filtered = query(
        index,
        "heartbeat",
        options=config,
        providers=providers,
        categories=["operations"],
        doc_ids=[ids["ALPHA"]],
    )
    assert {hit.doc_id for hit in filtered.hits} == {ids["ALPHA"]}
    assert all(route.count == 1 for route in filtered.routes.values())
    requests = len(server.calls)
    assert not query(index, "heartbeat", options=config, providers=providers, doc_ids=[]).hits
    assert len(server.calls) == requests
    query(index, "heartbeat", options=config, providers=providers, mode="bm25")
    assert len(server.calls) == requests
    for mode in ["dense", "sparse", "rrf"]:
        requests = len(server.embeddings("query"))
        single = query(index, "heartbeat", options=config, providers=providers, mode=mode)
        assert single.hits
        assert len(server.embeddings("query")) == requests + 1
        assert "reranker" not in single.routes

    with load_generation(Path(first.generation_dir)) as generation:
        ledger = CandidateLedger(
            generation,
            "heartbeat",
            options=config,
            providers=providers,
            topic_order=["T1"],
            max_pool=20,
            rerank_top_n=10,
        )
        requests = len(server.embeddings("query"))
        answer_evidence = ledger.collect(
            [Subquery("heartbeat", ("T1",))], round_index=1, uncovered_topics=["T1"]
        )
        assert len(server.embeddings("query")) == requests + 1
        assert answer_evidence.hits[0].source_name == "BETA.docx"

    assert build_index(artifacts, index, options=config, providers=providers).reused
    assert len(server.embeddings("document")) == 1
    changed = Document()
    changed.add_paragraph("ALPHA heartbeat starts after registration; updated synthetic content.")
    changed.save(source / "ALPHA.docx")
    parsed_new = parse_document(
        source / "ALPHA.docx",
        parsed,
        options=ParseOptions(
            allowed_roots=[source],
            category="operations",
            office_render="never",
            ocr_mode="disabled",
        ),
    )
    updated = build_index(
        [Path(parsed_new.artifact_dir)], index, options=config, providers=providers
    )
    assert updated.documents == 3 and updated.cache_hits == {"dense": 2, "sparse": 2}
    assert len(server.embeddings("document")) == 2
    assert len(server.embeddings("document")[-1]["input"]["texts"]) == 1
    assert Path(first.generation_dir).is_dir() and read_current(index)[0] == updated.generation_id

    server.fail = True
    before = len(server.embeddings("query"))
    with pytest.raises(DocQAError) as caught:
        query(index, "heartbeat", options=config, providers=providers)
    assert caught.value.code == "RETRIEVAL_FAILED"
    assert len(server.embeddings("query")) == before + 1
    # Explicit opt-in keeps both failed route diagnostics, never labels BM25 as full hybrid.
    degraded = query(
        index,
        "heartbeat",
        options=config.model_copy(update={"allow_degraded": True}),
        providers=providers,
        allow_degraded=True,
    )
    assert degraded.status == "degraded"
    assert all(degraded.routes[name].state == "failed" for name in ["dense", "sparse"])


def test_endpoint_variable_forwarded_to_headless_worker_and_not_to_profile(tmp_path, monkeypatch):
    config = options()
    for role in ["dense", "sparse"]:
        getattr(config, role).endpoint_env = "TEST_DASHSCOPE_ENDPOINT"
    monkeypatch.setenv("TEST_DASHSCOPE_ENDPOINT", "https://models.invalid/joint")
    monkeypatch.setenv("UNRELATED_SECRET", "must-not-forward")
    runner = HeadlessRunner(
        tmp_path / "index", tmp_path / "agent", AgentOptions(retrieval=config), tmp_path / "runs"
    )
    env = runner._environment()
    assert env["TEST_DASHSCOPE_ENDPOINT"] == "https://models.invalid/joint"
    assert "UNRELATED_SECRET" not in env
    original = make_providers(options())[0].fingerprint()
    assert make_providers(config)[0].fingerprint() == original
    for role in ["dense", "sparse"]:
        getattr(runner.options.retrieval, role).endpoint_env = "PYTHONPATH"
    with pytest.raises(DocQAError):
        runner._environment()


def test_two_route_generation_requires_rebuild_and_keeps_old_sources(corpus, tmp_path):
    _, _, artifacts, _ = corpus
    index = tmp_path / "upgrade"
    old = build_index(
        artifacts,
        index,
        options=DocBaseOptions(
            mode="demo",
            recall_routes=["dense", "bm25"],
            vector_backend="sqlite_vec",
        ),
    )
    config, server = options("sqlite_vec"), NativeServer()
    providers = make_providers(config, transport=httpx.MockTransport(server))
    with pytest.raises(DocQAError) as caught:
        query(index, "heartbeat", options=config, providers=providers)
    assert caught.value.code == "PROFILE_MISMATCH"
    assert not server.calls and read_current(index)[0] == old.generation_id
    new = build_index(artifacts, index, options=config, providers=providers)
    assert new.generation_id != old.generation_id and new.documents == 3
    assert Path(old.generation_dir).is_dir()
    assert query(index, "heartbeat", options=config, providers=providers).status == "ok"


@pytest.mark.parametrize("kind", ["missing_sparse", "service_error", "oversize"])
def test_build_never_publishes_half_a_joint_embedding(corpus, tmp_path, kind):
    _, _, artifacts, _ = corpus
    index = tmp_path / "atomic"
    old = build_index(artifacts, index, options=DocBaseOptions(mode="demo"))
    config = options()

    def respond(_):
        if kind == "service_error":
            return httpx.Response(503, json={"message": "do not expose provider body"})
        if kind == "oversize":
            return httpx.Response(200, content=b"x" * 1024)
        return httpx.Response(
            200,
            json={
                "output": {"embeddings": [embedding_row(i, sparse_embedding=[]) for i in range(3)]}
            },
        )

    if kind == "oversize":
        config.dense.max_response_bytes = config.sparse.max_response_bytes = 128
    providers = make_providers(config, transport=httpx.MockTransport(respond))
    with pytest.raises(DocQAError):
        build_index(artifacts, index, options=config, providers=providers)
    assert read_current(index)[0] == old.generation_id
    assert len(list((index / "generations").iterdir())) == 1


def test_endpoint_variable_reaches_background_ingestion_process(tmp_path, monkeypatch):
    config = options()
    for role in ["dense", "sparse"]:
        getattr(config, role).endpoint_env = "TEST_DASHSCOPE_ENDPOINT"
    monkeypatch.setenv("TEST_DASHSCOPE_ENDPOINT", "https://models.invalid/joint")
    monkeypatch.setenv("UNRELATED_SECRET", "must-not-forward")
    workspace = object.__new__(Workspace)
    workspace.options = WebOptions(state=tmp_path / "web", index=tmp_path / "index")
    workspace.options.state.mkdir()
    workspace.process, workspace._last_spawn = None, 0
    captured = []

    def launch(*args, **kwargs):
        captured.append(kwargs)
        return SimpleNamespace(poll=lambda: None)

    monkeypatch.setattr("docqa.web.workspace.subprocess.Popen", launch)
    workspace._spawn_service(
        SimpleNamespace(options=BackgroundOptions(retrieval=config), root=tmp_path / "service")
    )
    assert len(captured) == 1
    assert captured[0]["env"]["TEST_DASHSCOPE_ENDPOINT"] == "https://models.invalid/joint"
    assert "UNRELATED_SECRET" not in captured[0]["env"]
