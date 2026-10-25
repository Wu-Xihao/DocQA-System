"""Real image assets cross DOCX -> evidence -> saved citation -> HTTP safely.

VL and support decisions below are explicit test doubles. They verify transport,
ownership and integrity mechanisms, never genuine image understanding or answer
quality. Both the index and saved answer retain their synthetic labels.
"""

from __future__ import annotations

import json
import time
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, PngImagePlugin

from docqa.agent.config import AgentOptions
from docqa.agent.pipeline import (
    create_evidence_session,
    finalize_answer,
    retrieve_in_session,
)
from docqa.agent.sessions import SessionStore
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document, verify_artifact
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.docbase.storage import _io_path
from docqa.domain.answers import AnswerDraft, DraftClaim, SupportDecision, SupportReport
from docqa.domain.documents import VisionDescription
from docqa.web.app import create_app
from docqa.web.config import WebOptions

QUESTION = "IMG-QA-TEST 图片包含什么图形？"
DESCRIPTION = "IMG-QA-TEST：图片包含蓝色背景和红色矩形。仅用于自动化测试。"


class MockVision:
    def fingerprint(self):
        return {"adapter": "web-image-test-double", "revision": "1", "synthetic": True}

    def describe(self, image, *, ocr_text="", context=""):
        assert Path(image).is_file()
        return VisionDescription(status="ok", image_type="test", observations=[DESCRIPTION])


class MockSupportVerifier:
    def fingerprint(self):
        return {"adapter": "web-image-support-test-double", "synthetic": True}

    def verify(self, question, draft, items, *, deadline):
        assert question == QUESTION
        assert deadline > time.monotonic()
        assert len(draft.claims) == 1
        claim = draft.claims[0]
        item = next(value for value in items if value.evidence_id == claim.evidence_ids[0])
        assert item.origin == "vl_generated" and item.figure_ids
        assert claim.text == item.text and DESCRIPTION in item.text
        return SupportReport(
            claim_checks=[
                SupportDecision(
                    item_id=claim.claim_id,
                    state="supported",
                    evidence_ids=claim.evidence_ids,
                    reason="Explicit test-double decision; not semantic evaluation",
                )
            ],
            topic_checks=[
                SupportDecision(
                    item_id=topic.topic_id,
                    state="supported",
                    evidence_ids=claim.evidence_ids,
                    reason="Explicit test-double topic decision",
                )
                for topic in draft.required_topics
            ],
        )


@pytest.fixture(scope="module")
def image_corpus(tmp_path_factory):
    root = tmp_path_factory.mktemp("web-images")
    buffer = BytesIO()
    picture = Image.new("RGB", (64, 48), "blue")
    ImageDraw.Draw(picture).rectangle((12, 10, 50, 35), fill="red")
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("PrivateTestMetadata", "internal-test-value-not-for-preview")
    picture.save(buffer, format="PNG", pnginfo=metadata)
    buffer.seek(0)
    source = root / "图文测试.docx"
    document = Document()
    document.add_heading("图文引用自动化测试", 1)
    document.add_paragraph("正文介绍：以下图片用于测试图像引用与安全回源。")
    document.add_picture(buffer)
    document.save(source)
    parsed = parse_document(
        source,
        root / "parsed",
        options=ParseOptions(office_render="never", ocr_mode="disabled", category="image-test"),
        vision_client=MockVision(),
    )
    assert parsed.status == "complete" and parsed.eligible_for_indexing
    assert verify_artifact(parsed.artifact_dir)["valid"]
    retrieval = DocBaseOptions(mode="demo")
    built = build_index([Path(parsed.artifact_dir)], root / "kb", options=retrieval)
    return root / "kb", AgentOptions(mode="demo", retrieval=retrieval), built.generation_id


def bootstrap(client):
    result = client.get("/api/v1/bootstrap")
    assert result.status_code == 200, result.text
    client.headers["X-DocQA-Token"] = result.json()["csrf_token"]


def wait_job(client, job_id):
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        response = client.get("/api/v1/jobs/" + job_id)
        assert response.status_code == 200, response.text
        job = response.json()
        if job["state"] in {"succeeded", "failed", "cancelled"}:
            return job
        time.sleep(0.025)
    raise AssertionError("Image question did not finish")


def test_image_citation_serves_safe_png_with_scope_and_integrity_checks(image_corpus, tmp_path):
    index, agent, generation = image_corpus
    options = WebOptions(index=index, state=tmp_path / "web", agent=agent)
    retained = {}

    class SavedImageRunner:
        def run(self, job_id, request, *, cancelled, progress):
            assert not cancelled()
            assert request["generation_id"] == generation
            output = options.state / "agent" / job_id
            caller = "web:" + job_id
            progress("running")
            session = create_evidence_session(
                index,
                request["question"],
                output=output,
                options=agent,
                caller=caller,
                generation_id=request["generation_id"],
                categories=request["categories"],
                doc_ids=request["doc_ids"],
                required_topics=request["topics"],
            )
            pack = retrieve_in_session(
                session.session_id,
                request["question"],
                output=output,
                options=agent,
                caller=caller,
            )
            image_item = next(
                item
                for item in pack.items
                if item.origin == "vl_generated" and item.figure_ids and DESCRIPTION in item.text
            )
            draft = AnswerDraft(
                requested_format=request["output_format"],
                claims=[
                    DraftClaim(
                        claim_id="image_claim",
                        topic_ids=[session.required_topics[0].topic_id],
                        text=image_item.text,
                        kind="quotation",
                        evidence_ids=[image_item.evidence_id],
                        quotes={image_item.evidence_id: image_item.text},
                    )
                ],
                used_evidence_ids=[image_item.evidence_id],
            )
            answer = finalize_answer(
                session.session_id,
                draft,
                output=output,
                options=agent,
                caller=caller,
                verifier=MockSupportVerifier(),
            )
            assert answer.synthetic and not answer.verified and not answer.production_ready
            assert answer.citations[0].origin == "vl_generated"
            store = SessionStore(output)
            retained["source"] = store.resolve_citation(
                answer.answer_id, answer.citations[0].citation_id
            )
            progress("complete")
            return {
                "answer_id": answer.answer_id,
                "session_id": answer.session_id,
                "status": answer.status,
            }

    app = create_app(options, runner_factory=lambda _job: SavedImageRunner())
    with TestClient(app, base_url="http://127.0.0.1:8765", raise_server_exceptions=False) as owner:
        bootstrap(owner)
        accepted = owner.post("/api/v1/questions", json={"question": QUESTION})
        assert accepted.status_code == 202, accepted.text
        job = wait_job(owner, accepted.json()["job_id"])
        assert job["state"] == "succeeded", job
        answer = job["answer"]
        assert answer["synthetic"] and not answer["verified"] and not answer["production_ready"]
        citation = answer["citations"][0]
        assert citation["origin"] == "vl_generated"
        assert citation["resources"]["images"]
        image_url = citation["resources"]["images"][0]
        meta_url = citation["resources"]["metadata"]
        response = owner.get(image_url)
        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == "image/png"
        assert response.headers["x-content-type-options"] == "nosniff"
        assert response.content.startswith(b"\x89PNG\r\n\x1a\n")
        assert b"internal-test-value-not-for-preview" not in response.content
        with Image.open(BytesIO(response.content)) as preview:
            assert preview.size == (64, 48) and preview.mode == "RGBA"
            assert "PrivateTestMetadata" not in preview.info
            assert preview.getpixel((20, 20)) == (255, 0, 0, 255)
        metadata = owner.get(meta_url)
        assert metadata.status_code == 200, metadata.text
        public = metadata.json()
        assert public["resources"]["images"] == citation["resources"]["images"]
        assert public["origin"] == "vl_generated" and DESCRIPTION in public["quote"]
        serialized = json.dumps(public, ensure_ascii=False)
        for private in (
            str(index),
            str(tmp_path),
            "artifact_dir",
            "markdown_path",
            "snapshot_path",
            "image_paths",
        ):
            assert private not in serialized
        assert any(locator["locator_type"] == "docx_object" for locator in public["locators"])

        # A second signed browser session cannot adopt the owner's known URLs.
        other = TestClient(app, base_url="http://127.0.0.1:8765", raise_server_exceptions=False)
        try:
            bootstrap(other)
            for url in (image_url, meta_url, citation["resources"]["snapshot"]):
                denied = other.get(url)
                assert denied.status_code == 404, denied.text
        finally:
            other.close()
        base = image_url.split("?", 1)[0]
        assert owner.get(base + "?kind=image&image=999").status_code == 404
        assert owner.get(base + "?kind=image&image=-1").status_code == 422

        # Cookie-authenticated <img> links work without a token header; JSON
        # metadata still requires the normal browser API token.
        token = owner.headers.pop("X-DocQA-Token")
        assert owner.get(image_url).status_code == 200
        owner.headers["X-DocQA-Token"] = token

        asset = _io_path(Path(retained["source"]["images"][0]["path"]))
        original = asset.read_bytes()
        try:
            asset.write_bytes(b"tampered-source-image")
            for url in (image_url, meta_url, citation["resources"]["snapshot"]):
                rejected = owner.get(url)
                assert rejected.status_code == 503, rejected.text
                assert rejected.json()["error"]["code"] == "GENERATION_CORRUPT"
                assert "tampered-source-image" not in rejected.text
        finally:
            asset.write_bytes(original)
