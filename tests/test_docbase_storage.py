"""Immutable-package gates, filtered exact dense retrieval and SQLite CAS publication."""

from __future__ import annotations

import json
import sqlite3
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import numpy as np
import pytest

from docqa.docbase.storage import (
    _validate_postings,
    load_generation,
    publish,
    read_current,
    search_dense,
    verify_generation,
    write_generation,
)
from docqa.domain.errors import DocQAError


def _profiles():
    return {
        "mode": "demo",
        "dense": {"dimension": 2, "synthetic": True},
        "sparse": {},
        "reranker": {},
    }


def _empty(directory, identity="generation1", build_key="build1"):
    return write_generation(
        directory,
        generation_id=identity,
        chunks=[],
        dense_vectors=[],
        bm25={
            "schema_version": "docqa.bm25.v1",
            "row_count": 0,
            "doc_lengths": [],
            "postings": {},
            "avgdl": 0.0,
        },
        sparse={
            "schema_version": "docqa.sparse.v1",
            "row_count": 0,
            "nonzero_counts": [],
            "postings": {},
        },
        members=[],
        profiles=_profiles(),
        build_key=build_key,
        empty=True,
    )


def _rehash(directory, relative):
    path = directory / "manifest.json"
    manifest = json.loads(path.read_bytes())
    data = (directory / relative).read_bytes()
    manifest["files"][relative] = {"size": len(data), "sha256": sha256(data).hexdigest()}
    path.write_text(json.dumps(manifest), encoding="utf-8")


def test_empty_generation_round_trip_uses_readonly_mmap(tmp_path):
    directory = tmp_path / "candidate"
    manifest = _empty(directory)
    assert manifest["synthetic"] is True
    assert manifest["production_ready"] is False
    assert verify_generation(directory)["chunk_count"] == 0
    with load_generation(directory) as generation:
        assert generation.vectors.shape == (0, 2)
        assert not generation.vectors.flags.writeable
        assert generation.members == []
    assert generation.vectors._mmap.closed


def test_accidental_empty_generation_is_rejected(tmp_path):
    with pytest.raises(DocQAError, match="Cannot write generation") as error:
        write_generation(
            tmp_path / "bad",
            generation_id="bad",
            chunks=[],
            dense_vectors=[],
            bm25={},
            sparse={},
            members=[],
            profiles=_profiles(),
            build_key="bad",
        )
    assert error.value.code == "GENERATION_INVALID"


@pytest.mark.parametrize("relative", ["dense/vectors.npy", "catalog.sqlite", "bm25/index.json"])
def test_modified_file_digest_is_rejected(tmp_path, relative):
    directory = tmp_path / "candidate"
    _empty(directory)
    with (directory / relative).open("ab") as stream:
        stream.write(b"modified")
    with pytest.raises(DocQAError) as error:
        verify_generation(directory)
    assert error.value.code == "GENERATION_CORRUPT"


def test_unlisted_file_is_rejected(tmp_path):
    directory = tmp_path / "candidate"
    _empty(directory)
    (directory / "unlisted").write_text("stray", encoding="utf-8")
    with pytest.raises(DocQAError):
        verify_generation(directory)


@pytest.mark.parametrize(
    "relative", ["../outside", "C:/outside", "dense\\vectors.npy", "dense/./vectors.npy"]
)
def test_manifest_path_escaping_is_rejected(tmp_path, relative):
    directory = tmp_path / "candidate"
    _empty(directory)
    path = directory / "manifest.json"
    manifest = json.loads(path.read_bytes())
    manifest["files"][relative] = {"size": 0, "sha256": "0" * 64}
    path.write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(DocQAError):
        verify_generation(directory)


def test_recomputed_digest_does_not_hide_invalid_row_mapping(tmp_path):
    directory = tmp_path / "candidate"
    _empty(directory)
    (directory / "dense/row_ids.jsonl").write_text(
        '{"row_id":0,"chunk_id":"unknown"}\n', encoding="utf-8"
    )
    _rehash(directory, "dense/row_ids.jsonl")
    with pytest.raises(DocQAError):
        verify_generation(directory)


def test_recomputed_digest_does_not_hide_profile_change(tmp_path):
    directory = tmp_path / "candidate"
    _empty(directory)
    (directory / "profiles.json").write_text('{"dense":{"dimension":3}}', encoding="utf-8")
    _rehash(directory, "profiles.json")
    with pytest.raises(DocQAError):
        verify_generation(directory)


def test_object_numpy_array_is_rejected_even_after_rehash(tmp_path):
    directory = tmp_path / "candidate"
    _empty(directory)
    np.save(directory / "dense/vectors.npy", np.array([{"unsafe": True}], dtype=object))
    _rehash(directory, "dense/vectors.npy")
    with pytest.raises(DocQAError):
        verify_generation(directory)


def test_dense_filter_happens_before_top_k_and_ties_are_stable():
    vectors = np.array([[1, 0], [0, 1], [0, 1], [-1, 0]], dtype=np.float32)
    # The most similar global row is excluded; all allowed rows still compete.
    assert search_dense(vectors, [1, 0], {1, 2, 3}, 2) == [(1, 0.0), (2, 0.0)]
    assert search_dense(vectors, [1, 0], [], 2) == []


@pytest.mark.parametrize("query", [[0, 0], [1, 1], [float("nan"), 0], [1], ["x", "y"]])
def test_dense_rejects_nonunit_or_invalid_query(query):
    with pytest.raises(DocQAError) as error:
        search_dense(np.array([[1, 0]], dtype=np.float32), query, {0}, 1)
    assert error.value.code == "MODEL_OUTPUT_INVALID"


@pytest.mark.parametrize("eligible", [[-1], [1], [True], ["0"]])
def test_dense_rejects_invalid_filter(eligible):
    with pytest.raises(DocQAError):
        search_dense(np.array([[1, 0]], dtype=np.float32), [1, 0], eligible, 1)


@pytest.mark.parametrize(
    "change",
    ["duplicate", "out_of_range", "tf", "lengths", "sparse_nan", "sparse_counts", "sparse_token"],
)
def test_posting_semantic_gates(change):
    bm25 = {
        "schema_version": "docqa.bm25.v1",
        "row_count": 2,
        "doc_lengths": [2, 1],
        "postings": {"term": [[0, 2], [1, 1]]},
        "avgdl": 1.5,
    }
    sparse = {
        "schema_version": "docqa.sparse.v1",
        "row_count": 2,
        "nonzero_counts": [1, 0],
        "postings": {"12": [[0, 0.5]]},
    }
    _validate_postings(bm25, sparse, 2)
    if change == "duplicate":
        bm25["postings"]["term"] = [[0, 1], [0, 1]]
    elif change == "out_of_range":
        bm25["postings"]["term"] = [[2, 1]]
    elif change == "tf":
        bm25["postings"]["term"][0][1] = 1.5
    elif change == "lengths":
        bm25["doc_lengths"] = [1, 1]
    elif change == "sparse_nan":
        sparse["postings"]["12"][0][1] = float("nan")
    elif change == "sparse_counts":
        sparse["nonzero_counts"] = [0, 1]
    else:
        sparse["postings"] = {"01": [[0, 0.5]]}
    with pytest.raises(ValueError):
        _validate_postings(bm25, sparse, 2)


def test_publication_is_single_sqlite_cas_and_sequence_blocks_aba(tmp_path):
    root = tmp_path / "kb"
    assert read_current(root) == (None, 0)
    for name in ("g1", "g2"):
        _empty(root / "generations" / name, name, name)
    assert publish(root, "g1", None, 0, "g1") == 1
    assert read_current(root) == ("g1", 1)
    assert publish(root, "g2", "g1", 1, "g2") == 2
    assert publish(root, "g1", "g2", 2, "g1") == 3
    with pytest.raises(DocQAError) as error:
        publish(root, "g2", "g1", 1, "g2")
    assert error.value.code == "PUBLISH_STALE"
    assert read_current(root) == ("g1", 3)
    assert not (root / "active.json").exists()
    assert verify_generation(root / "generations/g2")["valid"]


def test_wrong_publication_build_key_cannot_change_pointer(tmp_path):
    root = tmp_path / "kb"
    _empty(root / "generations/g1", "g1", "build1")
    with pytest.raises(DocQAError) as error:
        publish(root, "g1", None, 0, "another-build")
    assert error.value.code == "GENERATION_INVALID"
    assert read_current(root) == (None, 0)


@pytest.fixture
def source_generation(tmp_path, request):
    from docx import Document
    from PIL import Image

    from docqa.doc2md.config import ParseOptions
    from docqa.doc2md.pipeline import parse_document
    from docqa.docbase.analyzer import Analyzer
    from docqa.docbase.chunker import chunk_artifact, chunk_profile
    from docqa.docbase.lexical import build_bm25, build_sparse
    from docqa.docbase.storage import _io_path
    from docqa.domain.documents import DocumentIR, VisionDescription
    from docqa.domain.rag import ChunkOptions

    long_image = getattr(request, "param", False)
    source = tmp_path / (("DocQA_Architecture_" * 5 if long_image else "文档") + ".docx")
    document = Document()
    document.add_heading("系统设计", 1)
    document.add_paragraph("SD5993 调用 FE_HEART_Start 开始心跳。")
    document.add_paragraph("检索只返回范围内的文档。")
    if long_image:
        buffer = BytesIO()
        Image.new("RGB", (40, 40), "red").save(buffer, format="PNG")
        buffer.seek(0)
        document.add_picture(buffer)
    document.save(source)

    class FakeVision:
        def fingerprint(self):
            return {"provider": "test", "revision": "1"}

        def describe(self, image, *, ocr_text="", context=""):
            return VisionDescription(status="ok", observations=["Visible test image"])

    parsed = parse_document(
        source,
        _io_path(tmp_path / "parsed"),
        options=ParseOptions(office_render="never", ocr_mode="disabled"),
        vision_client=FakeVision() if long_image else None,
    )
    assert parsed.status == "complete"
    artifact = Path(parsed.artifact_dir)
    ir = DocumentIR.model_validate_json((artifact / "document.ir.json").read_bytes())
    source_map = json.loads((artifact / "source_map.json").read_bytes())
    options = ChunkOptions()
    chunks = chunk_artifact(ir, source_map, options)
    analyzer = Analyzer()
    profiles = {
        **_profiles(),
        "chunk": chunk_profile(options),
        "analyzer": analyzer.profile(),
        "bm25": {"k1": 1.2, "b": 0.75},
    }
    member = {
        "doc_id": ir.doc_id,
        "revision_id": ir.revision_id,
        "artifact_id": ir.artifact_id,
        "source_name": ir.source_name,
        "category": ir.category,
        "artifact_ref": f"sources/artifacts/{ir.artifact_id}",
    }
    directory = tmp_path / "candidate"
    write_generation(
        directory,
        generation_id="full",
        chunks=chunks,
        dense_vectors=[[1, 0] for _ in chunks],
        bm25=build_bm25(chunks, analyzer),
        sparse=build_sparse([{row + 1: 0.5} for row in range(len(chunks))]),
        members=[member],
        profiles=profiles,
        build_key="full-build",
        source_packages={ir.artifact_id: artifact},
    )
    return directory, artifact


def test_source_package_is_self_contained_and_survives_original_removal(source_generation):
    directory, original = source_generation
    # Renaming the original artifact directory proves the generation uses its own copy.
    original.rename(original.with_name(original.name + "-moved"))
    with load_generation(directory) as generation:
        assert generation.chunks
        assert all(span.locators for chunk in generation.chunks for span in chunk.source_spans)
        assert verify_generation(directory)["member_count"] == 1


@pytest.mark.parametrize("source_generation", [True], indirect=True)
def test_long_image_paths_load_without_windows_long_path_setting(source_generation):
    directory, _ = source_generation
    with load_generation(directory) as generation:
        assert any(len(str(path)) > 264 for path in generation.directory.rglob("*.png"))
        assert any(chunk.figure_ids for chunk in generation.chunks)
        assert verify_generation(directory)["valid"]


def test_rehashed_bm25_statistics_cannot_hide_changed_text_postings(source_generation):
    directory, _ = source_generation
    path = directory / "bm25/index.json"
    index = json.loads(path.read_bytes())
    first = next(iter(index["postings"]))
    index["postings"]["invented_term"] = index["postings"].pop(first)
    path.write_text(json.dumps(index), encoding="utf-8")
    _rehash(directory, "bm25/index.json")
    with pytest.raises(DocQAError) as error:
        verify_generation(directory)
    assert error.value.code == "GENERATION_CORRUPT"


@pytest.mark.parametrize("tamper", ["source_digest", "display_text", "category", "locator"])
def test_rehashed_chunk_and_row_identity_cannot_hide_source_divergence(source_generation, tamper):
    from docqa.domain.rag import Chunk, chunk_identity

    directory, _ = source_generation
    connection = sqlite3.connect(directory / "catalog.sqlite")
    try:
        row_id, payload = connection.execute(
            "SELECT row_id,payload FROM chunks ORDER BY row_id LIMIT 1"
        ).fetchone()
        value = json.loads(payload)
        if tamper == "source_digest":
            value["source_spans"][0]["source_text_sha256"] = "0" * 64
        elif tamper == "display_text":
            value["display_text"] = "虚构内容"
        elif tamper == "category":
            value["category"] = "invisible-other-category"
        else:
            value["source_spans"][0]["locators"][0]["display_label"] = "不存在的位置"
        value["chunk_id"] = chunk_identity(Chunk.model_validate(value))
        with connection:
            connection.execute(
                "UPDATE chunks SET chunk_id=?,payload=? WHERE row_id=?",
                (value["chunk_id"], json.dumps(value), row_id),
            )
    finally:
        connection.close()
    row_path = directory / "dense/row_ids.jsonl"
    rows = [json.loads(line) for line in row_path.read_bytes().splitlines()]
    rows[row_id]["chunk_id"] = value["chunk_id"]
    row_path.write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    _rehash(directory, "catalog.sqlite")
    _rehash(directory, "dense/row_ids.jsonl")
    with pytest.raises(DocQAError) as error:
        verify_generation(directory)
    assert error.value.code == "GENERATION_CORRUPT"
