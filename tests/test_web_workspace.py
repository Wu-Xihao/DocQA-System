"""Real local HTTP workspace flows with demo models; no external model requests."""

import io
import time
from zipfile import ZipFile

import pytest
import test_agent_integration as shared
from docx import Document
from fastapi.testclient import TestClient

from docqa.agent.config import AgentOptions
from docqa.docbase.config import DocBaseOptions
from docqa.web.app import create_app
from docqa.web.config import WebOptions


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    return shared.corpus.__wrapped__(tmp_path_factory)


def client_for(options):
    return TestClient(create_app(options), base_url="http://127.0.0.1:8765")


def connect(client):
    bootstrap = client.get("/api/v1/bootstrap").json()
    client.headers["X-DocQA-Token"] = bootstrap["csrf_token"]
    return bootstrap


def wait(client, path, *, ready="ready", maximum=25):
    deadline = time.monotonic() + maximum
    while time.monotonic() < deadline:
        response = client.get(path)
        assert response.status_code == 200, response.text
        result = response.json()
        if result["state"] == ready:
            return result
        assert result["state"] not in {"failed", "cancelled"}, result
        time.sleep(0.15)
    pytest.fail("Task did not finish")


def send(client, conversation, key="request-unique-key-0001", **values):
    payload = {"question": "SD5993 心跳如何启动？", "client_request_id": key, **values}
    response = client.post(f"/api/v1/conversations/{conversation}/messages", json=payload)
    assert response.status_code == 202, response.text
    return response.json()


def test_conversation_messages_are_persisted_idempotent_and_owner_bound(corpus, tmp_path):
    index, agent, _ = corpus
    options = WebOptions(index=index, state=tmp_path, agent=agent)
    with client_for(options) as client:
        connect(client)
        conversation = client.post("/api/v1/conversations", json={}).json()
        job = send(client, conversation["id"])
        duplicate = send(client, conversation["id"])
        assert duplicate["job_id"] == job["job_id"]
        done = wait(client, "/api/v1/jobs/" + job["job_id"], ready="succeeded")
        assert done["answer"]["citations"]
        assert done["answer"]["performance"]["model_tasks"] >= 3
        assert "chat.verify" in done["answer"]["performance"]["operations"]
        assert "trace" not in done["answer"]
        messages = client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()[
            "items"
        ]
        assert [m["role"] for m in messages] == ["user", "assistant"]
        assert messages[1]["job"]["answer"]["question"] == "SD5993 心跳如何启动？"
        second = send(
            client, conversation["id"], key="request-unique-key-0002", question="调用之后要做什么？"
        )
        wait(client, "/api/v1/jobs/" + second["job_id"], ready="succeeded")
        record = client.app.state.manager.store.get(second["job_id"])
        assert record["payload"]["context_questions"] == ["SD5993 心跳如何启动？"]
        assert record["payload"]["generation_id"] == job["generation_id"]
        assert (
            client.post(
                f"/api/v1/conversations/{conversation['id']}/messages",
                json={"question": "different", "client_request_id": "request-unique-key-0001"},
            ).status_code
            == 409
        )
        # A different browser session cannot enumerate this conversation.
        client.cookies.clear()
        connect(client)
        assert client.get(f"/api/v1/conversations/{conversation['id']}").status_code == 404
        assert client.get(f"/api/v1/conversations/{conversation['id']}/messages").status_code == 404


def test_scope_empty_and_foreign_document_cannot_become_all(corpus, tmp_path):
    index, agent, _ = corpus
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        assert (
            client.post(
                f"/api/v1/conversations/{cid}/messages",
                json={
                    "question": "x",
                    "client_request_id": "scope-empty-123456",
                    "scope_mode": "selected",
                    "doc_ids": [],
                },
            ).status_code
            == 422
        )
        assert (
            client.post(
                f"/api/v1/conversations/{cid}/messages",
                json={
                    "question": "x",
                    "client_request_id": "scope-other-123456",
                    "doc_ids": ["absent"],
                },
            ).status_code
            == 409
        )
        assert not client.get("/api/v1/jobs").json()["jobs"]


def test_history_migration_rename_and_permanent_delete(corpus, tmp_path):
    index, agent, _ = corpus
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        job = client.post("/api/v1/questions", json={"question": "SD5993"}).json()
        wait(client, "/api/v1/jobs/" + job["job_id"], ready="succeeded")
        first = client.get("/api/v1/conversations").json()["items"]
        second = client.get("/api/v1/conversations").json()["items"]
        assert len(first) == len(second) == 1 and first[0]["id"] == second[0]["id"]
        c = first[0]
        updated = client.patch(
            "/api/v1/conversations/" + c["id"],
            json={"title": "新标题", "pinned": True, "version": c["version"]},
        ).json()
        assert updated["title"] == "新标题" and updated["pinned"]
        assert (
            client.patch(
                "/api/v1/conversations/" + c["id"], json={"title": "stale", "version": c["version"]}
            ).status_code
            == 409
        )
        client.delete("/api/v1/conversations/" + c["id"])
        assert not client.get("/api/v1/conversations").json()["items"]
        assert client.post(f"/api/v1/conversations/{c['id']}/restore", json={}).status_code == 404


def test_studio_artifacts_previews_and_real_file_exports(corpus, tmp_path):
    index, agent, _ = corpus
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        question = send(client, cid)
        answer = wait(client, "/api/v1/jobs/" + question["job_id"], ready="succeeded")["answer"]
        for kind, formats in [
            ("table", ["csv", "json"]),
            ("mindmap", ["svg", "png"]),
            ("report", ["docx", "md"]),
            ("slides", ["pptx", "md"]),
        ]:
            response = client.post(
                "/api/v1/studio/jobs",
                json={"conversation_id": cid, "answer_id": answer["answer_id"], "kind": kind},
            )
            assert response.status_code == 202, response.text
            result = wait(client, "/api/v1/studio/jobs/" + response.json()["id"])
            artifact = client.get("/api/v1/artifacts/" + result["artifact_id"]).json()
            assert (
                artifact["content"]["facts"]
                and artifact["generation_id"] == answer["generation_id"]
            )
            for format_name in formats:
                export = client.post(
                    f"/api/v1/artifacts/{artifact['id']}/exports", json={"format": format_name}
                ).json()
                exported = wait(client, "/api/v1/exports/" + export["id"])
                data = client.get(f"/api/v1/exports/{export['id']}/content").content
                assert data and exported["state"] == "ready"
                if format_name == "md" and kind == "slides":
                    value = data.decode("utf-8")
                    for group in artifact["content"]["tree"]["children"]:
                        assert "## " + group["label"] in value
                        for fact in group["children"]:
                            assert fact["text"] in value
                            assert all("[" + source + "]" in value for source in fact["sources"])
                if format_name in {"docx", "pptx"}:
                    with ZipFile(io.BytesIO(data)) as package:
                        assert "[Content_Types].xml" in package.namelist()
                if format_name == "svg":
                    assert b"<svg" in data and b"<script" not in data
                if format_name == "png":
                    from PIL import Image

                    with Image.open(io.BytesIO(data)) as picture:
                        assert picture.width == 2100 and picture.height >= 480


def test_upload_binary_stream_commit_service_and_publication(tmp_path):
    service = tmp_path / "service"
    retrieval = DocBaseOptions(
        mode="demo", recall_routes=["dense", "bm25"], incomplete_policy="native_only"
    )
    options = WebOptions(
        index=service / "index",
        service_state=service,
        state=tmp_path / "web",
        upload_root=tmp_path / "sources",
        service_mode="external",
        agent=AgentOptions(mode="demo", retrieval=retrieval),
    )
    data = io.BytesIO()
    document = Document()
    document.add_paragraph("UPLOAD_TEST 本地上传资料每隔30秒发送心跳。")
    document.save(data)
    with ZipFile(data, "a") as package:
        package.writestr("test-padding.bin", b"x" * 100000)
    raw = data.getvalue()
    assert len(raw) > 65536
    with client_for(options) as client:
        connect(client)
        record = client.post(
            "/api/v1/uploads", json={"name": "原始资料.docx", "size": len(raw)}
        ).json()
        identity = record["id"]
        result = client.put(
            f"/api/v1/uploads/{identity}/content",
            content=raw,
            headers={"Content-Type": "application/octet-stream"},
        )
        assert result.status_code == 200, result.text
        committed = client.post(f"/api/v1/uploads/{identity}/commit", json={})
        assert committed.status_code == 202, committed.text
        assert (
            client.post(f"/api/v1/uploads/{identity}/commit", json={}).json()["doc_id"]
            == committed.json()["doc_id"]
        )
        workspace = client.app.state.workspace
        service = workspace.service()
        actual_build = service.build_function

        def failed_build(*args, **kwargs):
            from docqa.domain.errors import DocQAError

            raise DocQAError("MODEL_AUTH_FAILED", "Simulated unavailable embedding credentials")

        service.build_function = failed_build
        service.run_once()
        workspace.refresh_uploads()
        failure = client.get("/api/v1/uploads/" + identity).json()
        assert failure["state"] == "failed" and failure["error_code"] == "INDEX_BUILD_FAILED"
        assert (
            client.post(
                f"/api/v1/documents/{committed.json()['doc_id']}/reprocess", json={}
            ).status_code
            == 202
        )
        service.build_function = actual_build
        service.run_once()
        workspace.refresh_uploads()
        done = client.get("/api/v1/uploads/" + identity).json()
        assert done["state"] == "ready", done
        assert client.get("/api/v1/documents").json()["documents"]
        # Pin an old conversation, then publish another uploaded document.
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        old = client.post(f"/api/v1/conversations/{cid}/upgrade", json={}).json()
        second_document = Document()
        second_document.add_paragraph("NEW_UPLOAD 第二份资料已完成入库。")
        output = io.BytesIO()
        second_document.save(output)
        second_raw = output.getvalue()
        second = client.post(
            "/api/v1/uploads", json={"name": "新增资料.docx", "size": len(second_raw)}
        ).json()
        second_url = "/api/v1/uploads/" + second["id"]
        assert (
            client.put(
                second_url + "/content",
                content=second_raw,
                headers={"Content-Type": "application/octet-stream"},
            ).status_code
            == 200
        )
        committed_second = client.post(second_url + "/commit", json={}).json()
        service.run_once()
        workspace.refresh_uploads()
        old_view = client.get(f"/api/v1/documents?conversation_id={cid}").json()
        assert old_view["generation_id"] == old["generation_id"]
        assert old_view["latest_generation_id"] != old["generation_id"]
        assert len(old_view["documents"]) == 1
        assert len(old_view["library_documents"]) == 2
        assert committed_second["doc_id"] in {d["doc_id"] for d in old_view["library_documents"]}
        assert (
            client.post(
                f"/api/v1/conversations/{cid}/messages",
                json={
                    "question": "NEW_UPLOAD",
                    "doc_ids": [committed_second["doc_id"]],
                    "scope_mode": "selected",
                    "client_request_id": "new-library-old-scope-0001",
                },
            ).status_code
            == 409
        )
        # Repeated UI polling reuses checked immutable catalog metadata.
        original_catalog = workspace.catalog

        def unexpected_catalog(*args, **kwargs):
            raise AssertionError("Unchanged display metadata should be cached")

        workspace.catalog = unexpected_catalog
        try:
            assert client.get(f"/api/v1/documents?conversation_id={cid}").status_code == 200
        finally:
            workspace.catalog = original_catalog
        upgraded = client.post(f"/api/v1/conversations/{cid}/upgrade", json={}).json()
        assert upgraded["generation_id"] == old_view["latest_generation_id"]
        assert len(client.get(f"/api/v1/documents?conversation_id={cid}").json()["documents"]) == 2
        assert (
            client.post("/api/v1/uploads", json={"name": "bad.exe", "size": 4}).status_code == 415
        )


def test_upload_ownership_and_declared_size_validation(tmp_path):
    s = tmp_path / "s"
    options = WebOptions(
        index=s / "index",
        service_state=s,
        state=tmp_path / "w",
        upload_root=tmp_path / "input",
        service_mode="external",
        agent=AgentOptions(mode="demo", retrieval=DocBaseOptions(mode="demo")),
    )
    with client_for(options) as client:
        connect(client)
        identity = client.post("/api/v1/uploads", json={"name": "x.docx", "size": 2}).json()["id"]
        response = client.put(
            f"/api/v1/uploads/{identity}/content",
            content=b"bigger",
            headers={"Content-Type": "application/octet-stream"},
        )
        assert response.status_code == 413
        client.cookies.clear()
        connect(client)
        assert client.get("/api/v1/uploads/" + identity).status_code == 404


def test_automatic_ingestion_native_only_images_and_integrity(tmp_path):
    """A real child Service publishes, then images stay reference-only and owner-bound."""
    from pathlib import Path

    from PIL import Image

    service = tmp_path / "service"
    retrieval = DocBaseOptions(
        mode="demo", recall_routes=["dense", "bm25"], incomplete_policy="native_only"
    )
    options = WebOptions(
        index=service / "index",
        service_state=service,
        state=tmp_path / "web",
        upload_root=tmp_path / "sources",
        service_mode="auto",
        agent=AgentOptions(mode="demo", retrieval=retrieval),
    )
    picture = io.BytesIO()
    Image.new("RGB", (300, 120), "#6657d9").save(picture, "PNG")
    document = Document()
    document.add_paragraph("DOCQA_IMAGE 系统处理包括解析、检索和问答三个阶段。")
    picture.seek(0)
    document.add_picture(picture)
    data = io.BytesIO()
    document.save(data)
    raw = data.getvalue()
    with client_for(options) as client:
        connect(client)
        workspace = client.app.state.workspace
        # Keep the native-only policy without invoking OCR or Office services.
        from docqa.doc2md.config import ParseOptions
        from docqa.service.config import ServiceOptions

        folder = options.upload_root
        folder.mkdir(parents=True)
        import json

        config = tmp_path / "service.json"
        config.write_text(
            json.dumps(
                ServiceOptions(
                    allowed_source_roots=[folder],
                    retrieval=retrieval,
                    parse=ParseOptions(office_render="never", ocr_mode="disabled"),
                    settle_seconds=0,
                ).model_dump(mode="json")
            ),
            encoding="utf-8",
        )
        workspace.options.service_config = config
        record = client.post("/api/v1/uploads", json={"name": "流程.docx", "size": len(raw)}).json()
        path = f"/api/v1/uploads/{record['id']}"
        assert (
            client.put(
                path + "/content", content=raw, headers={"Content-Type": "application/octet-stream"}
            ).status_code
            == 200
        )
        assert client.post(path + "/commit", json={}).status_code == 202
        ready = wait(client, path, ready="ready_with_gaps", maximum=35)
        assert ready["doc_id"]
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        question = send(client, cid, question="DOCQA_IMAGE")
        answer = wait(client, "/api/v1/jobs/" + question["job_id"], ready="succeeded")["answer"]
        assert answer["status"] == "partial"
        catalog = client.get("/api/v1/documents").json()
        preview = client.get(f"/api/v1/documents/{ready['doc_id']}/content?conversation_id={cid}")
        assert preview.status_code == 200 and "DOCQA\\_IMAGE" in preview.text
        assert len(catalog["documents"]) == 1
        url = f"/api/v1/queries/{answer['answer_id']}/related-images"
        gallery = client.get(url).json()
        assert gallery["items"] and all(i["role"] == "reference" for i in gallery["items"])
        image = gallery["items"][0]
        assert "artifact" not in image and "relative_path" not in image
        assert client.get(image["thumbnail_url"]).headers["content-type"] == "image/png"
        assert client.get(url).json()["items"][0]["id"] == image["id"]
        descriptor = workspace.store.object_get("media", image["id"])
        asset = Path(descriptor["artifact"]) / descriptor["relative_path"]
        original = asset.read_bytes()
        try:
            asset.write_bytes(b"changed")
            assert (
                client.get(image["preview_url"]).json()["error"]["code"] == "SOURCE_HASH_MISMATCH"
            )
        finally:
            asset.write_bytes(original)
        client.cookies.clear()
        connect(client)
        assert client.get(image["preview_url"]).status_code == 404
        assert client.get(url).status_code == 404
        assert (
            client.get(
                f"/api/v1/documents/{ready['doc_id']}/content?conversation_id={cid}"
            ).status_code
            == 404
        )


def test_retries_do_not_revalidate_a_changed_catalog(corpus, tmp_path, monkeypatch):
    index, agent, _ = corpus
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        job = send(client, cid)
        wait(client, "/api/v1/jobs/" + job["job_id"], ready="succeeded")

        def unavailable(*args, **kwargs):
            raise AssertionError("An idempotent retry must not load a new catalog")

        monkeypatch.setattr(client.app.state.workspace, "catalog", unavailable)
        assert send(client, cid)["job_id"] == job["job_id"]
