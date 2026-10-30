"""Real counters, ownership and batch completion; no external requests."""

import base64
import io
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.domain.errors import DocQAError
from docqa.service.queue import JobQueue


def test_progress_is_owned_transient_and_cleared_on_retry(tmp_path):
    queue = JobQueue(tmp_path)
    job = queue.enqueue("parse", {}, "parse:progress")
    queue.claim("worker")
    with pytest.raises(DocQAError):
        queue.report_progress(job["job_id"], "other", "images", 1, 2)
    for current, total in [(True, 2), (3, 2), (-1, 2), (0, None)]:
        with pytest.raises(DocQAError):
            queue.report_progress(job["job_id"], "worker", "images", current, total)
    queue.report_progress(job["job_id"], "worker", "images", 1, 2)
    assert queue.get(job["job_id"])["result"]["progress"] == {
        "stage": "images",
        "current": 1,
        "total": 2,
    }
    queue.fail(job["job_id"], "worker", "TEST_RETRY", retryable=True)
    assert queue.get(job["job_id"])["result"] is None


def test_parser_and_builder_emit_real_unit_counts(tmp_path):
    path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_paragraph("PROGRESS_SYNTHETIC reliable native paragraph")
    for color in ("blue", "red"):
        data = io.BytesIO()
        Image.new("RGB", (32, 20), color).save(data, format="PNG")
        data.seek(0)
        document.add_picture(data)
    document.save(path)
    events = []
    result = parse_document(
        path,
        tmp_path / "parsed",
        options=ParseOptions(office_render="never", ocr_mode="disabled"),
        progress=lambda *v: events.append(v),
    )
    assert [event[1:] for event in events if event[0] == "images"] == [(0, 2), (1, 2), (2, 2)]
    assert events[-1] == ("parsed", 1, 1)
    events = []
    built = build_index(
        [Path(result.artifact_dir)],
        tmp_path / "index",
        options=DocBaseOptions(mode="demo", incomplete_policy="native_only"),
        progress=lambda *v: events.append(v),
    )
    assert built.documents == 1
    dense = [event for event in events if event[0] == "dense"]
    assert dense[0][1] == 0 and dense[-1][1] == dense[-1][2] > 0
    assert [e[1] for e in dense] == sorted(e[1] for e in dense)
    assert events[-1] == ("published", 1, 1)


def test_frontend_batch_reconciles_ready_failure_cancel_and_real_stage_counts():
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js unavailable")
    source = Path("src/docqa/web/static/js/upload-progress.js").read_bytes()
    url = "data:text/javascript;base64," + base64.b64encode(source).decode()
    script = (
        f"import {{uploadBatchStatus,uploadProgress}} from '{url}';\n"
        + """
import assert from 'node:assert/strict';
const batch={total:2,items:new Map([['a',{state:'ready'}],['b',{state:'parsing'}]]),processed:2,transferring:false};
assert.equal(uploadBatchStatus(batch).hidden,false);
assert.equal(uploadBatchStatus(batch).value,1);
batch.items.set('b',{state:'ready_with_gaps'});
assert.equal(uploadBatchStatus(batch).hidden,true);
batch.items.set('b',{state:'failed'});
assert.equal(uploadBatchStatus(batch).hidden,false);
assert.match(uploadBatchStatus(batch).text,/失败 1/);
batch.items.set('b',{state:'cancelled'});
assert.equal(uploadBatchStatus(batch).hidden,true);
batch.transferring=true;
assert.equal(uploadBatchStatus(batch).hidden,false);
assert.match(uploadBatchStatus(batch).text,/传输/);
const counted=uploadProgress({state:'parsing',progress:{stage:'images',current:2,total:5}});
assert.equal(counted.current,2);assert.equal(counted.total,5);
assert.match(counted.label,/2\\/5/);
assert.equal(uploadProgress({state:'indexing'}).current,null);
assert.equal(uploadProgress({state:'ready'}),null);
assert.equal(uploadProgress({state:'parsing',progress:{current:500,total:5}}).current,null);
"""
    )
    subprocess.run(
        [node, "--input-type=module"],
        input=script,
        text=True,
        encoding="utf-8",
        check=True,
        capture_output=True,
        timeout=10,
    )
