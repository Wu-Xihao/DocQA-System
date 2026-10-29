"""Distinguish discovered sources from published, quality-gated index members."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from docqa.doc2md.config import ParseOptions
from docqa.docbase.config import DocBaseOptions
from docqa.domain.errors import DocQAError
from docqa.service.config import ServiceOptions
from docqa.service.pipeline import Service


@pytest.fixture
def blocked(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    picture = BytesIO()
    Image.new("RGB", (80, 60), "white").save(picture, format="PNG")
    picture.seek(0)
    document = Document()
    document.add_paragraph("A source with a figure needs image interpretation.")
    document.add_picture(picture)
    document.save(source / "source.docx")
    service = Service(
        tmp_path / "state",
        options=ServiceOptions(
            allowed_source_roots=[source],
            settle_seconds=0,
            parse=ParseOptions(ocr_mode="disabled", office_render="never"),
            retrieval=DocBaseOptions(mode="demo", recall_routes=["dense", "bm25"]),
        ),
    )
    service.add_source(source, types=["docx"])
    service.run_once()
    return service


def test_zero_published_count_explains_discovery_and_quality_blockers(blocked):
    result = blocked.index_status()
    assert result["documents"] == result["indexed_documents"] == 0
    assert result["discovered_documents"] == result["blocked_documents"] == 1
    assert result["document_states"] == {"review_required": 1}
    assert result["ingestion_status"] == "blocked"
    detail = result["blockers"][0]
    assert detail["error_code"] == "SOURCE_NOT_READY"
    assert detail["warning_counts"]["VISION_DISABLED"] == 1
    assert "source.docx" not in json.dumps(detail)
    assert "A source with a figure" not in json.dumps(detail)


def test_malformed_quality_report_does_not_break_index_status(blocked):
    document = blocked.list_documents()[0]
    path = Path(document["latest_attempt_artifact"]) / "quality_report.json"
    path.write_bytes(b"not valid JSON")
    result = blocked.index_status()
    assert result["blocked_documents"] == 1
    assert result["blockers"][0]["diagnostic_error"] == "QUALITY_REPORT_UNAVAILABLE"


def test_diagnostic_report_rejects_paths_outside_managed_parse_root(blocked, tmp_path):
    document = blocked.list_documents()[0]
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "quality_report.json").write_text(
        json.dumps({"warnings": [{"code": "MUST_NOT_READ"}]}), encoding="utf-8"
    )
    document["latest_attempt_artifact"] = str(outside)
    blocked.registry.put("documents", document)
    result = blocked.index_status()
    assert result["blockers"][0]["warning_counts"] == {}
    assert result["blockers"][0]["diagnostic_error"] == "QUALITY_REPORT_UNAVAILABLE"


def test_ready_document_and_published_count_remain_consistent(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    document = Document()
    document.add_paragraph("A fully parsed text document.")
    document.save(source / "text.docx")
    service = Service(
        tmp_path / "state",
        options=ServiceOptions(
            allowed_source_roots=[source],
            settle_seconds=0,
            retrieval=DocBaseOptions(mode="demo", recall_routes=["dense", "bm25"]),
        ),
    )
    service.add_source(source)
    service.run_once()
    result = service.index_status()
    assert result["discovered_documents"] == result["indexed_documents"] == 1
    assert result["blocked_documents"] == 0
    assert result["ingestion_status"] == "ready"


def test_failed_parse_does_not_attribute_old_quality_warnings_to_new_failure(blocked):
    document = blocked.list_documents()[0]
    document.update(state="failed", error_code="STORAGE_IO_FAILED")
    blocked.registry.put("documents", document)
    detail = blocked.index_status()["blockers"][0]
    assert detail["error_code"] == "STORAGE_IO_FAILED"
    assert detail["warning_counts"] == {}


def test_parsed_document_with_failed_build_is_not_reported_as_ready(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    document = Document()
    document.add_paragraph("A fully parsed text document awaiting model access.")
    document.save(source / "text.docx")

    def unavailable(*args, **kwargs):
        raise DocQAError("MODEL_AUTH_FAILED", "Test model is unavailable")

    service = Service(
        tmp_path / "state",
        options=ServiceOptions(
            allowed_source_roots=[source],
            settle_seconds=0,
            retrieval=DocBaseOptions(mode="demo", recall_routes=["dense", "bm25"]),
        ),
        build_function=unavailable,
    )
    service.add_source(source)
    service.run_once()
    result = service.index_status()
    assert result["document_states"] == {"ready": 1}
    assert result["blocked_documents"] == result["indexed_documents"] == 0
    assert result["ingestion_status"] == "build_failed"
    assert result["desired_revision"] > result["assembled_revision"]
