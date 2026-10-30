"""Self-contained synthetic image/VL browser acceptance; no real model or user server."""

import io
import socket
import sys
import tempfile
import threading
import time
from pathlib import Path

import uvicorn
from docx import Document
from PIL import Image
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[2]
DESCRIPTION = "IMAGE_UI_SYNTH：蓝色矩形与橙色圆形组成示意图。"
INERT_TEXT = "<script>window.__vlExecuted=true</script>"


def main():
    sys.path.insert(0, str(ROOT / "src"))
    from docqa.agent.config import AgentOptions
    from docqa.doc2md.config import ParseOptions
    from docqa.doc2md.pipeline import parse_document
    from docqa.docbase.config import DocBaseOptions
    from docqa.docbase.pipeline import build_index
    from docqa.domain.documents import VisionDescription, VisionEntity, VisionRelation
    from docqa.web.app import create_app
    from docqa.web.config import WebOptions

    class Vision:
        count = 0

        def fingerprint(self):
            return {"adapter": "synthetic-browser-vision", "synthetic": True}

        def describe(self, *args, **kwargs):
            self.count += 1
            if self.count > 1:
                return VisionDescription(status="partial", uncertainties=["小字无法识别"])
            return VisionDescription(
                status="ok",
                observations=[DESCRIPTION],
                entities=[VisionEntity(id="a", label="解析"), VisionEntity(id="b", label="检索")],
                relations=[
                    VisionRelation(source="a", target="b", direction="directed", label="下一步")
                ],
                interpretations=["仅用于界面隔离测试。", INERT_TEXT],
            )

    output = ROOT / "runtime/image-description-browser"
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="img-ui-", dir=ROOT.parent) as temporary:
        folder = Path(temporary)
        source = folder / "synthetic.docx"
        document = Document()
        document.add_paragraph("IMAGE_UI_SYNTH 文档说明：这是用于测试图片显示的合成资料。")
        for color in ("blue", "orange"):
            image = io.BytesIO()
            Image.new("RGB", (360, 220), color).save(image, format="PNG")
            image.seek(0)
            document.add_picture(image)
        document.save(source)
        parsed = parse_document(
            source,
            folder / "parsed",
            options=ParseOptions(office_render="never", ocr_mode="disabled"),
            vision_client=Vision(),
        )
        retrieval = DocBaseOptions(
            mode="demo",
            vector_backend="sqlite_vec",
            recall_routes=["dense", "bm25"],
            incomplete_policy="validated_parts",
        )
        build_index([Path(parsed.artifact_dir)], folder / "index", options=retrieval)
        listener = socket.socket()
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
        options = WebOptions(
            port=port,
            index=folder / "index",
            state=folder / "web",
            agent=AgentOptions(mode="demo", retrieval=retrieval),
        )
        server = uvicorn.Server(
            uvicorn.Config(create_app(options), host="127.0.0.1", port=port, log_level="error")
        )
        thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
        thread.start()
        try:
            deadline = time.monotonic() + 15
            while not server.started and thread.is_alive() and time.monotonic() < deadline:
                time.sleep(0.02)
            assert server.started, "Isolated test server did not start"
            base = f"http://127.0.0.1:{port}"
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(
                    executable_path=r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                    headless=True,
                )
                context = browser.new_context(viewport={"width": 1440, "height": 1000})
                context.route(
                    "**/*",
                    lambda route: (
                        route.continue_()
                        if route.request.url.startswith(base + "/")
                        else route.abort()
                    ),
                )
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(base)
                expect(page.locator(".document-name").first).to_be_visible()
                assert "缺口" not in page.locator("body").inner_text()
                page.locator(".document-name").first.click()
                page.get_by_role("tab", name="全部图片", exact=True).click()
                expect(page.locator(".reader-image-card")).to_have_count(2)
                first = page.locator(".reader-image-card").first
                expect(first.locator(".image-analysis-summary")).to_have_text(DESCRIPTION)
                first.locator("summary").click()
                expect(first.locator(".image-analysis-body")).to_contain_text(INERT_TEXT)
                assert page.evaluate("window.__vlExecuted === undefined")
                first.get_by_role("button", name="放大查看").click()
                expect(page.locator(".reader-image-dialog .image-analysis")).to_contain_text(
                    "解析 → 检索：下一步"
                )
                page.screenshot(path=str(output / "reader-desktop.png"))
                page.set_viewport_size({"width": 390, "height": 844})
                page.screenshot(path=str(output / "reader-mobile.png"))
                assert page.locator(".reader-image-dialog").evaluate(
                    "n => n.scrollWidth <= n.clientWidth + 1"
                ), page.locator(".reader-image-dialog").evaluate(
                    "n => ({client:n.clientWidth, scroll:n.scrollWidth, items:[...n.querySelectorAll('*')].filter(x=>x.scrollWidth>x.clientWidth+1).map(x=>({tag:x.tagName, class:x.className, client:x.clientWidth, scroll:x.scrollWidth}))})"
                )
                page.locator(".reader-image-dialog").get_by_role(
                    "button", name="关闭", exact=True
                ).click()
                page.locator("#close-preview").click()
                page.set_viewport_size({"width": 1440, "height": 1000})
                page.locator("#question-input").fill("IMAGE_UI_SYNTH")
                page.locator("#question-input").press("Enter")
                expect(page.locator(".answer-body .claim").first).to_be_visible(timeout=30000)
                assert "缺口" not in page.locator("body").inner_text()
                page.get_by_role("button", name="相关图片", exact=True).last.click()
                expect(page.locator("#images-list .image-card").first).to_be_visible()
                expect(page.locator("#images-list .image-analysis-summary").first).to_have_text(
                    DESCRIPTION
                )
                page.locator("#images-list .image-card > button").first.click()
                expect(page.locator("#lightbox-analysis")).to_contain_text(DESCRIPTION)
                expect(page.locator("#lightbox-analysis")).to_contain_text(INERT_TEXT)
                page.screenshot(path=str(output / "related-desktop.png"))
                page.set_viewport_size({"width": 390, "height": 844})
                assert page.locator("#image-dialog").evaluate(
                    "n => n.scrollWidth <= n.clientWidth + 1"
                )
                page.screenshot(path=str(output / "related-mobile.png"))
                assert page.evaluate("window.__vlExecuted === undefined")
                assert not errors, errors
                browser.close()
            print(
                "PASS: no coverage banners; document/related images and both enlarged previews show VL; mobile fits; VL text is inert; no browser JS errors or external requests."
            )
        finally:
            server.should_exit = True
            thread.join(timeout=10)
            listener.close()
            assert not thread.is_alive(), "Isolated test server failed to stop"


if __name__ == "__main__":
    main()
