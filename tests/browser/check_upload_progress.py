"""Isolated actual upload -> parser/image counters -> vector counters -> hidden banner."""

import io
import socket
import sys
import tempfile
import threading
from pathlib import Path

import uvicorn
import yaml
from docx import Document
from PIL import Image
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))


def main():
    from docqa.agent.config import AgentOptions
    from docqa.doc2md.config import ParseOptions
    from docqa.docbase.config import DocBaseOptions
    from docqa.service.config import ServiceOptions
    from docqa.web.app import create_app
    from docqa.web.config import WebOptions

    images, images_release, vectors, vectors_release = (threading.Event() for _ in range(4))
    output = ROOT / "runtime/upload-progress-browser"
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="upload-ui-", dir=ROOT.parent) as temporary:
        root = Path(temporary)
        retrieval = DocBaseOptions(mode="demo", incomplete_policy="native_only")
        service_config = root / "service.yaml"
        service_config.write_text(
            yaml.safe_dump(
                ServiceOptions(
                    allowed_source_roots=[root / "uploads"],
                    settle_seconds=0,
                    parse=ParseOptions(office_render="never", ocr_mode="disabled"),
                    retrieval=retrieval,
                ).model_dump(mode="json")
            ),
            encoding="utf-8",
        )
        listener = socket.socket()
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
        app = create_app(
            WebOptions(
                port=port,
                index=root / "service/index",
                state=root / "web",
                service_state=root / "service",
                service_config=service_config,
                upload_root=root / "uploads",
                service_mode="external",
                agent=AgentOptions(mode="demo", retrieval=retrieval),
            )
        )
        server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="error"))
        thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
        thread.start()
        service = app.state.workspace.service()
        original = service.queue.report_progress

        def progress(identity, owner, stage, current=None, total=None):
            original(identity, owner, stage, current, total)
            app.state.workspace.refresh_uploads()
            if stage == "images" and current == 1 and not images.is_set():
                images.set()
                assert images_release.wait(30)
            if stage == "dense" and current == 0 and not vectors.is_set():
                vectors.set()
                assert vectors_release.wait(30)

        service.queue.report_progress = progress
        document = Document()
        document.add_paragraph("UPLOAD_PROGRESS_NATIVE synthetic native body")
        for color in ("blue", "red"):
            picture = io.BytesIO()
            Image.new("RGB", (32, 20), color).save(picture, format="PNG")
            picture.seek(0)
            document.add_picture(picture)
        data = io.BytesIO()
        document.save(data)
        worker = None
        try:
            with sync_playwright() as pw:
                browser = pw.chromium.launch(
                    executable_path=r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                    headless=True,
                )
                page = browser.new_page(viewport={"width": 1440, "height": 1000})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(f"http://127.0.0.1:{port}")
                expect(page.locator("html")).to_have_attribute("data-app-state", "bound")
                page.locator("#file-input").set_input_files(
                    {
                        "name": "progress.docx",
                        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "buffer": data.getvalue(),
                    }
                )
                expect(page.locator("#folder-upload-status")).to_contain_text(
                    "解析入库", timeout=10000
                )
                expect(page.locator(".upload-row .ingestion-progress")).to_be_visible()
                worker = threading.Thread(target=service.run_once, daemon=True)
                worker.start()
                assert images.wait(15)
                expect(page.locator(".upload-stage")).to_contain_text("1/2", timeout=15000)
                expect(page.locator(".ingestion-progress")).to_have_attribute("value", "1")
                expect(page.locator(".ingestion-progress")).to_have_attribute("max", "2")
                page.screenshot(path=str(output / "parsing.png"))
                images_release.set()
                assert vectors.wait(15)
                expect(page.locator(".upload-stage")).to_contain_text("向量化", timeout=15000)
                page.screenshot(path=str(output / "indexing.png"))
                vectors_release.set()
                worker.join(20)
                assert not worker.is_alive()
                app.state.workspace.refresh_uploads()
                expect(page.locator("#document-count")).to_have_text("1", timeout=15000)
                expect(page.locator("#folder-upload-status")).to_be_hidden(timeout=15000)
                expect(page.locator("#folder-upload-status")).to_have_text("")
                expect(page.locator("#folder-upload-progress")).to_be_hidden()
                expect(page.locator(".upload-row")).to_have_count(0)
                page.screenshot(path=str(output / "ready.png"))
                assert not errors, errors
                browser.close()
            print(
                "PASS: real upload; live image/vector counters; published library; completed banner and progress auto-hide; no external model requests."
            )
        finally:
            images_release.set()
            vectors_release.set()
            if worker:
                worker.join(15)
            server.should_exit = True
            thread.join(10)
            listener.close()


if __name__ == "__main__":
    main()
