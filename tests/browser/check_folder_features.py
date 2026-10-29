"""Isolated demo browser acceptance for folder ingestion, scopes and study views."""

import json
import sys
import time
from pathlib import Path

from docx import Document
from playwright.sync_api import expect, sync_playwright
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(root / "src"))
sys.path.insert(0, str(root / "tests"))
from test_pdf_parser import _write_pdf  # noqa: E402

base = "http://127.0.0.1:8789"
output = root / "runtime/folder-features-browser"
output.mkdir(exist_ok=True)
folder = output / ("batch-" + str(time.time_ns()))
(folder / "nested").mkdir(parents=True)
for i in range(23):
    doc = Document()
    doc.add_paragraph(f"FOLDER_{i} synthetic document.")
    doc.save(folder / "nested" / f"paper-{i}.docx")
_write_pdf(
    folder / "scope.PDF", [b"BT /F1 12 Tf 20 250 Td (SCOPE_PDF synthetic PDF document.) Tj ET"]
)
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[6])
slide.shapes.add_textbox(
    Inches(1), Inches(1), Inches(7), Inches(2)
).text = "SCOPE_PPTX synthetic slide."
slides.save(folder / "scope.pptx")
(folder / "ignored.txt").write_text("Not an upload type", encoding="utf-8")
checks, errors = [], []
with sync_playwright() as p:
    browser = p.chromium.launch(
        executable_path=r"C:\Program Files\Google\Chrome\Application\chrome.exe", headless=True
    )
    context = browser.new_context(viewport={"width": 1680, "height": 1050})
    context.route(
        "**/*",
        lambda route: route.continue_() if route.request.url.startswith(base) else route.abort(),
    )
    page = context.new_page()
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(base)
    boot = context.request.get(base + "/api/v1/bootstrap").json()
    assert boot["settings"]["synthetic"]
    expect(page.locator("html")).to_have_attribute("data-app-state", "bound")
    expect(page.locator(".document-name").first).to_be_visible()
    page.locator(".document-name").filter(has_text="Synthetic research.docx").first.click()
    page.get_by_role("tab", name="全部图片", exact=True).click()
    expect(page.locator(".reader-image-card")).to_have_count(2)
    expect(page.locator(".reader-image-card img").first).to_be_visible()
    for _ in range(50):
        if page.locator(".reader-image-card img").first.evaluate("img=>img.naturalWidth>0"):
            break
        page.wait_for_timeout(100)
    assert page.locator(".reader-image-card img").first.evaluate("img=>img.naturalWidth>0")
    page.locator(".reader-image-card").first.get_by_role("button", name="放大查看").click()
    expect(page.locator(".reader-image-dialog")).to_be_visible()
    page.locator(".reader-image-dialog").get_by_role("button", name="关闭", exact=True).click()
    page.locator("#close-preview").click()
    checks.append("all extracted images tab, lazy thumbnails and enlarged preview")

    # Three explicit topics give the demo a multi-branch checked answer without a model call.
    token = boot["csrf_token"]
    headers = {"X-DocQA-Token": token}
    cid = context.request.post(base + "/api/v1/conversations", headers=headers, data={}).json()[
        "id"
    ]
    job = context.request.post(
        base + f"/api/v1/conversations/{cid}/messages",
        headers=headers,
        data={
            "question": "DOCQA_IMAGE / SD5993 / RECOVERY",
            "topics": ["DOCQA_IMAGE", "SD5993", "RECOVERY"],
            "client_request_id": "browser-studio-" + str(time.time_ns()),
            "doc_types": ["docx"],
        },
    ).json()
    until = time.monotonic() + 35
    while time.monotonic() < until:
        result = context.request.get(base + "/api/v1/jobs/" + job["job_id"], headers=headers).json()
        if result["state"] == "succeeded":
            break
        assert result["state"] not in {"failed", "cancelled"}, result
        page.wait_for_timeout(200)
    assert result["state"] == "succeeded" and result["answer"]["claims"]
    page.reload()
    expect(page.locator(".history-open").first).to_be_visible(timeout=20000)
    page.locator(".history-open").filter(has_text="DOCQA_IMAGE").first.click()
    expect(page.locator(".message-actions")).to_have_count(1, timeout=20000)
    if not page.locator("#right-panel").is_visible():
        page.locator("#toggle-right").click()
    page.locator(".message-actions").get_by_role("button", name="生成成果").click()
    page.locator(".studio-action").filter(has_text="思维导图").click()
    expect(page.locator(".mindmap .graph-node")).to_have_count(7, timeout=30000)
    paths = page.locator(".mindmap svg path[d]").evaluate_all(
        "nodes=>nodes.filter(n=>n.getAttribute('d').includes(' C')).length"
    )
    assert paths == 6
    page.locator(".mindmap .graph-tools").get_by_role("button", name="＋", exact=True).click()
    page.locator(".mindmap .graph-tools").get_by_role("button", name="适应画布").click()
    page.get_by_role("button", name="全屏画布", exact=True).click()
    expect(page.get_by_role("button", name="退出全屏", exact=True)).to_be_visible()
    page.get_by_role("button", name="退出全屏", exact=True).click()
    expect(page.get_by_role("button", name="全屏画布", exact=True)).to_be_visible()
    page.locator(".mindmap .graph-node").last.click()
    expect(page.locator(".map-selected")).to_contain_text("RECOVERY")
    page.screenshot(path=str(output / "mindmap.png"))
    page.locator(".map-outline > summary").click()
    page.get_by_role("button", name="折叠全部", exact=True).click()
    assert not page.locator(".map-branch").first.evaluate("el=>el.open")
    page.locator("#close-preview").click()
    checks.append("bilateral mindmap, curved branches, zoom, full node content and citations")
    page.locator(".studio-action").filter(has_text="演示文稿").click()
    expect(page.locator('.slide-page[data-layout="cover"]')).to_be_visible(timeout=30000)
    page.screenshot(path=str(output / "slides-cover.png"))
    page.get_by_role("button", name="下一页 ›", exact=True).click()
    expect(page.locator('.slide-page[data-layout="agenda"]')).to_be_visible()
    page.get_by_role("button", name="下一页 ›", exact=True).click()
    expect(page.locator(".slide-card").first).to_be_visible()
    expect(page.locator(".slide-card .citation-chip").first).to_be_visible()
    page.screenshot(path=str(output / "slides-content.png"))
    page.locator("#close-preview").click()
    checks.append("slide cover, agenda, structured cited cards and source appendix")

    page.locator("#new-chat").click()
    before_ids = {
        d["doc_id"]
        for d in context.request.get(base + "/api/v1/documents", headers=headers).json()[
            "documents"
        ]
    }
    with page.expect_file_chooser() as chosen:
        page.locator("#folder-upload").click()
    chosen.value.set_files(str(folder))
    expect(page.locator("#folder-upload-status")).to_contain_text("成功 25 · 失败 0", timeout=60000)
    # Parsing/index publication is separate from transfer progress.
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        docs = context.request.get(base + "/api/v1/documents", headers=headers).json()
        if len([d for d in docs["documents"] if d["doc_id"] not in before_ids]) == 25:
            break
        page.wait_for_timeout(400)
    new_docs = [d for d in docs["documents"] if d["doc_id"] not in before_ids]
    assert len(new_docs) == 25 and len([d for d in new_docs if "paper-" in d["source_name"]]) == 23
    assert not any(d["source_name"].endswith("ignored.txt") for d in docs["documents"])
    checks.append(
        "real browser folder picker recursively transfers 25 supported documents and ignores TXT"
    )
    page.locator("#types-all").uncheck()
    expect(page.locator("#ask-button")).to_be_disabled()
    page.locator('[data-doc-type="pdf"]').check()
    expect(page.locator("#documents-list .document-row")).to_have_count(
        sum(d["source_name"].lower().endswith(".pdf") for d in docs["documents"]), timeout=20000
    )
    page.locator("#question-input").fill("SCOPE_PDF")
    page.locator("#question-input").press("Enter")
    expect(page.locator(".message-actions")).to_have_count(1, timeout=30000)
    expect(page.locator(".answer-document").first).to_contain_text("scope.pdf")
    page.locator('[data-doc-type="pptx"]').check()
    expect(page.locator("#documents-list .document-row")).to_have_count(
        sum(d["source_name"].lower().endswith((".pdf", ".pptx")) for d in docs["documents"])
    )
    checks.append("one/two format selection, empty selection disables send, PDF answer cites PDF")
    page.set_viewport_size({"width": 390, "height": 844})
    page.locator("#toggle-left").click()
    expect(page.locator("#folder-upload")).to_be_visible()
    page.screenshot(path=str(output / "mobile-types.png"))
    assert not errors, errors
    browser.close()
report = {"status": "passed", "checks": checks, "browser_errors": errors, "real_model_calls": False}
(output / "report.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(report, ensure_ascii=True))
