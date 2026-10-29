from __future__ import annotations

import json
import shutil
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image, PngImagePlugin
from pydantic import ValidationError
from typer.testing import CliRunner

from docqa.cli.main import app
from docqa.doc2md.common import ParseContext
from docqa.doc2md.config import ParseOptions, VisionConfig, load_options
from docqa.doc2md.pipeline import batch_parse, parse_document, verify_artifact
from docqa.doc2md.security import inspect_snapshot, io_path
from docqa.domain.documents import VisionDescription, VisionEntity
from docqa.domain.errors import DocQAError


def write_docx(path: Path, *, image=False, text="中文😀 SOURCE-TEXT-001"):
    document = Document()
    document.add_heading("文档标题", 1)
    document.add_paragraph(text)
    if image:
        buffer = BytesIO()
        Image.new("RGB", (60, 40), "red").save(buffer, format="PNG")
        buffer.seek(0)
        document.add_picture(buffer)
    document.save(path)
    return path


def options(**overrides):
    return ParseOptions(ocr_mode="disabled", office_render="never", **overrides)


def test_text_roundtrip_idempotency_and_version(tmp_path):
    source = write_docx(tmp_path / "原文.docx")
    root = tmp_path / "out"
    first = parse_document(source, root, options=options())
    assert first.status == "complete" and first.eligible_for_indexing
    assert "中文😀 SOURCE" in Path(first.markdown_path).read_text(encoding="utf-8")
    assert verify_artifact(first.artifact_dir)["valid"]
    same = parse_document(source, root, options=options())
    assert same.reused and same.artifact_id == first.artifact_id
    write_docx(source, text="SECOND-VERSION")
    changed = parse_document(source, root, options=options())
    assert changed.doc_id == first.doc_id
    assert changed.revision_id != first.revision_id
    assert verify_artifact(first.artifact_dir)["valid"]
    assert verify_artifact(changed.artifact_dir)["valid"]


def test_same_bytes_different_paths_are_separate_documents(tmp_path):
    one = write_docx(tmp_path / "one.docx")
    two = tmp_path / "two.docx"
    two.write_bytes(one.read_bytes())
    a = parse_document(one, tmp_path / "out", options=options())
    b = parse_document(two, tmp_path / "out", options=options())
    assert a.doc_id != b.doc_id
    assert a.artifact_id != b.artifact_id


def test_image_without_vision_is_explicit_review(tmp_path):
    result = parse_document(
        write_docx(tmp_path / "image.docx", image=True), tmp_path / "out", options=options()
    )
    assert result.status == "review_required"
    assert not result.eligible_for_indexing
    assert any(warning.code == "VISION_DISABLED" for warning in result.warnings)
    assert verify_artifact(result.artifact_dir)["valid"]


def test_verify_relocated_artifact_with_paths_beyond_windows_max_path(tmp_path):
    parsed = parse_document(
        write_docx(tmp_path / "long-path.docx", image=True),
        tmp_path / "short",
        options=options(),
    )
    relocated = tmp_path / "relocated"
    while len(str(relocated)) < 270:
        relocated /= "long-artifact-path"
    shutil.copytree(tmp_path / "short", io_path(relocated))
    directory = relocated / "artifacts" / parsed.artifact_id
    assert len(str(directory / "manifest.json")) > 260
    assert verify_artifact(directory)["valid"]
    ir = json.loads(io_path(directory / "document.ir.json").read_bytes())
    asset = directory / ir["assets"][0]["relative_path"]
    io_path(asset).write_bytes(b"tampered")
    with pytest.raises(DocQAError) as failure:
        verify_artifact(directory)
    assert failure.value.code == "ARTIFACT_CORRUPT"


def test_injected_vision_is_recorded_and_cached(tmp_path):
    class FakeVision:
        calls = 0

        def fingerprint(self):
            return {"adapter": "unit-test-only", "revision": "1"}

        def describe(self, image, *, ocr_text="", context=""):
            self.calls += 1
            assert Path(image).is_file()
            return VisionDescription(status="ok", observations=["TEST DOUBLE ONLY"])

    fake = FakeVision()
    source = write_docx(tmp_path / "image.docx", image=True)
    result = parse_document(source, tmp_path / "out", options=options(), vision_client=fake)
    assert result.status == "complete"
    assert fake.calls == 1
    assert "TEST DOUBLE ONLY" in Path(result.markdown_path).read_text(encoding="utf-8")
    assert parse_document(source, tmp_path / "out", options=options(), vision_client=fake).reused
    assert fake.calls == 1


def test_batch_continues_after_corrupt_file(tmp_path):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    write_docx(inputs / "good.docx")
    (inputs / "bad.pdf").write_text("not a pdf", encoding="utf-8")
    (inputs / "ignored.txt").write_text("ignored", encoding="utf-8")
    result = batch_parse(inputs, tmp_path / "out", options=options())
    assert (result["attempted"], result["complete"], result["failed"]) == (2, 1, 1)
    assert result["failures"][0]["error"]["code"] == "FORMAT_MISMATCH"


def test_nested_batch_output_rejected(tmp_path):
    with pytest.raises(DocQAError, match="must not contain"):
        batch_parse(tmp_path, tmp_path / "out")


def test_allowed_root_blocks_other_directory(tmp_path):
    source = write_docx(tmp_path / "input.docx")
    allowed = tmp_path / "allowed"
    allowed.mkdir()
    with pytest.raises(DocQAError) as caught:
        parse_document(source, tmp_path / "out", options=options(allowed_roots=[allowed]))
    assert caught.value.code == "SOURCE_ACCESS_DENIED"


def test_tampered_cached_asset_fails_without_overwrite(tmp_path):
    source = write_docx(tmp_path / "input.docx")
    result = parse_document(source, tmp_path / "out", options=options())
    Path(result.markdown_path).write_text("tampered", encoding="utf-8")
    with pytest.raises(DocQAError) as caught:
        parse_document(source, tmp_path / "out", options=options())
    assert caught.value.code == "ARTIFACT_CORRUPT"
    assert Path(result.markdown_path).read_text(encoding="utf-8") == "tampered"


@pytest.mark.parametrize("payload", ["false", "[]", "42", "vision:\n  mode: http\n"])
def test_invalid_config_not_silently_defaulted(tmp_path, payload):
    path = tmp_path / "bad.yaml"
    path.write_text(payload, encoding="utf-8")
    with pytest.raises(DocQAError) as caught:
        load_options(path)
    assert caught.value.code == "INVALID_CONFIG"


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-32"])
def test_dtd_rejected_for_multibyte_xml(tmp_path, encoding):
    path = tmp_path / "unsafe.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(
            "word/document.xml",
            (
                "<!DOCTYPE document [<!ENTITY payload 'injected'>]><document>&payload;</document>"
            ).encode(encoding),
        )
    with pytest.raises(DocQAError) as caught:
        inspect_snapshot(path, "docx", ParseOptions())
    assert caught.value.code == "UNSAFE_XML"


def test_raw_image_occurrences_survive_pixel_deduplication(tmp_path):
    context = ParseContext(
        tmp_path / "source.docx", tmp_path / "work", "doc", "revision", "source.docx"
    )

    def encoded(tag):
        stream = BytesIO()
        pnginfo = PngImagePlugin.PngInfo()
        pnginfo.add_text("tag", tag)
        Image.new("RGB", (20, 20), "blue").save(stream, format="PNG", pnginfo=pnginfo)
        return stream.getvalue()

    a = context.save_image(encoded("first"), "a")
    b = context.save_image(encoded("second"), "b")
    assert a.asset_id == b.asset_id
    assert a.derived_from != b.derived_from
    assert len(context.raw_assets) == 2


def test_inverted_vision_box_rejected():
    with pytest.raises(ValidationError):
        VisionEntity(id="node", label="node", bbox=(0.8, 0.0, 0.2, 1.0))
    with pytest.raises(ValidationError):
        VisionConfig(mode="http")


def test_cli_parse_verify_and_strict_exit_codes(tmp_path):
    runner = CliRunner()
    source = write_docx(tmp_path / "input.docx")
    config = tmp_path / "config.yaml"
    config.write_text("office_render: never\nocr_mode: disabled\n", encoding="utf-8")
    command = [
        "doc2md",
        "parse",
        str(source),
        "--output",
        str(tmp_path / "out"),
        "--config",
        str(config),
        "--strict",
    ]
    result = runner.invoke(app, command)
    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    verified = runner.invoke(app, ["doc2md", "verify", payload["artifact_dir"]])
    assert verified.exit_code == 0 and json.loads(verified.stdout)["valid"]
    image = write_docx(tmp_path / "image.docx", image=True)
    command[2] = str(image)
    reviewed = runner.invoke(app, command)
    assert reviewed.exit_code == 5, reviewed.output
    assert json.loads(reviewed.stdout)["status"] == "review_required"
    invalid = runner.invoke(app, ["doc2md", "parse", str(tmp_path / "missing.pdf")])
    assert invalid.exit_code == 4


def test_snapshot_io_failure_is_classified(tmp_path):
    source = write_docx(tmp_path / "source.docx")
    with patch("docqa.doc2md.pipeline._snapshot", side_effect=OSError("test disk failure")):
        with pytest.raises(DocQAError) as caught:
            parse_document(source, tmp_path / "out", options=options())
    assert caught.value.code == "STORAGE_IO_FAILED"


def test_document_markdown_syntax_is_escaped(tmp_path):
    source = write_docx(
        tmp_path / "source.docx", text="~~~\n<script>alert(1)</script>\n---\n1. item\n- item"
    )
    result = parse_document(source, tmp_path / "out", options=options())
    markdown = Path(result.markdown_path).read_text(encoding="utf-8")
    assert "<script>" not in markdown
    assert "\\~\\~\\~" in markdown
    assert "1\\. item" in markdown
    assert "\\- item" in markdown


@pytest.mark.parametrize("payload", ["[]", "null", "true", '"string"'])
def test_manifest_root_type_returns_classified_error(tmp_path, payload):
    source = write_docx(tmp_path / "source.docx")
    result = parse_document(source, tmp_path / "out", options=options())
    (Path(result.artifact_dir) / "manifest.json").write_text(payload, encoding="utf-8")
    with pytest.raises(DocQAError) as caught:
        verify_artifact(result.artifact_dir)
    assert caught.value.code == "ARTIFACT_CORRUPT"


def test_failed_vision_can_recover_without_changing_credentials_profile(tmp_path):
    class RecoveringVision:
        calls = 0

        def fingerprint(self):
            return {"adapter": "unit-test-recovery", "revision": "1"}

        def describe(self, image, *, ocr_text="", context=""):
            self.calls += 1
            if self.calls == 1:
                raise DocQAError("VISION_AUTH_FAILED", "Simulated missing credentials")
            return VisionDescription(status="ok", observations=["TEST DOUBLE ONLY"])

    vision = RecoveringVision()
    source = write_docx(tmp_path / "source.docx", image=True)
    first = parse_document(source, tmp_path / "out", options=options(), vision_client=vision)
    assert first.status == "partial"
    second = parse_document(source, tmp_path / "out", options=options(), vision_client=vision)
    assert second.status == "complete" and not second.reused
    assert second.artifact_id != first.artifact_id
    assert verify_artifact(first.artifact_dir)["valid"]
    assert parse_document(source, tmp_path / "out", options=options(), vision_client=vision).reused
    assert vision.calls == 2


def test_cli_invalid_adapter_configuration_is_json_error(tmp_path):
    source = write_docx(tmp_path / "source.docx")
    config = tmp_path / "bad.yaml"
    config.write_text("ocr_languages: 'eng,--help'", encoding="utf-8")
    response = CliRunner().invoke(
        app,
        [
            "doc2md",
            "parse",
            str(source),
            "--config",
            str(config),
            "--output",
            str(tmp_path / "out"),
        ],
    )
    assert response.exit_code == 2
    assert json.loads(response.stdout)["error"]["code"] == "INVALID_CONFIG"
