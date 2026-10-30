"""Isolated diagram fixtures: no business files or external model calls."""

import threading
import time
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from test_pdf_parser import _ctx, _write_pdf

from docqa.doc2md.common import ParseContext
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.parsers.pdf import parse_pdf
from docqa.doc2md.pipeline import _add_office_render, _enrich, parse_document, verify_artifact
from docqa.domain.documents import DocumentIR, VisionDescription
from docqa.domain.errors import DocQAError


class Vision:
    def __init__(self):
        self.calls = 0
        self.active = 0
        self.peak = 0
        self.lock = threading.Lock()

    def fingerprint(self):
        return {"adapter": "isolated-regions-test", "synthetic": True}

    def describe(self, image, *, ocr_text="", context=""):
        assert Path(image).is_file()
        with self.lock:
            self.calls += 1
            self.active += 1
            self.peak = max(self.peak, self.active)
        time.sleep(0.05)
        with self.lock:
            self.active -= 1
        return VisionDescription(status="ok", observations=["Synthetic diagram observation"])


class OCR:
    def __init__(self):
        self.calls = 0

    def recognize(self, image):
        self.calls += 1
        return "Synthetic scanned text"

    def fingerprint(self):
        return {"adapter": "isolated-ocr-test"}


def test_multi_page_paper_has_two_cropped_diagrams_not_page_images(tmp_path):
    path = tmp_path / "paper.pdf"
    prose = b"BT /F1 10 Tf 30 270 Td (Synthetic paper native text) Tj ET "
    diagram = (
        b"30 140 60 35 re S 140 140 60 35 re S 90 157 m 140 157 l S "
        b"BT /F1 10 Tf 30 120 Td (Figure 1. Synthetic flow) Tj ET "
    )
    _write_pdf(path, [prose, prose + diagram, prose, prose + diagram, prose])
    vision = Vision()
    parsed = parse_document(
        path,
        tmp_path / "out",
        options=ParseOptions(ocr_mode="disabled", image_workers=2),
        vision_client=vision,
    )
    assert verify_artifact(parsed.artifact_dir)["valid"]
    ir = DocumentIR.model_validate_json(Path(parsed.artifact_dir, "document.ir.json").read_bytes())
    assert (
        len(ir.figures) == 2 and vision.calls == 1
    )  # Identical pixels + local context share cache.
    assert all(not f.full_page and f.locator.page in {2, 4} for f in ir.figures)
    assert all("Figure 1" in f.caption for f in ir.figures)
    assert all(a.width < 600 and a.height < 300 for a in ir.assets)
    assert "Synthetic diagram observation" in Path(parsed.markdown_path).read_text(encoding="utf-8")


def test_scan_is_ocr_source_without_full_page_vl_call(tmp_path):
    source = tmp_path / "scan.pdf"
    image = Image.new("RGB", (300, 400), "white")
    ImageDraw.Draw(image).text((40, 50), "Synthetic scan", fill="black")
    image.save(source, "PDF")
    ctx = _ctx(tmp_path, source)
    content = parse_pdf(ctx)
    vision, ocr = Vision(), OCR()
    _enrich(content, ctx, ParseOptions(), ocr, vision, tmp_path)
    assert vision.calls == 0 and ocr.calls == 1
    assert content.figures[0].metadata["exclude_gallery"]
    assert any(b.origin == "ocr" and b.text == "Synthetic scanned text" for b in content.blocks)
    assert any(w.code == "SCANNED_FIGURES_NEED_REVIEW" for w in content.warnings)
    assert not any(w.code == "VISION_DISABLED" for w in content.warnings)


def test_ocr_outage_does_not_repeat_timeout_for_all_images(tmp_path):
    source = tmp_path / "ocr-outage.pdf"
    _write_pdf(source, [f"30 100 {60 + i * 10} 40 re S".encode() for i in range(5)])
    ctx = _ctx(tmp_path, source)
    content = parse_pdf(ctx)

    class Down(OCR):
        def recognize(self, image):
            self.calls += 1
            raise DocQAError("OCR_TIMEOUT", "Isolated OCR timeout")

    ocr, vision = Down(), Vision()
    _enrich(content, ctx, ParseOptions(), ocr, vision, tmp_path)
    assert ocr.calls == 2 and vision.calls == 5
    assert sum(w.code == "OCR_BATCH_PAUSED" for w in content.warnings) == 3


def test_full_office_pipeline_preserves_native_picture_when_render_mapping_is_unknown(tmp_path):
    from io import BytesIO

    from docx import Document

    source = tmp_path / "office.docx"
    document = Document()
    document.add_paragraph("Synthetic native text")
    pixels = BytesIO()
    Image.new("RGB", (60, 40), "blue").save(pixels, format="PNG")
    pixels.seek(0)
    document.add_picture(pixels)
    document.save(source)
    rendered = tmp_path / "render.pdf"
    _write_pdf(rendered, [b"30 100 60 40 re S"])

    class Renderer:
        def fingerprint(self):
            return {"adapter": "isolated-office-renderer"}

        def render(self, source, output):
            output.mkdir(parents=True)
            target = output / "converted.pdf"
            target.write_bytes(rendered.read_bytes())
            return target

    parsed = parse_document(
        source,
        tmp_path / "out",
        options=ParseOptions(office_render="required", ocr_mode="disabled"),
        office_renderer=Renderer(),
        vision_client=Vision(),
    )
    assert verify_artifact(parsed.artifact_dir)["valid"]
    ir = DocumentIR.model_validate_json(Path(parsed.artifact_dir, "document.ir.json").read_bytes())
    assert len(ir.figures) == 2 and all(f.analysis_required for f in ir.figures)
    assert any(f.locator.locator_type == "derived_page" for f in ir.figures)
    assert any(f.locator.locator_type == "docx_object" for f in ir.figures)


@pytest.mark.parametrize("source_name", ["original.docx", "original.pptx"])
def test_office_render_imports_regions_and_keeps_derived_provenance(tmp_path, source_name):
    from docqa.domain.documents import ParsedContent

    rendered = tmp_path / "render.pdf"
    _write_pdf(
        rendered,
        [
            b"BT /F1 12 Tf 20 260 Td (Plain text) Tj ET",
            b"30 100 60 40 re S 130 100 60 40 re S 90 120 m 130 120 l S",
        ],
    )

    class Renderer:
        def render(self, source, output):
            output.mkdir(parents=True)
            target = output / "converted.pdf"
            target.write_bytes(rendered.read_bytes())
            return target

    ctx = ParseContext(
        tmp_path / source_name, tmp_path / "artifact", "doc", "revision", source_name
    )
    content = ParsedContent(metadata={"needs_office_render": True})
    _add_office_render(content, ctx, Renderer(), ParseOptions(), {})
    assert len(content.figures) == 1
    f = content.figures[0]
    assert not f.full_page and f.locator.page == 2 and f.locator.bbox != (0, 0, 1, 1)
    assert f.locator.locator_type == "derived_page"
    assert f.locator.native_ref["derived_pdf"] == "renders/document.pdf"
    assert source_name in f.locator.display_label


def test_region_analysis_is_bounded_parallel_and_progress_is_monotonic(tmp_path):
    source = tmp_path / "separate.pdf"
    _write_pdf(source, [b"30 100 60 40 re S", b"30 100 80 40 re S", b"30 100 100 40 re S"])
    ctx = _ctx(tmp_path, source)
    content = parse_pdf(ctx)
    progress = []
    vision = Vision()
    _enrich(
        content,
        ctx,
        ParseOptions(ocr_mode="disabled", image_workers=2),
        OCR(),
        vision,
        tmp_path,
        progress=lambda *args: progress.append(args),
    )
    assert vision.calls == 3 and vision.peak == 2
    assert [r[1] for r in progress] == [0, 1, 2, 3]
    assert [b.order for b in content.blocks] == list(range(len(content.blocks)))


@pytest.mark.parametrize(
    "code,expected",
    [("VISION_TRANSPORT_FAILED", 2), ("VISION_AUTH_FAILED", 1), ("VISION_SCHEMA_INVALID", 3)],
)
def test_failed_service_does_not_retry_every_region_of_a_long_document(tmp_path, code, expected):
    source = tmp_path / "outage.pdf"
    _write_pdf(source, [f"30 100 {60 + i * 10} 40 re S".encode() for i in range(6)])
    ctx = _ctx(tmp_path, source)
    content = parse_pdf(ctx)

    class Down(Vision):
        def describe(self, *args, **kwargs):
            self.calls += 1
            raise DocQAError(code, "Isolated service failure")

    vision = Down()
    _enrich(content, ctx, ParseOptions(ocr_mode="disabled"), OCR(), vision, tmp_path)
    assert vision.calls == expected
    assert all(f.semantic_state == "failed" and f.vision is None for f in content.figures)
    assert sum(w.code == "VISION_BATCH_PAUSED" for w in content.warnings) == 6 - expected


def test_regions_gallery_does_not_list_ocr_pages_or_replaced_archives(tmp_path):
    from docqa.domain.documents import FigureOccurrence
    from docqa.web import documents

    class IR:
        source_name = "synthetic.pdf"
        figures = [
            FigureOccurrence(
                figure_id="scan",
                asset_id="image",
                full_page=True,
                locator={
                    "locator_type": "pdf_page",
                    "page": 1,
                    "precision": "page_only",
                    "display_label": "scan",
                },
            ),
            FigureOccurrence(
                figure_id="archive",
                asset_id="image",
                metadata={"exclude_gallery": True},
                locator={
                    "locator_type": "pdf_page",
                    "page": 1,
                    "precision": "page_only",
                    "display_label": "archive",
                },
            ),
        ]

    from unittest.mock import patch

    with patch.object(documents, "_image_ir", return_value=IR()):
        result = documents.document_images(tmp_path, "/images")
    assert result["total"] == 0 and not result["items"]
