"""Multi-condition retrieval filters and the direct snippet search endpoint.

The corpus mixes DOCX, XLSX and CSV so the tests also prove spreadsheets go
through parsing, chunking and the demo hybrid index end to end. No real models.
"""

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from docqa.agent.config import AgentOptions
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.filters import ChunkFilter, coerce_filter
from docqa.docbase.pipeline import build_index
from docqa.docbase.retrieval import query
from docqa.docbase.storage import load_generation
from docqa.domain.errors import DocQAError
from docqa.web.app import create_app
from docqa.web.config import WebOptions


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    root = tmp_path_factory.mktemp("mixed")
    sources = root / "sources"
    sources.mkdir()
    word = sources / "协议说明.docx"
    document = Document()
    document.add_heading("心跳机制", 1)
    document.add_paragraph("SD5993 心跳启动：调用 FE_HEART_Start 后发送心跳报文。")
    document.add_heading("证书校验", 1)
    document.add_paragraph("TLS 连接需要校验证书链与有效期。")
    document.save(word)
    sheet = sources / "型号参数.xlsx"
    book = Workbook()
    book.active.title = "参数"
    book.active.append(["型号", "工作电压", "心跳周期"])
    book.active.append(["SD5993", "3.3V", "30s"])
    book.active.append(["SD6001", "5V", "60s"])
    book.save(sheet)
    table = sources / "备件清单.csv"
    table.write_text("名称,数量\n心跳模块,4\n证书卡,2\n", encoding="utf-8")
    artifacts = [
        Path(
            parse_document(
                path, root / "parsed", options=ParseOptions(allowed_roots=[sources])
            ).artifact_dir
        )
        for path in (word, sheet, table)
    ]
    retrieval = DocBaseOptions(mode="demo")
    build_index(artifacts, root / "kb", options=retrieval)
    return root / "kb", retrieval


def chunks(index):
    from docqa.docbase.retrieval import generation_directory

    with load_generation(generation_directory(index)) as generation:
        return list(generation.chunks)


def test_spreadsheets_are_indexed_with_sheet_locations(corpus):
    index, _ = corpus
    names = {chunk.source_name for chunk in chunks(index)}
    assert names == {"协议说明.docx", "型号参数.xlsx", "备件清单.csv"}
    sheet = [c for c in chunks(index) if c.source_name == "型号参数.xlsx"]
    labels = {loc.display_label for c in sheet for s in c.source_spans for loc in s.locators}
    assert any("工作表 参数 A2:C3" in label for label in labels)
    assert any("SD5993" in c.display_text and "3.3V" in c.display_text for c in sheet)


def test_filter_conditions_match_chunk_metadata(corpus):
    index, _ = corpus
    rows = chunks(index)

    def names(**values):
        selected = ChunkFilter(**values)
        return {c.source_name for c in rows if selected.matches(c)}

    assert names(doc_types=["xlsx", "csv"]) == {"型号参数.xlsx", "备件清单.csv"}
    assert names(content_types=["table"]) == {"型号参数.xlsx", "备件清单.csv"}
    # Sheet chunks also carry their worksheet heading as a text span.
    assert "协议说明.docx" in names(content_types=["text"])
    assert names(must_contain=["sd5993", "3.3v"]) == {"型号参数.xlsx"}
    assert names(source_name_contains="清单") == {"备件清单.csv"}
    assert names(section_contains="证书") == {"协议说明.docx"}
    assert names(origins=["ocr"]) == set()
    assert names(page_from=1) == set()  # DOCX/XLSX/CSV carry no physical pages.
    assert coerce_filter({}) is None
    with pytest.raises(ValueError):
        ChunkFilter(doc_types=["exe"])
    with pytest.raises(ValueError):
        ChunkFilter(page_from=5, page_to=2)
    with pytest.raises(ValueError):
        ChunkFilter(unknown=True)


def test_query_applies_conditions_before_recall(corpus):
    index, options = corpus
    result = query(
        index,
        "SD5993 心跳",
        options=options,
        filters={"doc_types": ["xlsx"], "must_contain": ["SD5993"]},
    )
    assert result.hits and {hit.source_name for hit in result.hits} == {"型号参数.xlsx"}
    assert result.applied_filter["conditions"] == {
        "doc_types": ["xlsx"],
        "must_contain": ["SD5993"],
    }
    empty = query(index, "SD5993", options=options, filters={"source_name_contains": "不存在"})
    assert empty.status == "empty" and empty.reason == "no_eligible_documents"
    plain = query(index, "SD5993 心跳", options=options)
    assert "conditions" not in plain.applied_filter
    with pytest.raises(DocQAError) as error:
        query(index, "x", options=options, filters={"doc_types": ["exe"]})
    assert error.value.code == "INVALID_REQUEST"


def test_search_endpoint_returns_located_snippets_and_links(corpus, tmp_path):
    index, retrieval = corpus
    options = WebOptions(
        index=index, state=tmp_path, agent=AgentOptions(mode="demo", retrieval=retrieval)
    )
    with TestClient(create_app(options), base_url="http://127.0.0.1:8765") as client:
        token = client.get("/api/v1/bootstrap").json()["csrf_token"]
        client.headers["X-DocQA-Token"] = token
        response = client.post(
            "/api/v1/search",
            json={"question": "心跳周期", "top_k": 5, "filters": {"content_types": ["table"]}},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == "ok" and body["elapsed_ms"] >= 0
        assert body["hits"] and all(
            hit["source_name"] in {"型号参数.xlsx", "备件清单.csv"} for hit in body["hits"]
        )
        first = body["hits"][0]
        assert first["rank"] == 1 and first["locations"] and first["block_ids"]
        markdown = client.get(first["links"]["markdown"])
        assert markdown.status_code == 200
        original = client.get(first["links"]["original"])
        assert original.status_code == 200 and original.content
        bad = client.post("/api/v1/search", json={"question": "x", "filters": {"x": 1}})
        assert bad.status_code in {400, 422}
        too_many = client.post("/api/v1/search", json={"question": "x", "top_k": 99})
        assert too_many.status_code in {400, 422}


def test_stats_endpoint_summarizes_metadata_only(corpus, tmp_path):
    index, retrieval = corpus
    options = WebOptions(
        index=index, state=tmp_path, agent=AgentOptions(mode="demo", retrieval=retrieval)
    )
    with TestClient(create_app(options), base_url="http://127.0.0.1:8765") as client:
        client.headers["X-DocQA-Token"] = client.get("/api/v1/bootstrap").json()["csrf_token"]
        assert client.post("/api/v1/search", json={"question": "心跳"}).status_code == 200
        body = client.get("/api/v1/stats").json()
    kb = body["knowledge_base"]
    assert kb["documents"] == 3 and kb["chunks"] > 0
    assert kb["formats"] == {"docx": 1, "xlsx": 1, "csv": 1}
    assert kb["content_types"]["table"] >= 2 and kb["origins"]["native"] == kb["chunks"]
    assert body["search"]["count"] == 1 and body["search"]["p50_ms"] is not None
    assert body["questions"] is None or body["questions"]["total"] == 0
    assert "SD5993" not in str(body)  # No document text leaks into statistics.
