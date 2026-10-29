"""Folder-scale upload, hard format scopes, image reader and AI editorial contracts.

All sources are generated locally; HTTP model calls use MockTransport only.
"""

import io
import json
import time
from pathlib import Path

import httpx
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from test_pdf_parser import _write_pdf
from test_web_workspace import client_for, connect, send, wait

from docqa.agent.config import AgentOptions, ChatServiceOptions
from docqa.agent.providers import HttpChatClient
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document, verify_artifact
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.docbase.retrieval import query
from docqa.docbase.storage import load_generation
from docqa.domain.documents import VisionDescription
from docqa.domain.errors import DocQAError
from docqa.domain.study import StudyPlan
from docqa.web.config import WebOptions


def docx_bytes(text="SD5993 心跳每隔30秒发送。", pictures=False):
    document = Document()
    document.add_paragraph(text)
    if pictures:
        for color in ["red", "blue"]:
            stream = io.BytesIO()
            Image.new("RGB", (80, 60), color).save(stream, format="PNG")
            stream.seek(0)
            document.add_picture(stream)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture(scope="module")
def formats(tmp_path_factory):
    root = tmp_path_factory.mktemp("format-scope")
    (root / "guide.docx").write_bytes(
        docx_bytes("SD5993 heartbeat every 30 seconds. 心跳每隔30秒发送。", pictures=True)
    )
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "SD5993 heartbeat every 30 seconds."
    from pptx.util import Inches

    slide.shapes.add_textbox(
        Inches(1), Inches(2), Inches(7), Inches(2)
    ).text = "SD5993 heartbeat every 30 seconds."
    slides.save(root / "talk.PPTX")
    _write_pdf(
        root / "paper.pdf", [b"BT /F1 12 Tf 20 250 Td (SD5993 heartbeat every 30 seconds.) Tj ET"]
    )
    artifacts = [
        Path(
            parse_document(
                root / name,
                root / "parsed",
                options=ParseOptions(
                    office_render="never",
                    ocr_mode="disabled",
                    render_pdf=False,
                ),
            ).artifact_dir
        )
        for name in ["guide.docx", "talk.PPTX", "paper.pdf"]
    ]
    retrieval = DocBaseOptions(
        mode="demo", recall_routes=["dense", "bm25"], incomplete_policy="native_only"
    )
    build_index(artifacts, root / "index", options=retrieval)
    return root / "index", AgentOptions(mode="demo", retrieval=retrieval)


@pytest.mark.parametrize("types", [["pdf"], ["docx", "pptx"], ["pdf", "docx", "pptx"], None])
def test_format_scope_is_resolved_before_worker_and_citations_stay_in_scope(
    formats, tmp_path, types
):
    index, agent = formats
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        catalog = client.get("/api/v1/documents").json()["documents"]
        allowed = {
            d["doc_id"]
            for d in catalog
            if types is None or Path(d["source_name"]).suffix.lower()[1:] in types
        }
        job = send(client, cid, doc_types=types, question="SD5993 heartbeat")
        stored = client.app.state.manager.store.get(job["job_id"])
        assert set(stored["payload"]["doc_ids"]) == allowed
        assert "doc_types" not in stored["payload"]
        answer = wait(client, "/api/v1/jobs/" + job["job_id"], ready="succeeded")["answer"]
        assert answer["citations"]
        assert all(c["doc_id"] in allowed for c in answer["citations"])
        messages = client.get(f"/api/v1/conversations/{cid}/messages").json()["items"]
        assert messages[0]["scope"]["doc_types"] == types


def test_conflicting_or_invalid_type_scopes_are_not_widened(formats, tmp_path):
    index, agent = formats
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        docs = client.get("/api/v1/documents").json()["documents"]
        word = next(d for d in docs if d["source_name"].endswith(".docx"))
        base = {"question": "SD5993", "client_request_id": "format-scope-invalid"}
        for types in [[], ["pdf", "pdf"], ["exe"], ["pdf", "pptx", "docx", "pdf"]]:
            assert (
                client.post(
                    f"/api/v1/conversations/{cid}/messages", json={**base, "doc_types": types}
                ).status_code
                == 422
            )
        response = client.post(
            f"/api/v1/conversations/{cid}/messages",
            json={**base, "doc_types": ["pdf"], "doc_ids": [word["doc_id"]]},
        )
        assert response.status_code == 409
        assert response.json()["error"]["code"] == "SCOPE_CHANGED"


def test_folder_batch_can_exceed_twenty_queued_files_and_keeps_paths(tmp_path):
    service = tmp_path / "service"
    agent = AgentOptions(
        mode="demo",
        retrieval=DocBaseOptions(
            mode="demo", recall_routes=["dense", "bm25"], incomplete_policy="validated_parts"
        ),
    )
    opts = WebOptions(
        index=service / "index",
        service_state=service,
        state=tmp_path / "web",
        upload_root=tmp_path / "sources",
        service_mode="external",
        agent=agent,
    )
    with client_for(opts) as client:
        connect(client)
        records = []
        for i in range(24):
            raw = docx_bytes(f"FOLDER_{i} SD5993 心跳每隔30秒发送。")
            created = client.post(
                "/api/v1/uploads",
                json={
                    "name": f"paper-{i}.docx",
                    "relative_path": f"资料/sub-{i % 3}/paper-{i}.docx",
                    "size": len(raw),
                },
            )
            assert created.status_code == 201, created.text
            record = created.json()
            assert record["relative_path"].startswith("资料/")
            assert (
                client.put(
                    f"/api/v1/uploads/{record['id']}/content",
                    content=raw,
                    headers={"Content-Type": "application/octet-stream"},
                ).status_code
                == 200
            )
            assert client.post(f"/api/v1/uploads/{record['id']}/commit", json={}).status_code == 202
            records.append(record)
        workspace = client.app.state.workspace
        workspace.service().run_once()
        workspace.refresh_uploads()
        assert len(client.get("/api/v1/documents").json()["documents"]) == 24
        for path in ["../evil.docx", "C:/evil.docx", "root//evil.docx", "root\\evil.docx"]:
            assert (
                client.post(
                    "/api/v1/uploads", json={"name": "evil.docx", "relative_path": path, "size": 10}
                ).status_code
                == 422
            )


def test_reader_lists_uncited_document_images_and_safe_previews(formats, tmp_path):
    index, agent = formats
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        word = next(
            d
            for d in client.get("/api/v1/documents").json()["documents"]
            if d["source_name"].endswith(".docx")
        )
        base = f"/api/v1/documents/{word['doc_id']}/images"
        first = client.get(base + "?limit=1").json()
        assert first["total"] == 2 and first["next_offset"] == 1
        second = client.get(base + "?offset=1").json()
        assert second["items"][0]["id"] != first["items"][0]["id"]
        image = first["items"][0]
        response = client.get(image["thumbnail"])
        assert response.status_code == 200 and response.headers["content-type"] == "image/png"
        with Image.open(io.BytesIO(response.content)) as picture:
            assert picture.width <= 480 and picture.height <= 480
        assert client.get(base + "/foreign/content").status_code == 404
        assert client.get(image["preview"] + "&size=90000").status_code == 422
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        job = send(client, cid, doc_types=["docx"])
        answer = wait(client, "/api/v1/jobs/" + job["job_id"], ready="succeeded")["answer"]
        citation = answer["citations"][0]
        url = citation["resources"]["images_gallery"]
        assert client.get(url).json()["total"] == 2
        client.cookies.clear()
        connect(client)
        assert client.get(url).status_code == 404


def test_successful_figure_components_survive_another_figures_failure(tmp_path):
    class MixedVision:
        count = 0

        def fingerprint(self):
            return {"adapter": "mixed-test"}

        def describe(self, *args, **kwargs):
            self.count += 1
            return (
                VisionDescription(status="ok", observations=["SAFE_FIGURE_SENTINEL"])
                if self.count == 1
                else VisionDescription(
                    status="partial",
                    observations=["UNSAFE_FIGURE_SENTINEL"],
                    uncertainties=["unreadable"],
                )
            )

    source = tmp_path / "mixed.docx"
    source.write_bytes(docx_bytes(pictures=True))
    result = parse_document(
        source,
        tmp_path / "parsed",
        options=ParseOptions(office_render="never", ocr_mode="disabled"),
        vision_client=MixedVision(),
    )
    options = DocBaseOptions(
        mode="demo", recall_routes=["dense", "bm25"], incomplete_policy="validated_parts"
    )
    built = build_index([Path(result.artifact_dir)], tmp_path / "index", options=options)
    with load_generation(Path(built.generation_dir)) as generation:
        text = "\n".join(c.index_text for c in generation.chunks)
        assert "SAFE_FIGURE_SENTINEL" in text and "UNSAFE_FIGURE_SENTINEL" not in text
        assert any(s.origin == "vl_generated" for c in generation.chunks for s in c.source_spans)
        assert generation.members[0]["source_quality"]["index_mode"] == "validated_parts"
    assert verify_artifact(Path(result.artifact_dir))["status"] != "complete"
    assert query(tmp_path / "index", "SAFE_FIGURE_SENTINEL", options=options).warnings


def test_new_policy_reads_native_only_generation_as_safe_subset(formats):
    index, agent = formats
    options = agent.retrieval.model_copy(update={"incomplete_policy": "validated_parts"})
    assert query(index, "SD5993", options=options).hits
    with pytest.raises(DocQAError):
        query(index, "SD5993", options=options.model_copy(update={"incomplete_policy": "reject"}))


@pytest.mark.parametrize(
    "ids,valid", [(["A", "B"], True), (["A", "C"], False), (["A", "A", "B"], False), (["A"], False)]
)
def test_ai_organizer_cannot_add_omit_or_repeat_claims(ids, valid):
    def handler(request):
        payload = json.loads(request.content)
        task = json.loads(payload["messages"][1]["content"])
        assert task["task"] == "organize_study"
        assert task["data"]["facts"][0]["text"] == "checked fact"
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {
                            "role": "assistant",
                            "content": json.dumps(
                                {"sections": [{"title": "研究方法", "claim_ids": ids}]}
                            ),
                        },
                    }
                ]
            },
        )

    client = HttpChatClient(
        ChatServiceOptions(
            endpoint="https://synthetic.invalid/chat/completions",
            model="mock",
            revision="1",
            retries=0,
        ),
        transport=httpx.MockTransport(handler),
    )
    facts = [{"id": "A", "text": "checked fact"}, {"id": "B", "text": "another checked fact"}]
    if valid:
        assert isinstance(
            client.organize_study("问题", facts, "mindmap", deadline=time.monotonic() + 2),
            StudyPlan,
        )
    else:
        with pytest.raises(DocQAError) as error:
            client.organize_study("问题", facts, "slides")
        assert error.value.code == "ARTIFACT_PLAN_INVALID"
