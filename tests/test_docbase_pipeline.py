"""Real doc2md artifacts exercise atomic DocBase rebuilds and source retention."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from typer.testing import CliRunner

from docqa.cli.main import app
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index, select_artifacts
from docqa.docbase.providers import make_providers
from docqa.docbase.retrieval import query, resolve_chunk
from docqa.docbase.storage import read_current, verify_generation
from docqa.domain.errors import DocQAError


def source(tmp_path, name="a", text="SD5993 调用 FE_HEART_Start 启动心跳。", category="protocol"):
    path = tmp_path / (name + ".docx")
    document = Document()
    document.add_heading("协议说明", level=1)
    document.add_paragraph(text)
    document.save(path)
    result = parse_document(
        path,
        tmp_path / "parsed",
        options=ParseOptions(category=category, office_render="never", ocr_mode="disabled"),
    )
    assert result.status == "complete"
    return Path(result.artifact_dir), path


@pytest.fixture
def options():
    return DocBaseOptions(mode="demo", analyzer_terms=["SD5993", "FE_HEART_Start"])


def test_build_query_reuse_resolve_and_packaged_sources(tmp_path, options):
    artifact, original = source(tmp_path)
    output = tmp_path / "kb"
    first = build_index(artifact, output, options=options)
    assert first.documents == 1 and first.chunks > 0
    assert first.synthetic and not first.production_ready
    assert verify_generation(Path(first.generation_dir))["valid"]
    again = build_index(artifact, output, options=options)
    assert again.reused and again.generation_id == first.generation_id
    assert read_current(output) == (first.generation_id, 1)
    result = query(output, "FE_HEART_Start 如何启动心跳？", options=options)
    assert result.status == "ok" and result.hits
    assert set(result.routes) == {"dense", "bm25", "sparse", "reranker"}
    resolved = resolve_chunk(output, result.hits[0].chunk_id)
    assert Path(resolved["snapshot_path"]).is_file()
    original.unlink()
    assert query(output, "SD5993", options=options).hits


def test_merge_replace_empty_and_retained_history(tmp_path, options):
    first_artifact, _ = source(tmp_path, "a", "第一份 SD5993 文档")
    second_artifact, _ = source(tmp_path, "b", "第二份 TLS 文档", "security")
    output = tmp_path / "kb"
    first = build_index(first_artifact, output, options=options)
    second = build_index(second_artifact, output, options=options)
    assert second.documents == 2
    assert second.cache_hits["dense"] > 0 and second.cache_hits["sparse"] > 0
    third = build_index(second_artifact, output, options=options, replace=True)
    assert third.documents == 1
    assert query(output, "SD5993", options=options, generation_id=first.generation_id).hits
    with pytest.raises(DocQAError, match="explicit"):
        build_index(None, output, options=options, replace=True)
    empty = build_index(None, output, options=options, replace=True, allow_empty=True)
    assert empty.documents == empty.chunks == 0
    assert query(output, "问题", options=options).reason == "index_empty"
    assert verify_generation(Path(first.generation_dir))["valid"]


def test_failed_build_keeps_published_generation(tmp_path, options):
    artifact, _ = source(tmp_path)
    output = tmp_path / "kb"
    first = build_index(artifact, output, options=options)
    next_artifact, _ = source(tmp_path, "b", "新的内容触发模型调用")
    dense, sparse, reranker = make_providers(options)

    class FailedDense:
        def fingerprint(self):
            return dense.fingerprint()

        def encode(self, texts, input_type="document"):
            raise DocQAError("TEST_NETWORK_FAILURE", "synthetic failure")

    with pytest.raises(DocQAError) as error:
        build_index(
            next_artifact, output, options=options, providers=(FailedDense(), sparse, reranker)
        )
    assert error.value.code == "TEST_NETWORK_FAILURE"
    assert read_current(output) == (first.generation_id, 1)
    assert query(output, "心跳", options=options).hits


def test_explicit_source_selection_rejects_two_revisions(tmp_path, options):
    first, original = source(tmp_path)
    second, _ = source(tmp_path, text="已修改的文档 SD5993")
    assert first != second
    with pytest.raises(DocQAError) as error:
        build_index([first, second], tmp_path / "kb", options=options)
    assert error.value.code == "DUPLICATE_DOCUMENT"
    assert select_artifacts(tmp_path / "parsed") == [second]


def test_source_list_relative_paths_and_plain_md_rejected(tmp_path, options):
    artifact, _ = source(tmp_path)
    listing = tmp_path / "sources.json"
    listing.write_text(
        json.dumps(
            {
                "schema_version": "docqa.docbase-sources.v1",
                "artifacts": [artifact.relative_to(tmp_path).as_posix()],
            }
        ),
        encoding="utf-8",
    )
    assert select_artifacts(listing) == [artifact]
    markdown = tmp_path / "standalone.md"
    markdown.write_text("无来源契约", encoding="utf-8")
    with pytest.raises(DocQAError) as error:
        build_index(markdown, tmp_path / "kb", options=options)
    assert error.value.code == "INVALID_REQUEST"


def test_review_required_source_is_not_silently_indexed(tmp_path, options):
    from PIL import Image

    image = tmp_path / "figure.png"
    Image.new("RGB", (100, 100), "white").save(image)
    document = Document()
    document.add_paragraph("含图文档")
    document.add_picture(str(image))
    path = tmp_path / "image.docx"
    document.save(path)
    parsed = parse_document(
        path, tmp_path / "parsed", options=ParseOptions(ocr_mode="disabled", office_render="never")
    )
    assert parsed.status == "review_required"
    with pytest.raises(DocQAError) as error:
        build_index(parsed.artifact_dir, tmp_path / "kb", options=options)
    assert error.value.code == "SOURCE_NOT_READY"


def test_cli_build_query_verify_and_errors_are_json(tmp_path):
    artifact, _ = source(tmp_path)
    output = tmp_path / "kb"
    runner = CliRunner()
    build = runner.invoke(
        app, ["docbase", "build", str(artifact), "-o", str(output), "--demo", "--no-progress"]
    )
    assert build.exit_code == 0, build.output
    assert json.loads(build.stdout)["synthetic"]
    response = runner.invoke(app, ["docbase", "query", "SD5993", "-i", str(output), "--demo"])
    assert response.exit_code == 0, response.output
    assert json.loads(response.stdout)["hits"]
    verified = runner.invoke(app, ["docbase", "verify", str(output)])
    assert verified.exit_code == 0 and json.loads(verified.stdout)["valid"]
    invalid = runner.invoke(
        app, ["docbase", "query", "SD5993", "-i", str(output), "--demo", "--top-k", "0"]
    )
    assert invalid.exit_code == 2, invalid.output
    assert json.loads(invalid.stdout)["error"]["code"] == "INVALID_REQUEST"
    unconfigured = runner.invoke(
        app, ["docbase", "build", str(artifact), "-o", str(tmp_path / "real"), "--no-progress"]
    )
    assert unconfigured.exit_code == 4, unconfigured.output
    assert json.loads(unconfigured.stdout)["error"]["code"] == "MODEL_NOT_CONFIGURED"
