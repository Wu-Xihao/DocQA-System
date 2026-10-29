"""Synthetic acceptance for storage migration, classification, monitoring and streams."""

import io
import json
import shutil
from pathlib import Path

import numpy as np
import pytest
import test_agent_integration as shared
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from test_web_workspace import client_for, connect, send, wait

from docqa.agent.config import AgentOptions
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document, verify_artifact
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.migration import migrate_vector_backend
from docqa.docbase.retrieval import query
from docqa.docbase.storage import _io_path, load_generation, read_current, search_dense
from docqa.docbase.vector_db import create, search, verify
from docqa.domain.documents import VisionDescription
from docqa.domain.errors import DocQAError
from docqa.sparse_service import create_app as sparse_app
from docqa.web.config import WebOptions


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    return shared.corpus.__wrapped__(tmp_path_factory)


def test_sqlite_vec_filters_before_knn_and_matches_cosine(tmp_path):
    rng = np.random.default_rng(18)
    vectors = rng.normal(size=(40, 8)).astype(np.float32)
    vectors /= np.linalg.norm(vectors, axis=1, keepdims=True)
    path = tmp_path / "vectors.sqlite"
    create(path, vectors)
    verify(path, vectors)
    eligible = [1, 8, 13, 22, 27]
    actual = search(path, vectors[0], eligible, 3, dimension=8, rows=40)
    expected = search_dense(vectors, vectors[0], eligible, 3)
    assert [v[0] for v in actual] == [v[0] for v in expected]
    assert np.allclose([v[1] for v in actual], [v[1] for v in expected], atol=1e-6)
    assert search(path, vectors[0], [], 3, dimension=8, rows=40) == []


def test_install_generation_retries_a_transient_windows_lock(tmp_path, monkeypatch):
    import os

    from docqa.docbase.storage import install_generation

    if os.name != "nt":
        pytest.skip("Windows file-sharing semantics")
    staging = tmp_path / "staging/candidate"
    staging.mkdir(parents=True)
    (tmp_path / "generations").mkdir()
    (staging / "sample").write_bytes(b"checked data")
    real = Path.replace
    attempts = []

    def locked_once(source, target):
        attempts.append(source)
        if len(attempts) == 1:
            raise PermissionError("simulated sharing lock")
        return real(source, target)

    monkeypatch.setattr(Path, "replace", locked_once)
    result = install_generation(tmp_path, staging, "a" * 32)
    assert len(attempts) == 2 and (result / "sample").read_bytes() == b"checked data"


def test_vector_backend_migration_preserves_evidence_and_never_needs_models(
    corpus, tmp_path, monkeypatch
):
    index, agent, _ = corpus
    copied = tmp_path / "index"
    shutil.copytree(_io_path(index), _io_path(copied))
    old = read_current(copied)[0]
    with load_generation(copied / "generations" / old) as generation:
        before = [c.chunk_id for c in generation.chunks]

    def forbidden(*a, **k):
        raise AssertionError("Migration cannot call models")

    monkeypatch.setattr("docqa.docbase.providers.make_providers", forbidden)
    result = migrate_vector_backend(copied)
    assert result["model_calls"] == 0 and result["generation_id"] != old
    with load_generation(copied / "generations" / result["generation_id"]) as generation:
        assert [c.chunk_id for c in generation.chunks] == before
        assert generation.profiles["vector_backend"] == "sqlite_vec"
        assert (generation.directory / "dense/vectors.sqlite").is_file()
    assert migrate_vector_backend(copied)["reused"]
    monkeypatch.undo()
    found = query(copied, "SD5993", options=agent.retrieval)
    assert found.hits
    path = copied / "generations" / result["generation_id"] / "dense/vectors.sqlite"
    path.write_bytes(b"changed database")
    with pytest.raises(DocQAError):
        load_generation(path.parent.parent)


def test_auto_category_and_named_figure_keep_vl_text_and_source_links(tmp_path):
    class Vision:
        def fingerprint(self):
            return {"adapter": "synthetic-local-test"}

        def describe(self, *args, **kwargs):
            return VisionDescription(
                status="ok", image_type="流程图", observations=["VLSENTINEL 解析后执行检索。"]
            )

    doc = Document()
    doc.add_paragraph("Abstract: This journal paper reports the methodology and references.")
    image = io.BytesIO()
    Image.new("RGB", (50, 50), "red").save(image, format="PNG")
    image.seek(0)
    doc.add_picture(image)
    source = tmp_path / "Research.docx"
    doc.save(source)
    result = parse_document(
        source,
        tmp_path / "parsed",
        options=ParseOptions(category="auto", office_render="never", ocr_mode="disabled"),
        vision_client=Vision(),
    )
    ir = json.loads(Path(result.artifact_dir, "document.ir.json").read_bytes())
    assert ir["category"] == "研究论文"
    assert ir["metadata"]["classification"]["method"] == "local_keyword_rules.v1"
    figure = ir["figures"][0]
    assert "Research_PageNA图001_Figure" in figure["metadata"]["trace_name"]
    assert _io_path(Path(result.artifact_dir, figure["metadata"]["trace_image_path"])).is_file()
    md = Path(result.markdown_path).read_text("utf8")
    assert "VLSENTINEL" in md and "打开原文" in md and "images/figures/" in md
    assert verify_artifact(Path(result.artifact_dir))


def test_monitor_registration_recursion_types_and_incremental_updates(tmp_path):
    sources = tmp_path / "uploads"
    (sources / "sub").mkdir(parents=True)
    file = sources / "sub/one.docx"

    def write(value):
        doc = Document()
        doc.add_paragraph(value)
        doc.save(file)

    write("Protocol: automatic monitoring first revision")
    (sources / "ignore.txt").write_text("ignore", encoding="utf8")
    service = tmp_path / "service"
    agent = AgentOptions(
        mode="demo",
        retrieval=DocBaseOptions(
            mode="demo",
            recall_routes=["dense", "bm25"],
            incomplete_policy="validated_parts",
            vector_backend="sqlite_vec",
        ),
    )
    opts = WebOptions(
        index=service / "index",
        service_state=service,
        state=tmp_path / "web",
        upload_root=sources,
        service_mode="external",
        agent=agent,
    )
    with client_for(opts) as client:
        connect(client)
        created = client.post(
            "/api/v1/maintenance/sources",
            json={
                "path": str(sources),
                "types": ["docx"],
                "recursive": True,
                "interval_seconds": 15,
            },
        )
        assert created.status_code == 201, created.text
        sid = created.json()["source_id"]
        assert (
            client.get(f"/api/v1/maintenance/sources/{sid}/files")
            .json()["files"][0]["path"]
            .endswith("one.docx")
        )
        worker = client.app.state.workspace.service()
        worker.run_once()
        original = worker.list_documents()[0]
        assert original["category"] == "技术规范"
        generation = read_current(service / "index")[0]
        write("Protocol: automatic monitoring changed second revision")
        worker.run_once()
        assert worker.list_documents()[0]["target_version"] > original["target_version"]
        assert read_current(service / "index")[0] != generation
        assert len(worker.list_documents()) == 1
        denied = client.post(
            "/api/v1/maintenance/sources",
            json={"path": str(tmp_path / "outside"), "types": ["pdf"]},
        )
        assert denied.status_code != 201
        assert (
            client.patch(f"/api/v1/maintenance/sources/{sid}", json={"enabled": False}).json()[
                "enabled"
            ]
            is False
        )


def test_sse_stream_contains_only_checked_paragraphs_and_finishes(corpus, tmp_path):
    index, agent, _ = corpus
    with client_for(WebOptions(index=index, state=tmp_path, agent=agent)) as client:
        connect(client)
        cid = client.post("/api/v1/conversations", json={}).json()["id"]
        job = send(client, cid)
        answer = wait(client, "/api/v1/jobs/" + job["job_id"], ready="succeeded")["answer"]
        response = client.get("/api/v1/events?conversation_id=" + cid)
        assert response.status_code == 200
        events = [
            json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")
        ]
        assert events[0]["kind"] == "answer_start" and events[-1]["kind"] == "answer_done"
        assert [e["claim"] for e in events if e["kind"] == "answer_delta"] == answer["claims"]
        assert events[-1]["job"]["answer"]["citations"] == answer["citations"]


def test_bge_sparse_service_protocol_is_learned_sparse_not_dense(monkeypatch):
    class Tokenizer:
        def __len__(self):
            return 1000

        def encode(self, text, **kw):
            return list(text)

    class Encoder:
        tokenizer = Tokenizer()

        def encode(self, texts, **kw):
            assert kw["return_sparse"] and not kw["return_dense"] and not kw["return_colbert_vecs"]
            return {"lexical_weights": [{"10": 0.8, "33": 0.4} for _ in texts]}

    payload = {
        "model": "BAAI/bge-m3",
        "revision": "local-v1",
        "tokenizer_revision": "local-v1",
        "input": ["synthetic sentence"],
        "input_type": "query",
        "truncate": False,
        "return_sparse": True,
    }
    with TestClient(sparse_app(encoder=Encoder())) as client:
        assert client.get("/health").json()["vocab_size"] == 1000
        assert client.post("/sparse", json=payload).json()["data"][0]["lexical_weights"] == {
            "10": 0.8,
            "33": 0.4,
        }
        assert client.post("/sparse", json={**payload, "revision": "wrong"}).status_code == 409
        assert client.post("/sparse", json={**payload, "truncate": True}).status_code == 422
        client.app.state.encoder.encode = lambda *a, **kw: {
            "lexical_weights": [{"10": float("nan"), "33": 0.4}]
        }
        assert client.post("/sparse", json=payload).status_code == 503
        monkeypatch.setenv("DOCQA_BGE_SPARSE_KEY", "synthetic-test-secret")
        assert client.post("/sparse", json=payload).status_code == 401


def test_three_route_build_and_query_via_sparse_http_adapter(corpus, tmp_path):
    import httpx

    from docqa.docbase.config import ServiceOptions
    from docqa.docbase.pipeline import build_index
    from docqa.docbase.providers import DemoDenseClient, DemoReranker, HttpSparseClient

    index, _, _ = corpus
    generation_id = read_current(index)[0]
    with load_generation(index / "generations" / generation_id) as generation:
        sources = [generation.directory / m["artifact_ref"] for m in generation.members]
    calls = []

    def sparse_http(request):
        body = json.loads(request.content)
        calls.append(body["input_type"])
        return httpx.Response(
            200,
            json={
                "data": [
                    {"index": i, "lexical_weights": {"10": 0.7, "44": 0.3}}
                    for i in range(len(body["input"]))
                ]
            },
        )

    options = DocBaseOptions(
        mode="demo",
        vector_backend="sqlite_vec",
        sparse=ServiceOptions(
            endpoint="http://127.0.0.1:18100/sparse",
            model="BAAI/bge-m3",
            revision="local-v1",
            tokenizer_revision="local-v1",
            vocab_size=1000,
        ),
    )
    providers = (
        DemoDenseClient(options.dense),
        HttpSparseClient(options.sparse, transport=httpx.MockTransport(sparse_http)),
        DemoReranker(options.reranker),
    )
    build_index(sources, tmp_path / "index", options=options, providers=providers)
    found = query(tmp_path / "index", "SD5993", options=options, providers=providers)
    assert {"dense", "bm25", "sparse", "reranker"}.issubset(found.routes)
    assert all(hit.rerank_score is not None for hit in found.hits)
    assert found.hits and "document" in calls and "query" in calls
    assert any("sparse" in hit.route_ranks for hit in found.hits)
