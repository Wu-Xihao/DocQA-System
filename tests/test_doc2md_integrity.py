"""Adversarial package tests bind metadata to IR, source bytes and projection.

Changed JSON files deliberately receive fresh manifest checksums: these cases
exercise cross-file consistency checks rather than only the byte-checksum gate.
"""

import json
import os
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image

from docqa.doc2md.common import digest
from docqa.doc2md.config import ParseOptions
from docqa.doc2md.markdown import render_markdown
from docqa.doc2md.pipeline import (
    _canonical_ir,
    _quality,
    parse_document,
    verify_artifact,
)
from docqa.domain.documents import DocumentIR, ImageAsset
from docqa.domain.errors import DocQAError


def read_json(path: Path):
    return json.loads(path.read_bytes())


def write_json(path: Path, value):
    path.write_bytes((json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))


def checksum(path: Path):
    return sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def artifact(tmp_path):
    source = tmp_path / "input" / "integrity.docx"
    source.parent.mkdir()
    document = Document()
    document.add_heading("Integrity fixture", level=1)
    document.add_paragraph("The input snapshot contains an immutable original fact.")
    document.add_paragraph("Independent paragraph with non-ASCII content: 中文定位。")
    document.save(source)
    result = parse_document(
        source,
        tmp_path / "output",
        options=ParseOptions(
            allowed_roots=[source.parent],
            office_render="never",
            ocr_mode="disabled",
            tesseract_executable="__docqa_integrity_no_tesseract__",
            office_executable="__docqa_integrity_no_office__",
        ),
    )
    directory = Path(result.artifact_dir)
    assert verify_artifact(directory)["valid"]
    assert result.status == "complete"
    return directory


def update_file(directory: Path, filename: str, mutate):
    path = directory / filename
    value = read_json(path)
    mutate(value)
    write_json(path, value)
    manifest_path = directory / "manifest.json"
    manifest = read_json(manifest_path)
    manifest["files"][filename] = checksum(path)
    write_json(manifest_path, manifest)


def assert_corrupt(directory: Path):
    with pytest.raises(DocQAError) as error:
        verify_artifact(directory)
    assert error.value.code == "ARTIFACT_CORRUPT"


@pytest.mark.parametrize("field", ["doc_id", "revision_id", "parse_profile_id", "parse_key"])
def test_manifest_identity_cannot_override_ir(artifact, field):
    path = artifact / "manifest.json"
    manifest = read_json(path)
    manifest[field] = str(uuid4()) if field == "doc_id" else "9" * 64
    write_json(path, manifest)
    assert_corrupt(artifact)


def test_manifest_cannot_swap_source_to_another_valid_content_addressed_blob(artifact):
    path = artifact / "manifest.json"
    manifest = read_json(path)
    root = artifact.parent.parent
    replacement = b"An unrelated source with a valid independent checksum"
    content_hash = sha256(replacement).hexdigest()
    blob = root / "blobs" / content_hash[:2] / (content_hash + ".docx")
    blob.parent.mkdir(parents=True, exist_ok=True)
    blob.write_bytes(replacement)
    manifest["snapshot_ref"] = blob.relative_to(root).as_posix()
    manifest["content_sha256"] = content_hash
    write_json(path, manifest)
    assert_corrupt(artifact)


def test_snapshot_with_same_bytes_cannot_escape_blob_root(artifact):
    path = artifact / "manifest.json"
    manifest = read_json(path)
    root = artifact.parent.parent
    original = root / manifest["snapshot_ref"]
    external = root.parent / "outside-source.docx"
    external.write_bytes(original.read_bytes())
    manifest["snapshot_ref"] = str(external)
    write_json(path, manifest)
    assert_corrupt(artifact)


@pytest.mark.parametrize("field", ["eligible_for_indexing", "status", "counts"])
def test_manifest_quality_cannot_override_derived_quality(artifact, field):
    path = artifact / "manifest.json"
    manifest = read_json(path)
    if field == "eligible_for_indexing":
        manifest["quality"][field] = False
    elif field == "status":
        manifest["quality"][field] = "review_required"
    else:
        manifest["quality"][field]["blocks"] += 1
    write_json(path, manifest)
    assert_corrupt(artifact)


def test_modified_quality_report_rejected_even_with_fresh_file_checksum(artifact):
    update_file(
        artifact, "quality_report.json", lambda value: value.update(eligible_for_indexing=False)
    )
    assert_corrupt(artifact)


def test_coordinated_quality_manifest_and_report_still_bound_to_ir(artifact):
    update_file(
        artifact,
        "quality_report.json",
        lambda value: value.update(status="review_required", eligible_for_indexing=False),
    )
    path = artifact / "manifest.json"
    manifest = read_json(path)
    manifest["quality"] = read_json(artifact / "quality_report.json")
    write_json(path, manifest)
    assert_corrupt(artifact)


@pytest.mark.parametrize(
    "mutation",
    [
        "omit_block",
        "offset",
        "locator",
        "anchor",
        "line_end",
        "text_hash",
        "origin",
        "text_range",
        "artifact_id",
        "schema",
    ],
)
def test_source_map_mutations_rejected_after_updating_manifest_hash(artifact, mutation):
    def mutate(value):
        entry = value["entries"][0]
        if mutation == "omit_block":
            value["entries"].pop()
        elif mutation == "offset":
            # Moving just past '<' leaves the anchor substring and line number
            # intact, so simple bounds/substring checks are insufficient.
            entry["start_byte"] += 1
        elif mutation == "locator":
            entry["locators"][0]["native_ref"]["object_path"] = "/w:document/w:body/w:p[999]"
        elif mutation == "anchor":
            entry["anchor"] = "block-incorrect-anchor"
        elif mutation == "line_end":
            entry["line_end"] += 3
        elif mutation == "text_hash":
            entry["source_text_sha256"] = "0" * 64
        elif mutation == "origin":
            entry["origin"] = "ocr"
        elif mutation == "text_range":
            entry["text_end_codepoint"] += 2
        elif mutation == "artifact_id":
            value["artifact_id"] = "0" * 64
        else:
            value["schema_version"] = "docqa.source-map.v999"

    update_file(artifact, "source_map.json", mutate)
    assert_corrupt(artifact)


def test_ir_text_cannot_change_under_original_artifact_identity(artifact):
    def mutate(value):
        value["blocks"][0]["text"] += " forged replacement fact"

    update_file(artifact, "document.ir.json", mutate)
    assert_corrupt(artifact)


def test_changed_profile_cannot_keep_original_profile_identity(artifact):
    def mutate(value):
        value["settings"]["category"] = "forged-profile-category"

    update_file(artifact, "profile.json", mutate)
    assert_corrupt(artifact)


@pytest.mark.parametrize(
    "filename",
    ["document.ir.json", "document.md", "source_map.json", "quality_report.json", "profile.json"],
)
def test_required_file_must_be_manifested_even_if_still_on_disk(artifact, filename):
    path = artifact / "manifest.json"
    manifest = read_json(path)
    del manifest["files"][filename]
    write_json(path, manifest)
    assert_corrupt(artifact)


def resign_modified_ir(directory: Path, ir: DocumentIR) -> Path:
    """Build a self-consistent candidate so path checks are reached independently.

    Production canonicalization/projection helpers calculate legitimate content
    identities; this helper does not duplicate their algorithms.
    """
    manifest = read_json(directory / "manifest.json")
    quality = _quality(ir)
    identity = digest(_canonical_ir(ir))
    ir.artifact_id = identity
    for block in ir.blocks:
        for locator in block.locators:
            locator.artifact_id = identity
    for figure in ir.figures:
        figure.locator.artifact_id = identity
    root = directory.parent.parent
    markdown, source_map = render_markdown(
        ir,
        snapshot_link=os.path.relpath(root / manifest["snapshot_ref"], directory),
        status=quality["status"],
    )
    write_json(directory / "document.ir.json", ir.model_dump(mode="json"))
    (directory / "document.md").write_bytes(markdown)
    write_json(directory / "source_map.json", source_map)
    write_json(directory / "quality_report.json", quality)
    manifest["artifact_id"] = identity
    manifest["quality"] = quality
    manifest["files"] = {filename: checksum(directory / filename) for filename in manifest["files"]}
    write_json(directory / "manifest.json", manifest)
    target = directory.with_name(identity)
    directory.rename(target)
    return target


@pytest.mark.parametrize("absolute", [False, True])
def test_image_asset_path_cannot_escape_artifact_even_with_fully_rebound_identity(
    artifact, absolute
):
    outside = artifact.parent.parent / "outside-image.png"
    Image.new("RGB", (2, 2), "white").save(outside)
    image_hash = checksum(outside)
    ir = DocumentIR.model_validate_json((artifact / "document.ir.json").read_bytes())
    ir.assets.append(
        ImageAsset(
            asset_id=image_hash,
            sha256=image_hash,
            relative_path=str(outside) if absolute else "../../outside-image.png",
            media_type="image/png",
            width=2,
            height=2,
        )
    )
    candidate = resign_modified_ir(artifact, ir)
    assert_corrupt(candidate)
