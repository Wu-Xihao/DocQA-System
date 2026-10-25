"""Exercise real local ingestion, publication and failure-safe reconciliation."""

from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path

import portalocker
import pytest
from docx import Document
from PIL import Image

from docqa.doc2md.config import ParseOptions
from docqa.doc2md.pipeline import parse_document
from docqa.docbase.config import DocBaseOptions
from docqa.docbase.pipeline import build_index
from docqa.docbase.retrieval import query
from docqa.domain.documents import VisionDescription
from docqa.domain.errors import DocQAError
from docqa.service.config import ServiceOptions, load_options
from docqa.service.pipeline import Service
from docqa.service.scanner import ScanSnapshot


class Clock:
    def __init__(self):
        self.value = 1000.0

    def __call__(self):
        return self.value


def write_doc(path, text="SD5993 调用 FE_HEART_Start 启动心跳。"):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


@pytest.fixture
def service(tmp_path):
    source = tmp_path / "in"
    source.mkdir()
    options = ServiceOptions(
        allowed_source_roots=[source],
        settle_seconds=0,
        retry_base_seconds=0,
        parse=ParseOptions(office_render="never", ocr_mode="disabled"),
        retrieval=DocBaseOptions(mode="demo"),
    )
    return Service(tmp_path / "s", options=options, clock=Clock()), source


def ingest(service):
    svc, source = service
    write_doc(source / "heart.docx")
    root = svc.add_source(source, types=("docx",), category="protocol")
    run = svc.run_once()
    assert run["failed"] == 0, run
    doc = svc.list_documents()[0]
    assert svc.document(doc["doc_id"])["target_satisfied"]
    return svc, source, root, doc


def test_real_ingest_query_reuse_and_metadata_touch(service):
    svc, source, _, doc = ingest(service)
    first = svc.index_status()
    assert first["documents"] == 1 and first["synthetic"]
    assert query(svc.index, "SD5993", options=svc.options.retrieval).hits
    os.utime(source / "heart.docx", None)
    svc.enqueue_scan()
    svc.run_once()
    assert svc.index_status()["generation_id"] == first["generation_id"]
    assert svc.registry.get("documents", doc["doc_id"])["target_version"] == 1
    assert svc.status()["unsatisfied_documents"] == []


def test_changed_source_preserves_identity_and_historical_generation(service):
    svc, source, _, doc = ingest(service)
    previous = svc.index_status()["generation_id"]
    write_doc(source / "heart.docx", "SD5993 新版心跳间隔为两秒。")
    svc.enqueue_scan()
    assert svc.run_once()["failed"] == 0
    current = svc.document(doc["doc_id"])
    assert current["target_version"] == 2 and current["target_satisfied"]
    assert current["published_revision_id"]
    assert svc.index_status()["generation_id"] != previous
    assert query(
        svc.index, "FE_HEART_Start", options=svc.options.retrieval, generation_id=previous
    ).hits


def test_bad_revision_retains_old_member_without_false_sync(service):
    svc, source, _, doc = ingest(service)
    previous = svc.index_status()["generation_id"]
    (source / "heart.docx").write_bytes(b"not an office document")
    svc.enqueue_scan()
    assert svc.run_once()["failed"] == 1
    assert svc.index_status()["generation_id"] == previous
    current = svc.document(doc["doc_id"])
    assert current["state"] == "failed" and not current["target_satisfied"]
    count = svc.queue.stats()["total"]
    svc.run_once()
    assert svc.queue.stats()["total"] == count


def test_stable_deletion_requires_two_complete_scans_and_reappearing_recovers(service):
    svc, source, _, doc = ingest(service)
    original = (source / "heart.docx").read_bytes()
    (source / "heart.docx").unlink()
    svc.enqueue_scan()
    svc.run_once()
    assert svc.index_status()["documents"] == 1

    svc.enqueue_scan()
    svc.run_once()
    assert svc.index_status()["documents"] == 0
    assert svc.document(doc["doc_id"])["removed"]
    (source / "heart.docx").write_bytes(original)
    svc.enqueue_scan()
    svc.run_once()
    assert svc.index_status()["documents"] == 1
    assert svc.document(doc["doc_id"])["target_satisfied"]


def test_incomplete_scan_never_tombstones_and_does_not_storm(service, monkeypatch):
    svc, source, root, doc = ingest(service)
    (source / "heart.docx").unlink()
    monkeypatch.setattr(
        svc,
        "_snapshot",
        lambda *a, **kw: ScanSnapshot(
            source, complete=False, errors=[{"code": "SOURCE_READ_FAILED"}]
        ),
    )
    svc.enqueue_scan()
    svc.run_once()
    assert not svc.document(doc["doc_id"])["removed"]
    count = svc.queue.stats()["total"]
    svc.run_once()
    assert svc.queue.stats()["total"] == count
    assert svc.registry.get("sources", root["source_id"])["scan_error"] == "SCAN_INCOMPLETE"


def test_missing_root_cannot_remove_members(service):
    svc, source, root, doc = ingest(service)
    source.rename(source.parent / "offline")
    svc.enqueue_scan(root["source_id"])
    assert svc.run_once()["failed"] == 1
    assert svc.index_status()["documents"] == 1
    assert not svc.document(doc["doc_id"])["removed"]


def test_explicit_remove_restore_and_category_update(service):
    svc, _, root, doc = ingest(service)
    svc.remove_document(doc["doc_id"])
    svc.run_once()
    assert svc.index_status()["documents"] == 0
    svc.enqueue_scan()
    svc.run_once()
    assert svc.index_status()["documents"] == 0
    svc.restore_document(doc["doc_id"])
    svc.run_once()
    assert svc.document(doc["doc_id"])["target_satisfied"]
    svc.update_source(root["source_id"], category="operations")
    svc.run_once()
    assert svc.document(doc["doc_id"])["category"] == "operations"
    assert query(svc.index, "SD5993", options=svc.options.retrieval, categories=["operations"]).hits


def test_source_removal_excludes_documents_but_preserves_files(service):
    svc, source, root, _ = ingest(service)
    svc.remove_source(root["source_id"])
    svc.run_once()
    assert svc.index_status()["documents"] == 0
    assert (source / "heart.docx").is_file()
    restored = svc.add_source(source, types=("docx",), category="protocol")
    assert restored["source_id"] == root["source_id"]
    svc.run_once()
    assert svc.index_status()["documents"] == 1
    # Re-registering with a narrower policy cannot revive excluded types.
    svc.remove_source(root["source_id"])
    svc.run_once()
    svc.add_source(source, types=("pdf",), category="protocol")
    svc.run_once()
    assert svc.index_status()["documents"] == 0


def test_new_root_recursive_search_single_file_and_authorization(service):
    svc, source = service
    write_doc(source / "a.docx")
    write_doc(source / "nested" / "b.docx")
    root = svc.add_source(source, recursive=False, types=("docx",))
    assert len(svc.discover(root["source_id"])["files"]) == 1
    with pytest.raises(DocQAError):
        svc.enqueue_parse(source / "nested" / "b.docx")
    with pytest.raises(DocQAError):
        svc.add_source(source / "nested")
    with pytest.raises(DocQAError):
        svc.add_source(svc.root)
    svc.enqueue_parse(source / "a.docx")
    svc.run_once()
    assert svc.index_status()["documents"] == 1


def test_polling_discovers_new_file_and_periodic_hash_checks_silent_change(service):
    svc, source, root, doc = ingest(service)
    write_doc(source / "b.docx", "TLS 证书需要验证。")
    svc.run_once()  # Metadata probe notices a new path without an explicit scan.
    assert svc.index_status()["documents"] == 2
    svc.clock.value += 1801
    svc.run_once()
    assert svc.registry.get("sources", root["source_id"])["last_success_at"] == svc.clock()
    assert svc.document(doc["doc_id"])["target_satisfied"]


def test_settling_waits_for_second_observation(service):
    svc, source = service
    svc.options.settle_seconds = 3
    write_doc(source / "a.docx")
    svc.add_source(source, types=("docx",))
    svc.run_once()
    assert svc.list_documents()[0]["state"] == "settling"
    assert svc.index_status()["documents"] == 0
    svc.clock.value += 3
    svc.run_once()
    assert svc.index_status()["documents"] == 1


def test_small_queue_drains_desired_outbox(service):
    svc, source = service
    svc.queue.max_pending = 1
    for i in range(3):
        write_doc(source / f"d{i}.docx", f"SD5993 文档 {i}")
    svc.add_source(source, types=("docx",))
    assert svc.run_once()["failed"] == 0
    assert svc.index_status()["documents"] == 3
    assert svc.queue.stats()["active"] == 0


def test_retry_exhaustion_does_not_restart_forever(service):
    svc, source = service

    def unavailable(*args, **kwargs):
        raise DocQAError("MODEL_TIMEOUT", "mock outage", retryable=True)

    svc.parse_function = unavailable
    write_doc(source / "a.docx")
    svc.add_source(source, types=("docx",))
    svc.run_once()
    failed = [j for j in svc.queue.list() if j["kind"] == "parse"][0]
    assert failed["attempt"] == 3 and failed["state"] == "failed"
    count = svc.queue.stats()["total"]
    svc.run_once()
    assert svc.queue.stats()["total"] == count


def test_permanent_build_failure_is_visible_and_explicitly_retryable(service):
    svc, source = service

    def unavailable(*args, **kwargs):
        raise DocQAError("MODEL_NOT_CONFIGURED", "missing service")

    svc.build_function = unavailable
    write_doc(source / "a.docx")
    svc.add_source(source, types=("docx",))
    assert svc.run_once()["failed"] == 1
    count = svc.queue.stats()["total"]
    svc.run_once()
    assert svc.queue.stats()["total"] == count
    svc.build_function = build_index
    svc.enqueue_build()
    svc.run_once()
    assert svc.index_status()["documents"] == 1


def test_generation_publish_rejects_a_changed_desired_state(service):
    svc, _, _, doc = ingest(service)
    previous = svc.index_status()["generation_id"]
    changed = False

    def build_with_change(*args, **kwargs):
        def progress(stage, current, total):
            nonlocal changed
            if not changed:
                changed = True
                svc.remove_document(doc["doc_id"])

        return build_index(*args, **kwargs, progress=progress)

    svc.build_function = build_with_change
    svc.enqueue_build()
    result = svc.process_next()
    assert result["error_code"] == "SERVICE_TARGET_STALE"
    assert svc.index_status()["generation_id"] == previous
    svc.build_function = build_index
    svc.run_once()
    assert svc.index_status()["documents"] == 0


def test_restart_recovers_jobs_only_after_owning_worker_lock(service):
    svc, source = service
    write_doc(source / "a.docx")
    svc.add_source(source, types=("docx",))
    claimed = svc.queue.claim("dead-process")
    with portalocker.Lock(svc.root / ".worker.lock", timeout=0):
        with pytest.raises(DocQAError) as error:
            svc.process_next()
        assert error.value.code == "SERVICE_ALREADY_RUNNING"
    restarted = Service(svc.root, options=svc.options, clock=svc.clock)
    restarted.run_once()
    assert restarted.queue.get(claimed["job_id"])["attempt"] == 2
    assert restarted.index_status()["documents"] == 1


def test_cancelled_parse_target_is_not_recreated(service):
    svc, source = service
    write_doc(source / "a.docx")
    svc.add_source(source, types=("docx",))
    svc.process_next()  # scan queues parse; scheduler not invoked
    job = next(item for item in svc.queue.list() if item["kind"] == "parse")
    svc.queue.cancel(job["job_id"])
    svc.run_once()
    assert svc.queue.latest(job["dedupe_key"])["state"] == "cancelled"
    svc.queue.retry(job["job_id"])
    svc.run_once()
    assert svc.index_status()["documents"] == 1


def test_mutation_rolls_back_source_document_and_revision_together(service, monkeypatch):
    svc, _, root, doc = ingest(service)
    revision = svc.registry.meta("desired_revision")
    original = svc.registry.put

    def crash(table, value):
        if table == "documents":
            raise RuntimeError("simulate process failure before transaction commit")
        original(table, value)

    monkeypatch.setattr(svc.registry, "put", crash)
    with pytest.raises(RuntimeError):
        svc.remove_source(root["source_id"])
    assert not svc.registry.get("sources", root["source_id"])["removed"]
    assert not svc.registry.get("documents", doc["doc_id"])["removed"]
    assert svc.registry.meta("desired_revision") == revision


@pytest.mark.parametrize("value", ["false", "[]", "42", "{unexpected: 1}"])
def test_invalid_configuration_root_rejected(tmp_path, value):
    path = tmp_path / "config.yaml"
    path.write_text(value, encoding="utf-8")
    with pytest.raises(DocQAError) as error:
        load_options(path)
    assert error.value.code == "INVALID_CONFIG"


def test_config_change_reparses_and_force_reanalysis_changes_artifact(service):
    svc, source, _, doc = ingest(service)
    initial = svc.document(doc["doc_id"])["published_artifact_id"]
    svc.reparse_images(doc["doc_id"])
    svc.run_once()
    assert svc.document(doc["doc_id"])["published_artifact_id"] != initial
    assert svc.images(doc["doc_id"])["figures"] == []
    svc.options.parse.analysis_revision = "2"
    svc.run_once()
    assert svc.document(doc["doc_id"])["target_satisfied"]
    assert svc.list_documents()[0]["analysis_epoch"] == 1


def test_cooperative_stop_between_ticks(service):
    svc, _ = service
    assert svc.run(max_ticks=1)["ticks"] == 1
    assert svc.request_stop()["stop_requested"]


def test_image_review_keeps_old_index_and_explicit_reanalysis_uses_new_cache_key(service):
    svc, source, _, doc = ingest(service)
    previous = svc.index_status()["generation_id"]
    package = Document()
    package.add_paragraph("SD5993 设备图示")
    buffer = BytesIO()
    Image.new("RGB", (80, 60), "white").save(buffer, format="PNG")
    buffer.seek(0)
    package.add_picture(buffer)
    package.save(source / "heart.docx")
    svc.enqueue_scan()
    svc.run_once()
    inspection = svc.images(doc["doc_id"])
    assert inspection["status"] == "review_required" and inspection["old_index_retained"]
    assert len(inspection["figures"]) == 1
    assert Path(inspection["figures"][0]["image_path"]).is_file()
    assert svc.index_status()["generation_id"] == previous

    class Vision:
        calls = 0

        def fingerprint(self):
            return {"adapter": "service-test-double", "revision": "1"}

        def describe(self, image, *, ocr_text="", context=""):
            self.calls += 1
            return VisionDescription(status="ok", observations=["测试用白色图片"])

    vision = Vision()

    def parse_with_vision(*args, **kwargs):
        return parse_document(*args, **kwargs, vision_client=vision)

    svc.parse_function = parse_with_vision
    svc.reparse_images(doc["doc_id"])
    assert svc.run_once()["failed"] == 0
    assert vision.calls == 1
    assert svc.images(doc["doc_id"])["figures"][0]["semantic_state"] == "described"
    assert svc.document(doc["doc_id"])["target_satisfied"]
    svc.reparse_images(doc["doc_id"])
    svc.run_once()
    assert vision.calls == 2


def test_pause_does_not_remove_old_members_and_single_file_does_not_scan_siblings(service):
    svc, source, root, doc = ingest(service)
    svc.update_source(root["source_id"], enabled=False)
    write_doc(source / "heart.docx", "改写但暂停摄入")
    svc.run_once()
    assert svc.index_status()["documents"] == 1
    svc.update_source(root["source_id"], enabled=True)
    svc.options.max_files = 1
    write_doc(source / "second.docx")
    job = svc.enqueue_parse(source / "heart.docx", force=True)
    assert job["kind"] == "parse"
    assert svc.document(doc["doc_id"])["desired_sha256"]
